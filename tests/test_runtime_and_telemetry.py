import ast
import inspect
import json
from pathlib import Path
import time
from types import SimpleNamespace

import pytest

import src.ui.control_panel as control_panel
import src.ui.canvas_gemini as canvas
import src.core.main as main
import pygame
import src.ui.telemetry_dashboard as telemetry_dashboard_module
from src.ui.canvas_gemini import H_Y, LANE
from src.core.signal_controller import SignalController
from src.ui.telemetry_dashboard import (
    HISTORY_MAX_POINTS,
    TelemetryDashboard,
    _gpu_none,
    build_excel_export_filename,
)
from src.telemetry import real_world_units as units
from src.telemetry.telemetry_exporter import (
    DEFAULT_TELEMETRY_PATH,
    DOWNSTREAM_BLOCKED_PX,
    TelemetryExporter,
)
from src.core.vehicle import Bus, DBL_LANE_INDEX, Vehicle
from tests.helpers import make_bus_for_leg, NODE_A, NODE_B


def lane_options():
    return {
        "EB": [H_Y - 0.5 * LANE, H_Y - 1.5 * LANE, H_Y - 2.5 * LANE],
        "WB": [H_Y + 0.5 * LANE, H_Y + 1.5 * LANE, H_Y + 2.5 * LANE],
    }


@pytest.fixture
def preserve_sim_random_state():
    state = main.random.getstate()
    yield
    main.random.setstate(state)
    main.reset_all_spawner_states()


def _collect_spawn_sequence(frame_count=240):
    vehicles = []
    config = {
        "model": "Poisson",
        "rate": 600,
        "turn_split": 0.7,
        "heavy_ratio": 0.25,
    }
    sequence = []
    for frame in range(frame_count):
        previous_count = len(vehicles)
        main.try_spawn_vehicle(
            vehicles,
            "EB",
            "EB",
            -20,
            lane_options()["EB"],
            config,
            min_gap=0,
        )
        if len(vehicles) == previous_count:
            continue
        vehicle = vehicles[-1]
        sequence.append(
            (
                frame,
                vehicle.is_heavy,
                vehicle.max_speed,
                vehicle.lane_index,
                vehicle.target_turn,
                vehicle.color,
            )
        )
    return sequence


def _seeded_spawn_sequence(seed):
    control_panel.global_config["random_seed"] = seed
    main.reset_traffic_generation()
    return _collect_spawn_sequence()


def test_same_seed_same_spawns(preserve_sim_random_state):
    first = _seeded_spawn_sequence(20260906)
    second = _seeded_spawn_sequence(20260906)

    assert first
    assert first == second


def test_demand_schedule_is_independent_of_other_random_consumers(
    preserve_sim_random_state,
):
    # Common random numbers: a controller arm that changes congestion changes
    # how often behavioural draws (lane changes) hit the global RNG. That must
    # not touch what any source offers, or baseline-vs-arm pairs compare two
    # different demand realisations.
    def run(perturb):
        control_panel.global_config["random_seed"] = 4242
        main.reset_traffic_generation()
        main._demand_draw_state["offers"].clear()
        vehicles, offered = [], []
        config = {"model": "Poisson", "rate": 600, "turn_split": 0.7, "heavy_ratio": 0.25}
        for _ in range(240):
            if perturb:
                main.random.random()
                main.spawner_states["WB"]["rng"].random()  # another source's stream
            main.try_spawn_vehicle(vehicles, "EB", "EB", -20, lane_options()["EB"], config, min_gap=0)
            offered.append(main.spawner_states["EB"]["requested_arrivals"])
        return offered, [(v.is_heavy, v.max_speed, v.lane_index, v.target_turn) for v in vehicles], main._demand_draw_hash_hex()

    assert run(False) == run(True)


def test_blocked_poisson_arrival_waits_at_the_source_instead_of_being_lost(monkeypatch):
    main.reset_all_spawner_states()
    state = main.spawner_states["EB"]
    monkeypatch.setattr(state["rng"], "random", lambda: 0.0)
    monkeypatch.setattr(state["rng"], "choice", lambda values: values[0])
    lanes = lane_options()
    config = {"model": "Poisson", "rate": 30, "turn_split": 1.0, "heavy_ratio": 0.0}
    vehicles = [Vehicle(-20, lanes["EB"][0], "EB", lane_index=0)]
    main.try_spawn_vehicle(vehicles, "EB", "EB", -20, lanes["EB"], config)
    assert state["requested_arrivals"] == 1 and state["pending_arrivals"] == 1
    assert len(vehicles) == 1
    vehicles.clear()
    main.try_spawn_vehicle(vehicles, "EB", "EB", -20, lanes["EB"], config)
    assert state["requested_arrivals"] == 2 and state["admitted_arrivals"] == 1
    assert state["pending_arrivals"] == 1
    main.reset_all_spawner_states()


def test_different_seed_differs(preserve_sim_random_state):
    first = _seeded_spawn_sequence(101)
    second = _seeded_spawn_sequence(202)

    assert first
    assert second
    assert first != second


def test_none_seed_nondeterministic(monkeypatch):
    seed_calls = []
    monkeypatch.setitem(control_panel.global_config, "random_seed", None)
    monkeypatch.setattr(main.random, "seed", lambda value: seed_calls.append(value))

    assert main.apply_configured_random_seed() is None
    assert seed_calls == []


def test_reset_reproduces_with_seed(preserve_sim_random_state):
    control_panel.global_config["random_seed"] = 8675309
    main.reset_traffic_generation()
    before_reset = _collect_spawn_sequence()

    # Consume more RNG and source state before exercising the production reset.
    _collect_spawn_sequence(60)
    main.reset_traffic_generation()
    after_reset = _collect_spawn_sequence()

    assert before_reset
    assert before_reset == after_reset
    assert all(value == 0.0 for value in main.bus_dispatch_counters.values())
    assert main.bus_sequence_counter == 0


def test_seed_config_accepts_integer_and_blank(monkeypatch):
    monkeypatch.setitem(control_panel.global_config, "random_seed", None)

    assert control_panel.set_random_seed(" 12345 ") == 12345
    assert control_panel.global_config["random_seed"] == 12345
    assert control_panel.set_random_seed("   ") is None
    assert control_panel.global_config["random_seed"] is None


def test_sim_launches_stopped():
    assert control_panel.global_config["is_running"] is False
    assert control_panel.global_config["start_requested"] is False
    source = inspect.getsource(main.main)
    assert "if is_running and not is_paused and not run_just_reset:" in source
    assert 'if not control_panel.global_config.get("is_running"' not in source


def test_start_performs_full_reset(tmp_path, monkeypatch):
    telemetry_log = tmp_path / "telemetry_log.jsonl"
    turn_log = tmp_path / "agent_turn_log.jsonl"
    telemetry_log.write_text("old telemetry\n", encoding="utf-8")
    turn_log.write_text("old turn\n", encoding="utf-8")
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", telemetry_log)
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", turn_log)
    seed_applications = []
    monkeypatch.setattr(
        main,
        "reset_traffic_generation",
        lambda: seed_applications.append(
            control_panel.global_config.get("random_seed")
        ),
    )

    class FakeSignals:
        frame_number = 99

        def reset_all_state(self):
            self.frame_number = 0

    class FakeTelemetry:
        reset_count = 0

        def reset_session(self):
            self.reset_count += 1

    signals = FakeSignals()
    telemetry = FakeTelemetry()
    vehicles = [object(), object()]
    for key in main.network_throughput:
        main.network_throughput[key] = 99
    monkeypatch.setitem(control_panel.global_config, "random_seed", 4242)

    frame = main.perform_full_reset(vehicles, signals, telemetry)

    assert frame == 0
    assert vehicles == []
    assert signals.frame_number == 0
    assert telemetry.reset_count == 1
    assert seed_applications == [4242]
    assert not telemetry_log.exists()
    assert not turn_log.exists()
    assert set(main.network_throughput.values()) == {0}


def test_stop_preserves_logs(tmp_path, monkeypatch):
    telemetry_log = tmp_path / "telemetry_log.jsonl"
    turn_log = tmp_path / "agent_turn_log.jsonl"
    telemetry_log.write_text("completed telemetry\n", encoding="utf-8")
    turn_log.write_text("completed turns\n", encoding="utf-8")
    monkeypatch.setitem(control_panel.global_config, "is_running", True)
    monkeypatch.setitem(control_panel.global_config, "start_requested", False)

    result = control_panel.request_start_stop()

    assert result == "STOPPED"
    assert control_panel.global_config["is_running"] is False
    assert telemetry_log.read_text(encoding="utf-8") == "completed telemetry\n"
    assert turn_log.read_text(encoding="utf-8") == "completed turns\n"


def test_start_after_stop_is_fresh(tmp_path, monkeypatch):
    telemetry_log = tmp_path / "telemetry_log.jsonl"
    turn_log = tmp_path / "agent_turn_log.jsonl"
    telemetry_log.write_text("run A\n", encoding="utf-8")
    turn_log.write_text("run A\n", encoding="utf-8")
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", telemetry_log)
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", turn_log)
    monkeypatch.setattr(main, "reset_traffic_generation", lambda: None)

    class FakeSignals:
        def reset_all_state(self):
            pass

    monkeypatch.setitem(control_panel.global_config, "is_running", True)
    assert control_panel.request_start_stop() == "STOPPED"
    assert telemetry_log.exists() and turn_log.exists()
    assert control_panel.request_start_stop() == "START_REQUESTED"

    frame = main.perform_full_reset([object()], FakeSignals())

    assert frame == 0
    assert not telemetry_log.exists()
    assert not turn_log.exists()


def test_pause_still_independent(monkeypatch):
    monkeypatch.setitem(control_panel.global_config, "is_running", True)
    monkeypatch.setitem(control_panel.global_config, "is_paused", False)

    assert control_panel.request_pause_resume() is True
    assert control_panel.global_config["is_running"] is True
    assert control_panel.request_pause_resume() is False
    assert control_panel.global_config["is_running"] is True

    assert control_panel.request_start_stop() == "STOPPED"
    assert control_panel.request_pause_resume() is False
    assert control_panel.global_config["is_running"] is False


def test_truck_passenger_count():
    truck = Vehicle(0, 0, "EB", is_heavy=True)
    car = Vehicle(0, 0, "EB", is_heavy=False)
    bus = make_bus_for_leg("R1_EB_A_NB", NODE_A, "PAX_BUS")

    assert truck.passengers == 1
    assert car.passengers == 4
    assert isinstance(bus, Bus)
    assert bus.passengers == 45


def test_throughput_counts_truck_as_one(monkeypatch):
    truck = Vehicle(canvas.WIDTH + 61, H_Y - 0.5 * LANE, "EB", is_heavy=True)
    truck.passed_nodes.add(canvas.INT_X[0])
    truck.update = lambda **kwargs: None

    class FakeRoot:
        def geometry(self, _value):
            pass

        def protocol(self, _name, _handler):
            pass

        def after(self, _delay, callback):
            self.callback = callback

        def mainloop(self):
            self.callback()

    class FakePane:
        """Stand-in for a mounted pane frame: no real widget behind it."""

        def __init__(self):
            self.scroll_canvas = object()

        def winfo_children(self):
            return []

        def bind(self, *_args, **_kwargs):
            pass

    class FakeProcess:
        def terminate(self):
            pass

        def wait(self, timeout=None):
            pass

    class FakeSignals:
        def __init__(self, *args, **kwargs):
            pass

        def update(self, vehicles=None):
            pass

        def is_discharge_active(self):
            return False

        def get_all_signals(self, _nodes):
            return {}

        def get_all_dbl_states(self, _nodes, _vehicles):
            return {}

    class FakeTelemetry:
        def __init__(self, *args, **kwargs):
            pass

        def export(self, **kwargs):
            return False

    spawned = False

    def spawn_truck(vehicles, *args, **kwargs):
        nonlocal spawned
        if not spawned:
            vehicles.append(truck)
            spawned = True

    approaches = {
        key: {"active": key == "EB"}
        for key in ("EB", "WB", "A_NB", "A_SB", "B_NB", "B_SB")
    }
    # Loop start, tick start, tick end (the re-arm measures its own work).
    monotonic_values = iter((100.0, 100.02, 100.025))

    monkeypatch.setattr(main, "reset_session_logs", lambda: None)
    monkeypatch.setattr(main, "SignalController", FakeSignals)
    monkeypatch.setattr(main, "TelemetryExporter", FakeTelemetry)
    monkeypatch.setattr(main.control_panel, "approach_configs", approaches)
    # The unified window shell (PanedWindow, scrollable panes, pygame-blit
    # canvas) is real Tk widget construction and needs a real Tk() as its
    # parent, which FakeRoot deliberately is not -- it exists only to let
    # mainloop() return after exactly one simulation_step() instead of
    # blocking on a real event loop. This test is about vehicle/telemetry
    # accounting, not window layout, so the whole shell builder is replaced
    # the same way create_dashboard_window always has been mocked out.
    monkeypatch.setattr(
        main, "build_main_window",
        lambda: (FakeRoot(), FakePane(), FakePane(), FakePane()),
    )
    monkeypatch.setattr(
        main.control_panel, "create_dashboard_window", lambda _pane: None
    )
    monkeypatch.setattr(main, "bind_pane_mousewheel", lambda *_a, **_k: None)
    # The telemetry dashboard now mounts in-process too (no longer a
    # subprocess), so main() constructs it directly; same reasoning as
    # create_dashboard_window above.
    monkeypatch.setattr(main, "TelemetryDashboard", lambda _pane: None)
    # The simulation canvas is real Tk widget construction too (a tk.Canvas
    # parented to simulation_pane), same reasoning as the other three mounted
    # components above -- FakePane is not a real widget and cannot be a
    # valid Tk master.
    class FakeCanvas:
        """The real canvas answers frame_is_due(): the loop renders only on
        a due tick. Always due here, so the assertions see every frame."""

        def frame_is_due(self):
            return True

    monkeypatch.setattr(
        main, "build_simulation_canvas",
        lambda _pane: (FakeCanvas(), lambda _surface, _draw=None: None),
    )
    monkeypatch.setattr(main.subprocess, "Popen", lambda *args, **kwargs: FakeProcess())
    monkeypatch.setattr(main.atexit, "register", lambda callback: callback)
    monkeypatch.setattr(main.time, "monotonic", lambda: next(monotonic_values))
    monkeypatch.setattr(main, "try_spawn_vehicle", spawn_truck)
    monkeypatch.setattr(main, "check_and_dispatch_buses", lambda *args: None)
    monkeypatch.setattr(main, "get_demand_telemetry", lambda: {})
    monkeypatch.setattr(main.canvas, "draw_network", lambda *args, **kwargs: None)
    monkeypatch.setattr(main.pygame, "init", lambda: None)
    monkeypatch.setattr(main.pygame.font, "init", lambda: None)
    monkeypatch.setattr(main.pygame.font, "SysFont", lambda *args, **kwargs: None)
    # No pygame window exists any more (the simulation renders onto an
    # offscreen Surface), so display.Info/set_mode/set_caption/flip and the
    # old QUIT-handling event loop are no longer called by main() at all --
    # closing is now MainWindow's single WM_DELETE_WINDOW handler instead.
    monkeypatch.setitem(main.control_panel.global_config, "is_running", True)
    monkeypatch.setitem(main.control_panel.global_config, "reset_triggered", False)
    monkeypatch.setitem(main.control_panel.global_config, "is_paused", False)
    monkeypatch.setitem(main.control_panel.global_config, "sim_speed", 1.0)
    for key in main.network_throughput:
        monkeypatch.setitem(main.network_throughput, key, 0)

    main.main()

    assert spawned is True
    assert truck.passed_nodes == {canvas.INT_X[0]}
    assert main.network_throughput["passengers_served_car"] == 1
    assert main.network_throughput["passengers_served_total"] == 1
    assert main.network_throughput["cars_served"] == 1
    assert main.network_throughput["vehicles_served_total"] == 1


def test_bus_dispatch_is_lane_aware():
    for config in control_panel.bus_routes_config.values():
        config["active"] = False
        config["manual_dispatch"] = False
    config = control_panel.bus_routes_config["R1_EB_A_NB"]
    config["active"] = True
    config["manual_dispatch"] = True
    vehicles = [Vehicle(-40, H_Y - 0.5 * LANE, "EB", lane_index=0)]
    main.check_and_dispatch_buses(vehicles, lane_options(), 1 / 60)
    assert len(vehicles) == 2
    assert vehicles[-1].route_id == "R1_EB_A_NB"


def test_zero_headway_resets_retained_time():
    route_id = "R1_EB_A_NB"
    for config in control_panel.bus_routes_config.values():
        config["active"] = False
    config = control_panel.bus_routes_config[route_id]
    config["active"] = True
    config["headway_sec"] = 0
    main.bus_dispatch_counters[route_id] = 12.5
    main.check_and_dispatch_buses([], lane_options(), 1 / 60)
    assert main.bus_dispatch_counters[route_id] == 0


def test_telemetry_uses_authoritative_passengers_and_per_node_signals(tmp_path):
    controller = SignalController({"green_time": 20})
    car = Vehicle(0, H_Y - 0.5 * LANE, "EB")
    exporter = TelemetryExporter(tmp_path / "state.json", 1)
    payload = exporter.build_payload(controller, [car], 60)
    assert payload["schema_version"] == 3
    assert payload["network_summary"]["passenger_volume"] == car.passengers
    assert set(payload["signal_state"]["nodes"]) == {str(NODE_A), str(NODE_B)}
    assert payload["signal_state"]["timing"] == {
        "frames_per_second": 60,
        "green_frames": 20,
        "yellow_frames": 60,
        "all_red_frames": 60,
        "nominal_cycle_frames": 280,
    }
    assert payload["simulation_time_seconds"] == 1.0


def test_telemetry_exposes_per_node_approach_passenger_queues(tmp_path):
    controller = SignalController({"green_time": 20})
    exporter = TelemetryExporter(tmp_path / "state.json", 1)
    node_a_eb = Vehicle(NODE_A - 100, H_Y - 0.5 * LANE, "EB", lane_index=0)
    node_b_eb = Vehicle(NODE_B - 100, H_Y - 0.5 * LANE, "EB", lane_index=0)
    node_b_eb.passed_nodes.add(canvas.INT_X[0])
    node_a_nb = Vehicle(
        canvas.INT_X[0] - 0.5 * LANE,
        H_Y + 100,
        "NB",
        lane_index=0,
        assigned_node_x=canvas.INT_X[0],
    )
    for vehicle in (node_a_eb, node_b_eb, node_a_nb):
        vehicle.speed = 0.0

    payload = exporter.build_payload(
        controller, [node_a_eb, node_b_eb, node_a_nb], 60
    )
    node_a = payload["signal_state"]["nodes"][str(NODE_A)]
    node_b = payload["signal_state"]["nodes"][str(NODE_B)]

    assert node_a["queues_passengers_est"] == {
        "EB": 4, "WB": 0, "NB": 4, "SB": 0
    }
    assert node_a["total_waiting_passengers_est"] == 8
    assert node_b["queues_passengers_est"] == {
        "EB": 4, "WB": 0, "NB": 0, "SB": 0
    }
    assert node_b["total_waiting_passengers_est"] == 4
    assert payload["network_summary"]["queues_passengers_est_by_node"] == {
        str(NODE_A): node_a["queues_passengers_est"],
        str(NODE_B): node_b["queues_passengers_est"],
    }


def _queued_car(x, direction="EB", is_heavy=False):
    vehicle = Vehicle(
        x, H_Y - 0.5 * LANE, direction, is_heavy=is_heavy, lane_index=0
    )
    vehicle.speed = 0.0
    return vehicle


def _queued_nb(x_offset, node_x, is_heavy=False):
    vehicle = Vehicle(
        node_x - 0.5 * LANE,
        H_Y + 100 + x_offset,
        "NB",
        is_heavy=is_heavy,
        lane_index=0,
        assigned_node_x=node_x,
    )
    vehicle.speed = 0.0
    return vehicle


def test_queue_passengers_uses_real_weights():
    exporter = TelemetryExporter(export_interval_frames=1)
    controller = SignalController({"green_time": 20})
    vehicles = [
        _queued_car(NODE_A - 100),
        _queued_car(NODE_A - 130),
        _queued_car(NODE_A - 160, is_heavy=True),
    ]
    assert [v.passengers for v in vehicles] == [4, 4, 1]

    payload = exporter.build_payload(controller, vehicles, 60)
    summary = payload["network_summary"]

    assert summary["queues"]["EB"] == 3
    assert summary["queues_passengers_est"]["EB"] == 2 * 4 + 1 * 1
    assert summary["queues_passengers_est"]["EB"] != 3 * 4
    assert summary["queues_passengers_est_by_node"][str(NODE_A)]["EB"] == 9


def test_queued_bus_counts_full_load():
    exporter = TelemetryExporter(export_interval_frames=1)
    controller = SignalController({"green_time": 20})
    bus = make_bus_for_leg("R1_EB_A_NB", NODE_A, "QUEUED_BUS")
    bus.speed = 0.0
    assert bus.passengers == 45

    by_node = exporter.compute_queue_passengers_by_node([bus])
    assert by_node[str(NODE_A)]["EB"] == 45

    payload = exporter.build_payload(controller, [bus], 60)
    summary = payload["network_summary"]
    assert summary["queues"]["EB"] == 1
    assert summary["queues_passengers_est"]["EB"] == 45
    node_a = payload["signal_state"]["nodes"][str(NODE_A)]
    assert node_a["queues_passengers_est"]["EB"] == 45
    assert node_a["total_waiting_passengers_est"] == 45


def test_queue_vehicle_count_unchanged():
    exporter = TelemetryExporter(export_interval_frames=1)
    bus = make_bus_for_leg("R1_EB_A_NB", NODE_A, "COUNT_BUS")
    bus.speed = 0.0
    vehicles = [bus, _queued_car(NODE_A - 150), _queued_car(NODE_A - 180, is_heavy=True)]

    counts = exporter.compute_queue_counts_by_node(vehicles)
    passengers = exporter.compute_queue_passengers_by_node(vehicles)

    # Three vehicles regardless of type; passengers reflect 45 + 4 + 1.
    assert counts[str(NODE_A)]["EB"] == 3
    assert passengers[str(NODE_A)]["EB"] == 50
    assert exporter.compute_queue_counts(vehicles)["EB"] == 3


def test_bus_distribution_by_node_approach_and_route():
    exporter = TelemetryExporter(export_interval_frames=1)
    # Mid-route leg (straight lane 1), not the route's natural left-turn/DBL
    # lane, so the DBL-lane count below isolates the bus explicitly moved in.
    bus_a = make_bus_for_leg("R2_EB_B_NB", NODE_A, "DIST_BUS_A")
    bus_b = make_bus_for_leg("R3_EB_ONLY", NODE_B, "DIST_BUS_B")
    bus_b.lane_index = DBL_LANE_INDEX
    car = _queued_car(NODE_A - 150)

    distribution = exporter.compute_bus_distribution([bus_a, bus_b, car])

    assert distribution["total_buses"] == 2
    assert distribution["buses_by_node_approach"][str(NODE_A)]["EB"] == 1
    assert distribution["buses_by_node_approach"][str(NODE_B)]["EB"] == 1
    assert distribution["buses_by_route"] == {
        "R2_EB_B_NB": 1,
        "R3_EB_ONLY": 1,
    }
    assert distribution["buses_in_dbl_lane"] == 1

    controller = SignalController({"green_time": 20})
    payload = exporter.build_payload(controller, [bus_a, bus_b, car], 60)
    assert payload["bus_distribution"] == distribution


def test_queue_passengers_per_node():
    exporter = TelemetryExporter(export_interval_frames=1)
    controller = SignalController({"green_time": 20})
    node_a_x, node_b_x = canvas.INT_X
    node_a_eb_truck = _queued_car(NODE_A - 100, is_heavy=True)
    node_b_eb_car = _queued_car(NODE_B - 100)
    node_b_eb_car.passed_nodes.add(node_a_x)
    node_a_nb_car = _queued_nb(0, node_a_x)
    node_b_nb_truck = _queued_nb(0, node_b_x, is_heavy=True)
    vehicles = [node_a_eb_truck, node_b_eb_car, node_a_nb_car, node_b_nb_truck]

    payload = exporter.build_payload(controller, vehicles, 60)
    node_a = payload["signal_state"]["nodes"][str(node_a_x)]
    node_b = payload["signal_state"]["nodes"][str(node_b_x)]

    assert node_a["queues"] == {"EB": 1, "WB": 0, "NB": 1, "SB": 0}
    assert node_a["queues_passengers_est"] == {"EB": 1, "WB": 0, "NB": 4, "SB": 0}
    assert node_a["total_waiting_passengers_est"] == 5
    assert node_b["queues"] == {"EB": 1, "WB": 0, "NB": 1, "SB": 0}
    assert node_b["queues_passengers_est"] == {"EB": 4, "WB": 0, "NB": 1, "SB": 0}
    assert node_b["total_waiting_passengers_est"] == 5

    summary = payload["network_summary"]
    assert summary["queues_passengers_est"] == {
        "EB": 5, "WB": 0, "A_NB": 4, "A_SB": 0, "B_NB": 1, "B_SB": 0
    }
    assert summary["queues_passengers_est_by_node"] == {
        str(node_a_x): node_a["queues_passengers_est"],
        str(node_b_x): node_b["queues_passengers_est"],
    }


def test_empty_approach_zero():
    exporter = TelemetryExporter(export_interval_frames=1)
    controller = SignalController({"green_time": 20})
    moving_car = Vehicle(NODE_A - 100, H_Y - 0.5 * LANE, "EB", lane_index=0)
    moving_car.speed = 1.0

    for vehicles in ([], [moving_car]):
        payload = exporter.build_payload(controller, vehicles, 60)
        summary = payload["network_summary"]
        assert all(value == 0 for value in summary["queues"].values())
        assert all(value == 0 for value in summary["queues_passengers_est"].values())
        for node in payload["signal_state"]["nodes"].values():
            assert node["queues_passengers_est"] == {
                "EB": 0, "WB": 0, "NB": 0, "SB": 0
            }
            assert node["total_waiting_passengers_est"] == 0


def _downstream_eb_cars(x_positions, lanes=(0,)):
    """Stopped EB cars on the road between the two nodes (past Node A)."""
    cars = []
    for lane_index in lanes:
        for x in x_positions:
            car = Vehicle(
                x, H_Y - (lane_index + 0.5) * LANE, "EB", lane_index=lane_index
            )
            car.passed_nodes.add(canvas.INT_X[0])
            car.speed = 0.0
            cars.append(car)
    return cars


def test_queue_length_metres_zero_when_nothing_queued():
    exporter = TelemetryExporter(export_interval_frames=1)
    controller = SignalController({"green_time": 20})
    moving_car = Vehicle(NODE_A - 100, H_Y - 0.5 * LANE, "EB", lane_index=0)
    moving_car.speed = 1.0

    for vehicles in ([], [moving_car]):
        by_node = exporter.compute_queue_length_by_node(vehicles)
        assert by_node == {
            str(NODE_A): {"EB": 0.0, "WB": 0.0, "NB": 0.0, "SB": 0.0},
            str(NODE_B): {"EB": 0.0, "WB": 0.0, "NB": 0.0, "SB": 0.0},
        }
        payload = exporter.build_payload(controller, vehicles, 60)
        assert payload["network_summary"]["queue_length_m_by_node"] == by_node
        for node in payload["signal_state"]["nodes"].values():
            assert node["queue_length_m"] == {
                "EB": 0.0, "WB": 0.0, "NB": 0.0, "SB": 0.0
            }


def test_queue_length_metres_matches_furthest_queued_vehicle():
    exporter = TelemetryExporter(export_interval_frames=1)
    controller = SignalController({"green_time": 20})
    # Three stopped EB cars short of Node A; the one at x=NODE_A - 160 is the tail.
    vehicles = [_queued_car(NODE_A - 100), _queued_car(NODE_A - 130), _queued_car(NODE_A - 160)]
    tail_px = max(
        v.distance_to_node_stop_bar(NODE_A, H_Y, canvas.ROAD_W, canvas.STOP)
        for v in vehicles
    )
    assert tail_px == vehicles[-1].distance_to_node_stop_bar(
        NODE_A, H_Y, canvas.ROAD_W, canvas.STOP
    )

    by_node = exporter.compute_queue_length_by_node(vehicles)
    assert by_node[str(NODE_A)]["EB"] == pytest.approx(units.px_to_m(tail_px), abs=0.1)
    assert by_node[str(NODE_A)]["EB"] > 0
    assert by_node[str(NODE_A)]["WB"] == 0.0
    assert by_node[str(NODE_B)]["EB"] == 0.0

    # A longer queue reports a longer length; the value rides the payload.
    longer = vehicles + [_queued_car(NODE_A - 220)]
    assert (
        exporter.compute_queue_length_by_node(longer)[str(NODE_A)]["EB"]
        > by_node[str(NODE_A)]["EB"]
    )
    payload = exporter.build_payload(controller, vehicles, 60)
    node_a = payload["signal_state"]["nodes"][str(NODE_A)]
    assert node_a["queue_length_m"] == by_node[str(NODE_A)]
    assert payload["network_summary"]["queue_length_m_by_node"] == by_node


def test_downstream_space_decreases_as_downstream_road_fills():
    exporter = TelemetryExporter(export_interval_frames=1)
    empty_space, empty_blocked = exporter.compute_downstream_space_by_node([])
    # Empty network: EB at Node A sees the full node-to-node stretch, EB at
    # Node B sees the stretch to the canvas edge; nothing is blocked.
    node_gap_px = canvas.INT_X[1] - canvas.INT_X[0] - canvas.ROAD_W
    edge_px = canvas.WIDTH - canvas.INT_X[1] - canvas.ROAD_W / 2
    assert empty_space[str(NODE_A)]["EB"] == pytest.approx(
        units.px_to_m(node_gap_px), abs=0.1
    )
    assert empty_space[str(NODE_B)]["EB"] == pytest.approx(
        units.px_to_m(edge_px), abs=0.1
    )
    assert empty_space[str(NODE_A)]["WB"] == empty_space[str(NODE_B)]["EB"]
    assert empty_space[str(NODE_B)]["WB"] == empty_space[str(NODE_A)]["EB"]
    assert not any(any(row.values()) for row in empty_blocked.values())

    # A standing queue whose tail creeps back toward Node A's box lowers the
    # EB receiving space of that lane, and touches nothing else. Moving
    # traffic on the stretch is not spillback and takes no room.
    previous = empty_space[str(NODE_A)]["EB"]
    for tail_x in (NODE_B - 200, NODE_B - 400, NODE_B - 600):
        cars = _downstream_eb_cars(range(tail_x, NODE_B - 66, 30), lanes=(0, 1, 2))
        space, _blocked = exporter.compute_downstream_space_by_node(cars)
        assert space[str(NODE_A)]["EB"] < previous
        previous = space[str(NODE_A)]["EB"]
        assert space[str(NODE_A)]["WB"] == empty_space[str(NODE_A)]["WB"]
        assert space[str(NODE_B)]["EB"] == empty_space[str(NODE_B)]["EB"]
        assert space[str(NODE_A)]["NB"] == empty_space[str(NODE_A)]["NB"]
    moving = _downstream_eb_cars(range(NODE_A + 80, NODE_B - 66, 30), lanes=(0, 1, 2))
    for car in moving:
        car.speed = 1.0
    assert exporter.compute_downstream_space_by_node(moving)[0] == empty_space


def test_receiving_space_is_per_lane_and_per_movement():
    """The audit's case: a left-turning bus whose exit lane is full while
    the approach's straight lanes are empty. The approach reads open, the
    bus's own receiving lane reads blocked, and the TSP gate refuses."""
    from src.core.signal_controller import TSP_DENY_DOWNSTREAM_BLOCKED
    from src.core.vehicle import receiving_space_px
    exporter = TelemetryExporter(export_interval_frames=1)
    # NB exit of an EB left at Node A lands on lane 2 of the NB road,
    # x = NODE_A - 2.5 lanes; queue it solid from the box edge northward.
    exit_x = NODE_A - 2.5 * LANE
    queue = []
    for y in range(int(H_Y - 66 - 9), 0, -30):
        car = Vehicle(exit_x, y, "NB", lane_index=2)
        car.speed = 0.0
        queue.append(car)
    by_lane = exporter.compute_downstream_space_by_lane(queue)
    assert by_lane[str(NODE_A)]["EB"] == [pytest.approx(NODE_B - NODE_A - 132)] * 3
    assert by_lane[str(NODE_A)]["NB"][2] < DOWNSTREAM_BLOCKED_PX
    assert by_lane[str(NODE_A)]["NB"][0] > DOWNSTREAM_BLOCKED_PX
    _space, blocked = exporter.compute_downstream_space_by_node(queue)
    assert blocked[str(NODE_A)]["EB"] is False and blocked[str(NODE_A)]["NB"] is False
    # Movement-specific: EB straight has the whole link, EB left has nothing.
    assert receiving_space_px(NODE_A, "EB", "STRAIGHT", H_Y - 1.5 * LANE, queue) > 600
    assert receiving_space_px(NODE_A, "EB", "LEFT", H_Y - 2.5 * LANE, queue) < 1.0

    bus = make_bus_for_leg("R1_EB_A_NB", NODE_A, "LEFT_BUS")  # EB left at A
    bus.x = NODE_A - 66 - 10 - bus.length / 2 - 150
    route = control_panel.bus_routes_config["R1_EB_A_NB"]
    route["active"] = route["tsp_enabled"] = True
    controller = SignalController({"green_time": 240, "is_running": True}, yellow_time=60, red_clearance_time=60)
    vehicles = queue + [bus]
    controller.update(vehicles)
    request = controller.nodes[NODE_A].active_request
    assert request is not None and request.tsp_requested
    payload = exporter.build_payload(controller, vehicles, 60)
    bus_row = next(b for b in payload["active_buses"] if b["bus_id"] == "LEFT_BUS")
    assert bus_row["receiving_blocked"] is True and bus_row["receiving_space_m"] == 0.0
    # Cross street green now: without the gate the bus's early green would cut it.
    controller.nodes[NODE_A].phase = 3
    controller.nodes[NODE_A].timer = 0
    for _ in range(30):
        controller.update(vehicles)
    assert request.tsp_gate_reason == TSP_DENY_DOWNSTREAM_BLOCKED
    assert controller.nodes[NODE_A].priority_state == "NORMAL"
    assert request.tsp_action == "none"
    # The lane clears: the gate lifts on the next frame.
    for car in queue:
        car.speed = 1.0
    controller.update(vehicles)
    assert request.tsp_gate_reason != TSP_DENY_DOWNSTREAM_BLOCKED


def test_downstream_blocked_when_downstream_road_nearly_full():
    exporter = TelemetryExporter(export_interval_frames=1)
    controller = SignalController({"green_time": 20})
    # Every lane of the EB road between the nodes packed nose to tail.
    jammed = _downstream_eb_cars(range(NODE_A + 80, NODE_B - 66, 30), lanes=(0, 1, 2))

    space, blocked = exporter.compute_downstream_space_by_node(jammed)
    assert space[str(NODE_A)]["EB"] < units.px_to_m(DOWNSTREAM_BLOCKED_PX)
    assert blocked[str(NODE_A)]["EB"] is True
    assert blocked[str(NODE_A)]["WB"] is False
    assert blocked[str(NODE_B)]["EB"] is False

    # One lane's worth of cars is not a blockage.
    _space, partial = exporter.compute_downstream_space_by_node(
        _downstream_eb_cars(range(NODE_A + 80, NODE_B - 66, 30))
    )
    assert partial[str(NODE_A)]["EB"] is False

    payload = exporter.build_payload(controller, jammed, 60)
    node_a = payload["signal_state"]["nodes"][str(NODE_A)]
    assert node_a["downstream_blocked"]["EB"] is True
    assert node_a["downstream_space_m"] == space[str(NODE_A)]
    summary = payload["network_summary"]
    assert summary["downstream_space_m_by_node"] == space
    assert summary["downstream_blocked_by_node"] == blocked
    # The downstream signal never touches the queue-based fields.
    assert node_a["queues"]["EB"] == 0


def test_telemetry_includes_congestion_demand_backlog(tmp_path):
    controller = SignalController({"green_time": 20})
    exporter = TelemetryExporter(tmp_path / "state.json", 1)
    demand = {
        "EB": {
            "model": main.CONGESTION_MODEL,
            "pending_arrivals": 7,
            "peak_active": True,
        },
        "WB": {"model": "Poisson", "pending_arrivals": 0},
    }

    payload = exporter.build_payload(
        controller, [], 1, demand_state=demand
    )

    assert payload["network_summary"]["pending_demand"] == 7
    assert payload["demand_generation"] == demand


def test_telemetry_export_atomically_writes_valid_schema(tmp_path):
    destination = tmp_path / "state.json"
    exporter = TelemetryExporter(destination, 1)
    controller = SignalController({"green_time": 20})
    throughput = {
        "passengers_served_total": 49,
        "passengers_served_bus": 45,
        "passengers_served_car": 4,
        "vehicles_served_total": 2,
        "buses_served": 1,
        "cars_served": 1,
    }
    assert exporter.export(controller, [], 3600, throughput_state=throughput)
    payload = json.loads(destination.read_text(encoding="utf-8"))
    assert payload["schema_version"] == 3
    assert payload["frame_number"] == 3600
    assert payload["network_throughput"]["passengers_served_total"] == 49
    assert payload["network_throughput"]["passengers_per_minute"] == 49.0
    assert not list(tmp_path.glob("tmp*"))


def test_session_telemetry_log_preserves_full_audit_snapshot(tmp_path, monkeypatch):
    destination = tmp_path / "telemetry.jsonl"
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", destination)
    monkeypatch.setattr(main, "_last_telemetry_log_frame", None)
    payload = {
        "timestamp": 123.5,
        "simulation_time_seconds": 1.0,
        "network_summary": {"total_vehicles": 1, "queues": {"EB": 1}},
        "network_throughput": {"passengers_served_total": 0},
        "vehicle_positions": [
            {"snapshot_id": "car-1", "x_px": 10.0, "y_px": 20.0}
        ],
    }

    assert main.log_telemetry_sample(main.TELEMETRY_LOG_INTERVAL, payload) is True
    logged = json.loads(destination.read_text(encoding="utf-8"))
    assert logged["timestamp"] == 123.5
    assert logged["telemetry_snapshot"] == payload


def test_schema_v3_exposes_bus_eta_routes_and_passenger_weighted_queues():
    controller = SignalController({"green_time": 20})
    exporter = TelemetryExporter(export_interval_frames=1)
    bus = make_bus_for_leg("R1_EB_A_NB", NODE_A, "ETA_BUS")
    bus.speed = 0.0
    car = Vehicle(0, H_Y - 0.5 * LANE, "EB")
    car.speed = 0.0
    throughput = {
        "passengers_served_total": 49,
        "passengers_served_bus": 45,
        "passengers_served_car": 4,
        "vehicles_served_total": 2,
        "buses_served": 1,
        "cars_served": 1,
    }

    payload = exporter.build_payload(
        controller,
        [bus, car],
        frame_number=3600,
        throughput_state=throughput,
    )

    assert payload["schema_version"] == 3
    assert {
        "signal_state",
        "active_buses",
        "approaching_buses",
        "vehicle_positions",
        "demand_generation",
        "network_discharge",
    } <= payload.keys()
    assert set(payload["routes"]) == set(control_panel.bus_routes_config)

    bus_state = payload["active_buses"][0]
    assert 0 <= bus_state["eta_to_stop_bar_sec_freeflow"] < float("inf")
    assert 0 <= bus_state["eta_to_stop_bar_sec_live"] <= 100.0

    route = payload["routes"]["R1_EB_A_NB"]
    assert route["buses_on_route"] == 1
    assert route["route_passengers_total"] == bus.passengers
    assert route["nearest_bus_id"] == "ETA_BUS"
    assert route["nearest_bus_eta_sec"] == bus_state["eta_to_stop_bar_sec_freeflow"]
    assert bus_state["lane_index"] == bus.lane_index
    assert bus_state["in_dbl_lane"] is True
    positions = payload["vehicle_positions"]
    assert len(positions) == 2
    assert positions[0]["snapshot_id"] == "ETA_BUS"
    assert positions[0]["vehicle_type"] == "bus"
    assert positions[0]["route_id"] == "R1_EB_A_NB"
    assert positions[0]["x_px"] == pytest.approx(bus.x)
    assert positions[1]["vehicle_type"] == "car"
    assert positions[1]["lane_index"] == car.lane_index
    assert positions[1]["speed_px_per_frame"] == 0.0
    summary = payload["network_summary"]
    assert summary["queues_passengers_est"] == exporter._flatten_queue_counts(
        exporter.compute_queue_passengers_by_node([bus, car])
    )
    assert "car_occupancy_assumed" not in summary


def test_recent_throughput_uses_rolling_window_and_clears_on_reset():
    controller = SignalController({"green_time": 20})
    exporter = TelemetryExporter(export_interval_frames=1)

    initial = exporter.build_payload(
        controller, [], frame_number=0, throughput_state={"passengers_served_total": 0}
    )
    rising = exporter.build_payload(
        controller,
        [],
        frame_number=600,
        throughput_state={"passengers_served_total": 100},
    )
    recent = exporter.build_payload(
        controller,
        [],
        frame_number=2400,
        throughput_state={"passengers_served_total": 160},
    )
    reset = exporter.build_payload(
        controller,
        [],
        frame_number=2460,
        throughput_state={"passengers_served_total": 0},
    )

    assert initial["network_throughput"]["passengers_per_minute_recent"] == 0.0
    assert rising["network_throughput"]["passengers_per_minute_recent"] == 600.0
    assert recent["network_throughput"]["passengers_per_minute_recent"] == 120.0
    assert reset["network_throughput"]["passengers_per_minute_recent"] == 0.0
    assert exporter._throughput_samples == [(41.0, 0)]


def test_dashboard_freshness_states():
    now = time.time()
    assert TelemetryDashboard.classify_status(None, now) == "CONNECTING"
    assert TelemetryDashboard.classify_status({"timestamp": now}, now) == "LIVE"
    assert TelemetryDashboard.classify_status({"timestamp": now, "simulation_paused": True}, now) == "PAUSED"
    assert TelemetryDashboard.classify_status({"timestamp": now - 10}, now) == "STALE"


def dashboard_sample(frame, simulation_time, vehicles=10, buses=1, queued=2, pending=3):
    return {
        "frame_number": frame,
        "simulation_time_seconds": simulation_time,
        "network_summary": {
            "total_vehicles": vehicles,
            "total_buses": buses,
            "passenger_volume": 0,
            "queues": {"EB": queued},
            "pending_demand": pending,
        },
    }


def test_dashboard_history_downsamples_duplicates_and_clears_on_reset():
    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.initialize_history_state()
    dashboard.draw_trend_charts = lambda: None

    assert dashboard.record_history_sample(dashboard_sample(60, 1.0))
    assert not dashboard.record_history_sample(dashboard_sample(60, 1.0))
    assert not dashboard.record_history_sample(dashboard_sample(90, 1.5))
    assert dashboard.record_history_sample(dashboard_sample(120, 2.0, queued=4))
    assert list(dashboard.history["road_queue"]) == [2, 4]

    assert dashboard.record_history_sample(
        dashboard_sample(0, 0.0, vehicles=4, buses=0, queued=1, pending=0)
    )
    assert list(dashboard.history["time"]) == [0.0]
    assert list(dashboard.history["vehicles"]) == [4]
    assert list(dashboard.queue_history) == [1]


def test_dashboard_history_is_bounded_in_memory():
    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.initialize_history_state()
    dashboard.draw_trend_charts = lambda: None

    for second in range(HISTORY_MAX_POINTS + 25):
        dashboard.record_history_sample(
            dashboard_sample(second * 60, float(second), queued=second)
        )

    assert len(dashboard.history["time"]) == HISTORY_MAX_POINTS
    assert dashboard.history["time"][0] == 25.0
    assert len(dashboard.queue_history) == 20


def test_dashboard_initial_window_fits_smaller_screens_and_remains_useful():
    # The dense summary layout fits in 430px, so the dashboard no longer
    # claims 780px of monitor height it does not need.
    assert TelemetryDashboard.initial_window_size(1920, 1080) == (900, 430)
    assert TelemetryDashboard.initial_window_size(1366, 768) == (900, 430)
    assert TelemetryDashboard.initial_window_size(640, 480) == (560, 360)


def test_startup_window_layout_tiles_large_and_standard_hd_desktops():
    large = main.calculate_startup_window_layout(2560, 1440)
    hd = main.calculate_startup_window_layout(1920, 1080)

    assert large == {
        "mode": "tiled",
        "canvas_position": (900, 30),
        "control_geometry": "880x1030+10+10",
        "telemetry_geometry": "1000x430+900+640",
    }
    assert hd == {
        "mode": "tiled",
        "canvas_position": (900, 30),
        "control_geometry": "880x1030+10+10",
        "telemetry_geometry": "1000x400+900+640",
    }


def test_startup_window_layout_uses_on_screen_cascade_when_space_is_small():
    layout = main.calculate_startup_window_layout(1366, 768)

    assert layout["mode"] == "cascade"
    assert layout["canvas_position"] == (183, 30)
    assert layout["control_geometry"] == "880x718+476+10"
    assert layout["telemetry_geometry"] == "900x430+233+30"


def test_dashboard_responsive_profile_shrinks_content_without_scrollbars():
    full = TelemetryDashboard.responsive_profile(900, 780)
    compact = TelemetryDashboard.responsive_profile(560, 500)

    assert not full["compact"]
    assert compact["compact"]
    assert compact["header_font"] < full["header_font"]
    assert compact["metric_value_font"] < full["metric_value_font"]
    assert compact["metric_row_height"] < full["metric_row_height"]
    assert compact["phase_height"] < full["phase_height"]
    assert compact["node_height"] < full["node_height"]
    assert compact["chart_height"] < full["chart_height"]


def test_dashboard_line_chart_renders_sampled_series_without_gui():
    class FakeCanvas:
        def __init__(self):
            self.operations = []

        def delete(self, *args):
            self.operations.append(("delete", args, {}))

        @staticmethod
        def winfo_width():
            return 500

        @staticmethod
        def winfo_height():
            return 160

        def create_line(self, *args, **kwargs):
            self.operations.append(("line", args, kwargs))
            return len(self.operations)

        def create_text(self, *args, **kwargs):
            self.operations.append(("text", args, kwargs))
            return len(self.operations)

        def create_oval(self, *args, **kwargs):
            self.operations.append(("oval", args, kwargs))
            return len(self.operations)

        @staticmethod
        def bbox(_item):
            return (0, 0, 120, 12)

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.initialize_history_state()
    dashboard.draw_trend_charts = lambda: None
    dashboard.record_history_sample(dashboard_sample(60, 1.0, vehicles=8))
    dashboard.record_history_sample(dashboard_sample(120, 2.0, vehicles=12))
    canvas_widget = FakeCanvas()

    dashboard.draw_line_chart(
        canvas_widget,
        (("vehicles", "Vehicles", "#123456"),),
    )

    assert any(
        operation == "line" and kwargs.get("smooth") is True
        for operation, _args, kwargs in canvas_widget.operations
    )
    assert any(
        operation == "text" and "Vehicles: 12.0" in kwargs.get("text", "")
        for operation, _args, kwargs in canvas_widget.operations
    )


def test_dashboard_phase_plan_and_marker_wrap_match_controller_cycle():
    timing = {
        "green_frames": 240,
        "yellow_frames": 60,
        "all_red_frames": 60,
    }
    cycle, segments = TelemetryDashboard.build_nominal_phase_segments(timing)

    assert cycle == 720
    assert segments["EW"][:2] == (
        (0, 240, "GREEN"),
        (240, 300, "YELLOW"),
    )
    assert segments["NS_A"][1:3] == (
        (360, 600, "GREEN"),
        (600, 660, "YELLOW"),
    )
    assert segments["NS_A"] == segments["NS_B"]
    assert TelemetryDashboard.phase_marker_fraction(0, cycle) == 0
    assert TelemetryDashboard.phase_marker_fraction(360, cycle) == 0.5
    assert TelemetryDashboard.phase_marker_fraction(720, cycle) == 0


def test_dashboard_phase_diagram_renders_live_priority_state_without_gui():
    class FakeLabel:
        def config(self, **kwargs):
            self.values = kwargs

    class FakeCanvas:
        def __init__(self):
            self.operations = []

        def _record(self, operation, args, kwargs):
            self.operations.append((operation, args, kwargs))
            return len(self.operations)

        def delete(self, *args):
            return self._record("delete", args, {})

        @staticmethod
        def winfo_width():
            return 900

        @staticmethod
        def winfo_height():
            return 145

        def create_text(self, *args, **kwargs):
            return self._record("text", args, kwargs)

        def create_rectangle(self, *args, **kwargs):
            return self._record("rectangle", args, kwargs)

        def create_line(self, *args, **kwargs):
            return self._record("line", args, kwargs)

        def create_polygon(self, *args, **kwargs):
            return self._record("polygon", args, kwargs)

        def create_oval(self, *args, **kwargs):
            return self._record("oval", args, kwargs)

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.phase_cycle_canvas = FakeCanvas()
    dashboard.phase_status_lbl = FakeLabel()
    data = dashboard_sample(180, 3.0)
    data["signal_state"] = {
        "timing": {
            "frames_per_second": 60,
            "green_frames": 240,
            "yellow_frames": 60,
            "all_red_frames": 60,
        },
        "nodes": {
            str(NODE_A): {
                "priority_state": "PRIORITY_ACTIVE",
                "signals": {"EB": "GREEN", "WB": "RED", "NB": "RED", "SB": "RED"},
            },
            str(NODE_B): {
                "priority_state": "NORMAL",
                "signals": {"EB": "GREEN", "WB": "GREEN", "NB": "RED", "SB": "RED"},
            },
        },
    }

    dashboard.draw_phase_cycle(data)

    assert sum(
        operation == "oval"
        for operation, _args, _kwargs in dashboard.phase_cycle_canvas.operations
    ) == 3
    assert any(
        operation == "polygon"
        for operation, _args, _kwargs in dashboard.phase_cycle_canvas.operations
    )
    assert dashboard.phase_status_lbl.values == {
        "text": "PRIORITY OVERRIDE",
        "fg": "#2D8CFF",
    }
    assert not any(
        operation == "text"
        and kwargs.get("text") in ("PRIORITY OVERRIDE", "NORMAL PLAN", "ALL RED ACTIVE")
        for operation, _args, kwargs in dashboard.phase_cycle_canvas.operations
    )


def test_dashboard_poll_reschedules_after_unexpected_error():
    class FakeRoot:
        def __init__(self):
            self.after_calls = []

        def after(self, delay, callback):
            self.after_calls.append((delay, callback))

    class FakeLabel:
        def config(self, **kwargs):
            self.last = kwargs

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.root = FakeRoot()
    dashboard.status_lbl = FakeLabel()
    dashboard.last_read_error = None
    dashboard.safe_read_telemetry = lambda: (_ for _ in ()).throw(RuntimeError("boom"))
    dashboard.poll_telemetry()
    assert dashboard.root.after_calls == [(250, dashboard.poll_telemetry)]
    assert dashboard.status_lbl.last["text"] == "TELEMETRY ERROR"


def test_gpu_poller_reads_nvidia_smi_and_degrades_gracefully(monkeypatch):
    calls = []

    class Result:
        returncode = 0
        stdout = "6144, 8192, 73, 187.5, 62\n"

    def fake_run(command, **kwargs):
        calls.append((command, kwargs))
        return Result()

    monkeypatch.setattr(telemetry_dashboard_module.subprocess, "run", fake_run)

    assert telemetry_dashboard_module.poll_gpu_stats() == {
        "vram_used_mb": 6144.0,
        "vram_total_mb": 8192.0,
        "gpu_util_pct": 73.0,
        "power_w": 187.5,
        "temp_c": 62.0,
    }
    assert calls[0][0][0] == "nvidia-smi"
    assert calls[0][1]["timeout"] == 2

    monkeypatch.setattr(
        telemetry_dashboard_module.subprocess,
        "run",
        lambda *args, **kwargs: (_ for _ in ()).throw(FileNotFoundError()),
    )
    assert telemetry_dashboard_module.poll_gpu_stats() == _gpu_none()


def test_llm_turn_reader_skips_partial_lines_and_deduplicates_poll(
    tmp_path, monkeypatch
):
    turn_log = tmp_path / "agent_turn_log.jsonl"
    record = {
        "turn": 1,
        "model": "test-model",
        "status": "OK",
        "flags": {},
    }
    turn_log.write_text(
        json.dumps(record) + "\nnot-json\n{\"turn\":",
        encoding="utf-8",
    )
    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.initialize_llm_monitor_state()

    records, size = dashboard.safe_read_agent_turns(turn_log)

    assert records == [record]
    assert size == turn_log.stat().st_size

    class FakeRoot:
        def __init__(self):
            self.after_calls = []

        def after(self, delay, callback):
            self.after_calls.append((delay, callback))

    dashboard.root = FakeRoot()
    dashboard.safe_read_ai_control = lambda: {
        "armed": True,
        "model": "test-model",
    }
    dashboard.safe_read_agent_turns = lambda: ([record], size)
    dashboard.update_llm_performance_display = lambda: None
    monkeypatch.setattr(
        telemetry_dashboard_module, "poll_gpu_stats", lambda: _gpu_none()
    )
    dashboard.poll_llm_performance()
    dashboard.poll_llm_performance()

    assert len(dashboard.llm_samples) == 1
    assert dashboard.llm_seen_turns == {1}
    assert dashboard.root.after_calls == [
        (
            telemetry_dashboard_module.LLM_POLL_MILLISECONDS,
            dashboard.poll_llm_performance,
        ),
        (
            telemetry_dashboard_module.LLM_POLL_MILLISECONDS,
            dashboard.poll_llm_performance,
        ),
    ]


def test_llm_sample_rollup_reconciles_status_tokens_and_gpu():
    flags_on = {
        "R1": {"tsp": True, "dbl": False},
        "R2": {"tsp": False, "dbl": True},
    }
    gpu = {
        "vram_used_mb": 6000,
        "vram_total_mb": 8192,
        "gpu_util_pct": 80,
        "power_w": 180,
        "temp_c": 60,
    }
    first = TelemetryDashboard.build_llm_sample(
        {
            "turn": 1,
            "model": "model-a",
            "status": "OK",
            "latency_ms": 1000,
            "output_tokens": 20,
            "tokens_per_sec": 10,
            "flags": flags_on,
        },
        gpu,
    )
    second = TelemetryDashboard.build_llm_sample(
        {
            "turn": 2,
            "model": "model-a",
            "status": "HELD_ALL_OFF",
            "latency_ms": 3000,
            "output_tokens": None,
            "tokens_per_sec": None,
            "flags": {},
        },
        {**gpu, "vram_used_mb": 7000, "power_w": 200},
    )

    summary = TelemetryDashboard.summarize_llm_samples([first, second])

    assert first["tsp_on_count"] == 1
    assert first["dbl_on_count"] == 1
    assert summary["turns"] == 2
    assert summary["ok_rate_pct"] == 50.0
    assert summary["held_rate_pct"] == 50.0
    assert summary["avg_latency_ms"] == 2000.0
    assert summary["avg_tokens_per_sec"] == 10.0
    assert summary["avg_vram_used_mb"] == 6500.0
    assert summary["peak_vram_used_mb"] == 7000.0
    assert summary["avg_power_w"] == 190.0
    assert summary["total_output_tokens"] == 20.0


def test_llm_excel_export_writes_samples_and_per_model_summary(tmp_path):
    from openpyxl import load_workbook

    class FakeLabel:
        def config(self, **kwargs):
            self.values = kwargs

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.llm_export_status_lbl = FakeLabel()
    dashboard.llm_samples = [
        {
            "turn": 1,
            "timestamp": 100.0,
            "model": "model-a",
            "status": "OK",
            "latency_ms": 1000.0,
            "input_tokens": 100,
            "output_tokens": 20,
            "tokens_per_sec": 10.0,
            "tsp_on_count": 1,
            "dbl_on_count": 0,
            "reason": "Serve the approaching bus.",
            "vram_used_mb": 6000.0,
            "gpu_util_pct": 80.0,
            "power_w": 180.0,
            "temp_c": 60.0,
            "pax_per_min_recent": 400.0,
        },
        {
            "turn": 2,
            "timestamp": 105.0,
            "model": "model-a",
            "status": "HELD_ALL_OFF",
            "latency_ms": 45000.0,
            "input_tokens": None,
            "output_tokens": None,
            "tokens_per_sec": None,
            "tsp_on_count": 0,
            "dbl_on_count": 0,
            "reason": "",
            "vram_used_mb": 7000.0,
            "gpu_util_pct": 95.0,
            "power_w": 210.0,
            "temp_c": 65.0,
            "pax_per_min_recent": 350.0,
        },
    ]
    destination = tmp_path / "llm_perf_test.xlsx"

    assert dashboard.export_llm_performance(destination) == destination

    workbook = load_workbook(destination, data_only=True)
    try:
        assert workbook.sheetnames == ["Samples", "Summary", "AI Decision Audit"]
        samples = workbook["Samples"]
        summary = workbook["Summary"]
        sample_headers = [cell.value for cell in samples[1]]
        summary_headers = [cell.value for cell in summary[1]]
        assert samples.max_row == 3
        assert sample_headers[:4] == ["turn", "timestamp", "model", "status"]
        assert sample_headers[-1] == "pax_per_min_recent"
        assert summary.max_row == 2
        summary_row = {
            header: summary.cell(row=2, column=index + 1).value
            for index, header in enumerate(summary_headers)
        }
        assert summary_row["model"] == "model-a"
        assert summary_row["turns"] == 2
        assert summary_row["guard_ok_pct"] == 50.0
        assert summary_row["held_pct"] == 50.0
        assert summary_row["avg_latency_ms"] == 23000.0
        assert summary_row["total_output_tokens"] == 20.0
        audit = workbook["AI Decision Audit"]
        assert [cell.value for cell in audit[1]] == main.AI_DECISION_AUDIT_HEADERS
        assert audit.max_row == 3
    finally:
        workbook.close()
    assert dashboard.llm_export_status_lbl.values["fg"] == "#2ECC71"

    dashboard.llm_samples = []
    empty_destination = tmp_path / "must_not_exist.xlsx"
    assert dashboard.export_llm_performance(empty_destination) is None
    assert not empty_destination.exists()
    assert dashboard.llm_export_status_lbl.values["text"] == "No samples yet"


def test_export_all_creates_core_sheets(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    class FakeLabel:
        def config(self, **kwargs):
            self.values = kwargs

    route_id = telemetry_dashboard_module.ROUTE_ORDER[0]
    decision_log = tmp_path / "agent_turn_log.jsonl"
    telemetry_log = tmp_path / "telemetry_log.jsonl"
    decision_log.write_text(
        json.dumps(
            {
                "turn": 3,
                "timestamp": 100.0,
                "model": "model-a",
                "status": "OK",
                "reason": "Serve route one.",
                "pax_per_min_recent": 180.0,
                "flags": {route_id: {"tsp": True, "dbl": False}},
                "locked_routes": [route_id],
                "minimap": "ROUTES:\n1) route one",
                "raw_output": '{"tsp":[true]}',
                "telemetry_snapshot": {
                    "timestamp": 99.5,
                    "frame_number": 118,
                    "simulation_time_seconds": 1.967,
                    "routes": {
                        route_id: {
                            "tsp_enabled": True,
                            "dbl_enabled": False,
                            "nearest_bus_priority_pending": True,
                            "dbl_lane_queue_ahead": 0,
                            "dbl_lane_obstructed": False,
                            "buses_on_route": 2,
                            "route_passengers_total": 90,
                        }
                    },
                    "signal_state": {"nodes": {}},
                    "network_summary": {"total_vehicles": 1},
                    "network_throughput": {
                        "passengers_per_minute_recent": 49.0
                    },
                    "vehicle_positions": [
                        {
                            "snapshot_id": "BUS-1",
                            "vehicle_type": "bus",
                            "route_id": route_id,
                            "x_px": 250.0,
                            "y_px": 245.0,
                        }
                    ],
                },
            }
        )
        + "\nnot-json\n",
        encoding="utf-8",
    )
    telemetry_log.write_text(
        json.dumps(
            {
                "frame": 120,
                "sim_time_s": 2.0,
                "passengers_served_total": 49,
                "passengers_served_bus": 45,
                "passengers_served_car": 4,
                "buses_served": 1,
                "cars_served": 1,
                "pax_per_min_cumulative": 1470.0,
                "pax_per_min_recent": 49.0,
                "vehicles_in_network": 8,
                "ai_armed": True,
                "ai_last_status": "OK",
                "queues_vehicles": {"EB": 2},
                "queues_passengers_est": {"EB": 8},
            }
        )
        + "\n{\"partial\":",
        encoding="utf-8",
    )
    monkeypatch.setattr(
        telemetry_dashboard_module, "AGENT_TURN_LOG_FILE", decision_log
    )
    monkeypatch.setattr(
        telemetry_dashboard_module, "TELEMETRY_LOG_FILE", telemetry_log
    )

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.export_all_status_lbl = FakeLabel()
    dashboard.llm_samples = [
        {
            "turn": 3,
            "timestamp": 100.0,
            "model": "model-a",
            "status": "OK",
            "latency_ms": 800.0,
            "input_tokens": 120,
            "output_tokens": 24,
            "tokens_per_sec": 30.0,
            "tsp_on_count": 1,
            "dbl_on_count": 0,
            "reason": "Serve route one.",
            "vram_used_mb": 5000.0,
            "gpu_util_pct": 70.0,
            "power_w": 150.0,
            "temp_c": 58.0,
            "pax_per_min_recent": 49.0,
        }
    ]
    destination = tmp_path / "combined.xlsx"

    assert dashboard.export_all(destination) == destination

    workbook = load_workbook(destination, data_only=True)
    try:
        # EXPORT ALL is now the only export, so it always carries the four
        # session sheets plus the operator inputs that produced them.
        session_sheets = [
            "Decisions",
            "Telemetry",
            "LLM Performance",
            "LLM Summary",
        ]
        assert workbook.sheetnames == session_sheets + [
            "AI Decision Audit",
            "Control Panel Inputs",
            "Control Panel Inputs (start)",
            "Bus Events",
            "Unit Conversions",
        ]
        assert all(workbook[name].max_row == 2 for name in session_sheets)
        inputs = workbook["Control Panel Inputs"]
        assert [cell.value for cell in inputs[1]] == [
            "section",
            "parameter",
            "value",
        ]
        assert inputs.max_row > 1
        decision_headers = [cell.value for cell in workbook["Decisions"][1]]
        assert workbook["Decisions"].cell(
            row=2,
            column=decision_headers.index(f"{route_id}_tsp") + 1,
        ).value is True
        telemetry_headers = [cell.value for cell in workbook["Telemetry"][1]]
        assert workbook["Telemetry"].cell(
            row=2,
            column=telemetry_headers.index("ai_armed") + 1,
        ).value is True
        assert workbook["LLM Performance"]["C2"].value == "model-a"
        assert workbook["LLM Summary"]["A2"].value == "model-a"
        audit_row = dict(
            zip(
                [cell.value for cell in workbook["AI Decision Audit"][1]],
                [cell.value for cell in workbook["AI Decision Audit"][2]],
            )
        )
        assert audit_row["model_reason"] == "Serve route one."
        assert audit_row["requested_tsp_routes"] == route_id
        assert audit_row["observation_minimap"] == "ROUTES:\n1) route one"
        assert audit_row["observed_tsp_routes"] == route_id
        assert audit_row["pending_priority_routes"] == route_id
        # Per-route load at the decision instant, flattened for reading the
        # decision against the state it was made from.
        assert audit_row["route_vehicles_json"] == f'{{"{route_id}":2}}'
        assert audit_row["route_passengers_json"] == f'{{"{route_id}":90}}'
        assert audit_row["vehicle_count"] == 1
        assert '"snapshot_id":"BUS-1"' in audit_row["vehicle_positions_json"]
    finally:
        workbook.close()
    assert dashboard.export_all_status_lbl.values["fg"] == "#2ECC71"


def test_export_all_missing_logs_ok(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    class FakeLabel:
        def config(self, **kwargs):
            self.values = kwargs

    monkeypatch.setattr(
        telemetry_dashboard_module,
        "AGENT_TURN_LOG_FILE",
        tmp_path / "missing_decisions.jsonl",
    )
    monkeypatch.setattr(
        telemetry_dashboard_module,
        "TELEMETRY_LOG_FILE",
        tmp_path / "missing_telemetry.jsonl",
    )
    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.export_all_status_lbl = FakeLabel()
    dashboard.llm_samples = [
        {"turn": 1, "model": "model-a", "status": "OK"}
    ]
    destination = tmp_path / "missing_logs.xlsx"

    assert dashboard.export_all(destination) == destination

    workbook = load_workbook(destination, data_only=True)
    try:
        assert workbook["Decisions"].max_row == 1
        assert workbook["Telemetry"].max_row == 1
        assert workbook["LLM Performance"].max_row == 2
        assert workbook["LLM Summary"].max_row == 2
    finally:
        workbook.close()


def test_export_all_empty_writes_nothing(tmp_path, monkeypatch):
    class FakeLabel:
        def config(self, **kwargs):
            self.values = kwargs

    monkeypatch.setattr(
        telemetry_dashboard_module,
        "AGENT_TURN_LOG_FILE",
        tmp_path / "missing_decisions.jsonl",
    )
    monkeypatch.setattr(
        telemetry_dashboard_module,
        "TELEMETRY_LOG_FILE",
        tmp_path / "missing_telemetry.jsonl",
    )
    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.export_all_status_lbl = FakeLabel()
    dashboard.llm_samples = []
    destination = tmp_path / "must_not_exist.xlsx"

    assert dashboard.export_all(destination) is None
    assert not destination.exists()
    assert dashboard.export_all_status_lbl.values == {
        "text": "Nothing to export yet",
        "fg": telemetry_dashboard_module.COLOR_WARNING,
    }


def test_combined_decisions_matches_session_columns(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    class FakeLabel:
        def config(self, **kwargs):
            self.values = kwargs

    decision_log = tmp_path / "agent_turn_log.jsonl"
    decision_log.write_text(
        json.dumps(
            {
                "turn": 1,
                "model": "model-a",
                "status": "OK",
                "flags": {},
            }
        )
        + "\n",
        encoding="utf-8",
    )
    missing_telemetry = tmp_path / "missing_telemetry.jsonl"
    monkeypatch.setattr(
        telemetry_dashboard_module, "AGENT_TURN_LOG_FILE", decision_log
    )
    monkeypatch.setattr(
        telemetry_dashboard_module, "TELEMETRY_LOG_FILE", missing_telemetry
    )
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", decision_log)
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", missing_telemetry)

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.export_all_status_lbl = FakeLabel()
    dashboard.llm_samples = []
    combined_destination = tmp_path / "combined.xlsx"
    session_destination = tmp_path / "session.xlsx"

    assert dashboard.export_all(combined_destination) == combined_destination
    assert main.export_session_excel(session_destination) == session_destination

    combined = load_workbook(combined_destination, data_only=True)
    session = load_workbook(session_destination, data_only=True)
    try:
        combined_headers = [cell.value for cell in combined["Decisions"][1]]
        session_headers = [cell.value for cell in session["Decisions"][1]]
        assert combined_headers == session_headers
        assert combined_headers == telemetry_dashboard_module.DECISION_EXPORT_HEADERS
    finally:
        combined.close()
        session.close()


def test_summary_snapshot_export_reconciles_live_telemetry(tmp_path):
    from openpyxl import load_workbook

    class FakeLabel:
        def config(self, **kwargs):
            self.values = kwargs

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.summary_export_status_lbl = FakeLabel()
    dashboard.latest_telemetry = {
        "timestamp": time.time(),
        "frame_number": 120,
        "simulation_time_seconds": 2.0,
        "simulation_paused": False,
        "simulation_speed": 1.0,
        "network_summary": {
            "total_vehicles": 10,
            "total_buses": 2,
            "passenger_volume": 122,
            "queues": {"EB": 3},
            "queues_passengers_est": {"EB": 12},
            "pending_demand": 1,
        },
        "network_throughput": {
            "passengers_served_total": 90,
            "passengers_per_minute_recent": 45.0,
        },
        "network_discharge": {"active": False, "status": "IDLE"},
        "demand_generation": {
            "EB": {
                "active": True,
                "model": "Poisson",
                "configured_rate_vpm": 20,
                "effective_rate_vpm": 20.0,
                "pending_arrivals": 1,
                "peak_active": False,
            }
        },
        "routes": {
            "R1_EB_A_NB": {
                "route_name": "EB to Node A",
                "active": True,
                "tsp_enabled": True,
                "dbl_enabled": False,
                "buses_on_route": 2,
            }
        },
        "signal_state": {
            "nodes": {
                str(NODE_A): {
                    "node_x": NODE_A,
                    "phase": "EW_GREEN",
                    "phase_index": 0,
                    "phase_timer_frames": 20,
                    "priority_state": "NORMAL",
                    "priority_timer_frames": 0,
                    "signals": {
                        "EB": "GREEN",
                        "WB": "GREEN",
                        "NB": "RED",
                        "SB": "RED",
                    },
                    "active_request": None,
                    "queued_requests": [],
                    "reservation_count": 1,
                }
            }
        },
    }
    destination = tmp_path / "summary.xlsx"

    assert dashboard.export_summary_snapshot(destination) == destination

    workbook = load_workbook(destination, data_only=True)
    try:
        assert workbook.sheetnames == [
            "Overview",
            "Queues & Demand",
            "Routes",
            "Signal Nodes",
            "AI Decision Audit",
        ]
        overview = {
            row[0]: row[1]
            for row in workbook["Overview"].iter_rows(
                min_row=2, values_only=True
            )
        }
        assert overview["network_summary.total_vehicles"] == 10
        assert overview["network_throughput.passengers_served_total"] == 90
        assert workbook["Queues & Demand"]["A2"].value == "EB"
        assert workbook["Routes"]["A2"].value == "R1_EB_A_NB"
        assert workbook["Signal Nodes"]["A2"].value == str(NODE_A)
        audit = workbook["AI Decision Audit"]
        audit_row = dict(
            zip([cell.value for cell in audit[1]], [cell.value for cell in audit[2]])
        )
        assert audit_row["observed_tsp_routes"] == "R1_EB_A_NB"
        assert audit_row["simulation_time_s"] == 2.0
    finally:
        workbook.close()
    assert dashboard.summary_export_status_lbl.values["fg"] == "#2ECC71"

    dashboard.latest_telemetry = None
    empty_destination = tmp_path / "no_summary.xlsx"
    assert dashboard.export_summary_snapshot(empty_destination) is None
    assert not empty_destination.exists()


def test_session_trends_export_uses_only_in_memory_history(tmp_path):
    from openpyxl import load_workbook

    class FakeLabel:
        def config(self, **kwargs):
            self.values = kwargs

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.initialize_history_state()
    dashboard.draw_trend_charts = lambda: None
    dashboard.trends_export_status_lbl = FakeLabel()
    dashboard.record_history_sample(
        dashboard_sample(60, 1.0, vehicles=8, buses=1, queued=2, pending=3)
    )
    dashboard.record_history_sample(
        dashboard_sample(120, 2.0, vehicles=12, buses=2, queued=6, pending=1)
    )
    destination = tmp_path / "trends.xlsx"

    assert dashboard.export_session_trends(destination) == destination

    workbook = load_workbook(destination, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Session Trends",
            "Session Charts",
            "Summary",
            "AI Decision Audit",
        ]
        trends = workbook["Session Trends"]
        assert trends.max_row == 3
        assert [cell.value for cell in trends[1]] == [
            "time_seconds",
            "vehicles",
            "buses",
            "road_queue",
            "pending_demand",
            "congestion_pct",
        ]
        assert "SessionTrendData" in trends.tables
        assert trends.tables["SessionTrendData"].ref == "A1:F3"
        assert trends.auto_filter.ref is None
        summary = {
            row[0]: row[1]
            for row in workbook["Summary"].iter_rows(
                min_row=2, values_only=True
            )
        }
        assert summary["samples"] == "=COUNT('Session Trends'!A2:A3)"
        assert summary["duration_seconds"] == "=B4-B3"
        assert summary["average_vehicles"] == (
            "=AVERAGE('Session Trends'!B2:B3)"
        )
        assert summary["peak_road_queue"] == "=MAX('Session Trends'!D2:D3)"
        assert workbook.calculation.fullCalcOnLoad is True
        assert workbook.calculation.forceFullCalc is True

        charts = workbook["Session Charts"]._charts
        assert len(charts) == 3
        assert [len(chart.series) for chart in charts] == [2, 2, 1]
        assert [chart.anchor._from.row for chart in charts] == [0, 22, 44]
        assert not workbook["Session Charts"]._images
        assert [chart.title.tx.rich.p[0].r[0].t for chart in charts] == [
            "Network Occupancy Over Time (vehicles)",
            "Queue Pressure Over Time (vehicles and arrivals)",
            "Network Congestion Over Time (%)",
        ]
        assert all(chart.x_axis.title is None for chart in charts)
        assert all(chart.x_axis.tickLblPos == "low" for chart in charts)
        assert all(chart.x_axis.delete is False for chart in charts)
        assert all(chart.y_axis.delete is False for chart in charts)
        assert all(chart.title.overlay is False for chart in charts)
        assert [
            series.tx.v
            for chart in charts
            for series in chart.series
        ] == [
            "Vehicles",
            "Buses",
            "Road queue",
            "Pending demand",
            "Congestion",
        ]
        assert [
            series.yVal.numRef.f
            for chart in charts
            for series in chart.series
        ] == [
            "'Session Trends'!$B$2:$B$3",
            "'Session Trends'!$C$2:$C$3",
            "'Session Trends'!$D$2:$D$3",
            "'Session Trends'!$E$2:$E$3",
            "'Session Trends'!$F$2:$F$3",
        ]
    finally:
        workbook.close()
    assert dashboard.trends_export_status_lbl.values["fg"] == "#2ECC71"

    dashboard.clear_history()
    empty_destination = tmp_path / "no_trends.xlsx"
    assert dashboard.export_session_trends(empty_destination) is None
    assert not empty_destination.exists()


def test_excel_export_filename_uses_requested_field_order():
    assert build_excel_export_filename(
        "gemini2.5", 300, 42, timestamp="20260914_112746"
    ) == "gemini2.5-assisted_5min_42seed_14092026_112746.xlsx"


def test_runtime_paths_are_repo_root_relative(monkeypatch):
    monkeypatch.undo()  # conftest's isolate_runtime_files: check the real defaults
    repo_root = main.BASE_DIR
    # DASHBOARD_PATH no longer exists: the telemetry dashboard mounts
    # in-process (TelemetryDashboard imported directly) instead of being
    # launched as a separate subprocess script.
    assert not hasattr(main, "DASHBOARD_PATH")
    assert main.TELEMETRY_PATH == repo_root / "data" / "traffic_state_telemetry.json"
    assert DEFAULT_TELEMETRY_PATH == repo_root / "data" / "traffic_state_telemetry.json"


def test_main_uses_monotonic_elapsed_time_with_bounded_catchup():
    source = Path(main.__file__).read_text(encoding="utf-8")
    assert "time.monotonic()" in source
    assert "max_catchup_seconds" in source
    assert "time_accumulator += elapsed * sim_speed" in source
    assert "max_steps_per_callback = 6" in source
    assert "steps_this_callback < max_steps_per_callback" in source


def test_control_panel_uses_normal_desktop_window_stacking():
    source = Path(control_panel.__file__).read_text(encoding="utf-8")
    tree = ast.parse(source)
    topmost_calls = [
        node
        for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Attribute)
        and node.func.attr == "attributes"
        and len(node.args) >= 2
        and isinstance(node.args[0], ast.Constant)
        and node.args[0].value == "-topmost"
    ]

    assert len(topmost_calls) == 1
    assert isinstance(topmost_calls[0].args[1], ast.Constant)
    assert topmost_calls[0].args[1].value is False
    assert "overrideredirect" not in source


def test_control_panel_exposes_congestion_peak_model():
    source = Path(control_panel.__file__).read_text(encoding="utf-8")
    assert main.CONGESTION_MODEL in source


def test_congestion_backlog_survives_blocked_spawn_and_drains_on_admission(
    monkeypatch,
):
    main.reset_all_spawner_states()
    rng = main.spawner_states["EB"]["rng"]
    monkeypatch.setattr(rng, "random", lambda: 0.0)
    monkeypatch.setattr(rng, "choice", lambda values: values[0])
    monkeypatch.setattr(rng, "uniform", lambda low, high: low)
    lanes = lane_options()
    config = {
        "model": main.CONGESTION_MODEL,
        "rate": 30,
        "turn_split": 1.0,
        "heavy_ratio": 0.0,
    }
    blocker = Vehicle(-20, lanes["EB"][0], "EB", lane_index=0)
    vehicles = [blocker]

    main.try_spawn_vehicle(
        vehicles, "EB", "EB", -20, lanes["EB"], config
    )
    state = main.spawner_states["EB"]
    assert state["pending_arrivals"] == 1
    assert state["admitted_arrivals"] == 0
    assert len(vehicles) == 1

    vehicles.clear()
    main.try_spawn_vehicle(
        vehicles, "EB", "EB", -20, lanes["EB"], config
    )
    assert state["pending_arrivals"] == 1
    assert state["requested_arrivals"] == 2
    assert state["admitted_arrivals"] == 1
    assert len(vehicles) == 1
    main.reset_all_spawner_states()


def test_active_dbl_retains_new_general_left_turn_at_source(monkeypatch):
    config = control_panel.bus_routes_config["R1_EB_A_NB"]
    monkeypatch.setitem(config, "dbl_enabled", True)
    bus = make_bus_for_leg("R1_EB_A_NB", NODE_A, "SPAWN_GUARD_BUS")
    controller = SignalController({"green_time": 999})
    controller.update([bus])
    assert controller.is_dbl_active_for_approach(NODE_A, "EB")

    main.reset_all_spawner_states()
    config = {"model": "Poisson", "rate": 60, "turn_split": 0.0, "heavy_ratio": 0.0}
    state = main.spawner_states["EB"]
    monkeypatch.setattr(state["rng"], "random", lambda: 0.5)  # a near-node left turn
    main._record_demand_offer("EB", state, 1, config)
    monkeypatch.setattr(main, "should_spawn_vehicle", lambda *args: True)
    spawned = []
    main.try_spawn_vehicle(
        spawned, "EB", "EB", -20, lane_options()["EB"], config,
        signal_controller=controller,
    )

    assert spawned == []
    assert state["pending_arrivals"] == 1  # retained at the source, not lost
    main.reset_all_spawner_states()


def test_congestion_peak_cycles_to_recovery_and_caps_backlog(monkeypatch):
    main.reset_all_spawner_states()
    state = main.spawner_states["EB"]
    monkeypatch.setattr(state["rng"], "random", lambda: 1.0)
    assert not main.should_spawn_vehicle("EB", main.CONGESTION_MODEL, 30)
    assert state["peak_active"] is True
    assert state["effective_rate_vpm"] == 120.0

    state["congestion_cycle_frame"] = main.CONGESTION_PEAK_SECONDS * 60
    assert not main.should_spawn_vehicle("EB", main.CONGESTION_MODEL, 30)
    assert state["peak_active"] is False
    assert state["effective_rate_vpm"] == 30.0

    state["congestion_cycle_frame"] = 0
    state["pending_arrivals"] = main.MAX_PENDING_ARRIVALS
    monkeypatch.setattr(state["rng"], "random", lambda: 0.0)
    assert not main.should_spawn_vehicle("EB", main.CONGESTION_MODEL, 30)
    assert state["pending_arrivals"] == main.MAX_PENDING_ARRIVALS
    assert state["overflow_arrivals"] == 1
    assert state["requested_arrivals"] == 1  # still counted as offered demand
    main.reset_all_spawner_states()


def test_llm_callbacks_write_runtime_control_instead_of_remaining_placeholders():
    source = Path(control_panel.__file__).read_text(encoding="utf-8")
    tree = ast.parse(source)
    functions = {node.name: node for node in ast.walk(tree) if isinstance(node, ast.FunctionDef)}
    for callback_name in (
        "on_run_llm",
        "on_strategy_selected",
        "on_decider_selected",
        "on_tick_seconds_changed",
    ):
        callback = functions[callback_name]
        assert not any(isinstance(node, ast.Pass) for node in ast.walk(callback))
        called_functions = {
            node.func.id
            for node in ast.walk(callback)
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)
        }
        assert "write_ai_control" in called_functions


def test_priority_telemetry_distinguishes_pending_from_active():
    """A request queued behind another bus is pending for TSP, but its DBL is
    in force from the frame it is raised: a lane reservation on the bus's own
    approach never waits for the node's one TSP slot."""
    control_panel.bus_routes_config["R1_EB_A_NB"]["dbl_enabled"] = True
    first = make_bus_for_leg("R1_EB_A_NB", NODE_A, "TELEMETRY_BUS_A")
    second = make_bus_for_leg("R1_EB_A_NB", NODE_A, "TELEMETRY_BUS_B")
    second.x -= 60
    controller = SignalController({"green_time": 100}, 2, 2)
    exporter = TelemetryExporter(export_interval_frames=1)

    controller.update([first, second])
    payload = exporter.build_payload(controller, [first, second], 1)
    by_id = {bus["bus_id"]: bus for bus in payload["active_buses"]}
    pending = by_id["TELEMETRY_BUS_B"]
    assert pending["priority_transitioning"] is True
    assert pending["priority_requested"] is True
    assert pending["dbl_priority_pending"] is False
    assert pending["priority_granted"] is False
    assert pending["dbl_active_triggered"] is True

    active = by_id["TELEMETRY_BUS_A"]
    assert active["priority_transitioning"] is False
    assert active["priority_requested"] is True
    assert active["dbl_priority_pending"] is False
    assert active["dbl_active_triggered"] is True
    # DBL alone never adjusts the signals, so nothing is "granted".
    assert active["priority_granted"] is False
    assert active["tsp_action"] == "none"
    assert controller.get_all_dbl_states()[NODE_A]["EB"] == "ACTIVE"


def test_dashboard_shows_active_and_pending_priority_separately():
    class FakeLabel:
        def config(self, **kwargs):
            self.values = kwargs

    class FakeRoot:
        @staticmethod
        def update_idletasks():
            pass

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.queue_history = []
    dashboard.root = FakeRoot()
    dashboard.vars = {
        key: FakeLabel()
        for key in (
            "vehicles", "buses", "passengers", "queued", "delay",
            "congestion", "tsp", "dbl", "timer"
        )
    }
    dashboard.node_a_canvas = object()
    dashboard.node_b_canvas = object()
    dashboard.draw_intersection = lambda *args, **kwargs: None
    data = {
        "frame_number": 1,
        "network_summary": {
            "total_vehicles": 2,
            "total_buses": 2,
            "passenger_volume": 90,
            "queues": {},
            "pending_demand": 4,
        },
        "active_buses": [
            {"tsp_active_triggered": True, "dbl_priority_pending": True},
            {"tsp_priority_pending": True, "dbl_active_triggered": True},
        ],
        "signal_state": {"nodes": {}},
    }

    dashboard.update_metrics(data)

    assert dashboard.vars["tsp"].values["text"] == "1/1"
    assert dashboard.vars["dbl"].values["text"] == "1/1"
    assert dashboard.vars["queued"].values["text"] == "0/4"


def test_webster_timing_readout_lives_in_the_dashboard_summary_tab():
    """Webster timing is reported by the dashboard, not the control panel:
    measured capacity plus each node's condition/cycle/demand ratio in the
    timing card, and the green split under each live signal diagram."""

    class FakeLabel:
        def __init__(self):
            self.values = {}
            self.packed = False

        def config(self, **kwargs):
            self.values.update(kwargs)

        def pack(self, **_kwargs):
            self.packed = True

        def pack_forget(self):
            self.packed = False

        def winfo_manager(self):
            return "pack" if self.packed else ""

    class FakeCanvas:
        def __init__(self):
            self.green_time_label = FakeLabel()

        def winfo_width(self):
            return 170

    class FakeCard:
        @staticmethod
        def winfo_width():
            return 320

    def fake_node_card():
        return {
            "card": FakeLabel(), "state": FakeLabel(), "title": FakeLabel(),
            "cycle": FakeLabel(), "ratio": FakeLabel(), "note": FakeLabel(),
        }

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.node_a_canvas = FakeCanvas()
    dashboard.node_b_canvas = FakeCanvas()
    dashboard.webster_status_lbl = FakeLabel()
    dashboard.webster_status_lbl.master = FakeCard()
    dashboard.webster_node_cards = {NODE_A: fake_node_card(), NODE_B: fake_node_card()}

    control_panel.global_config["calibrating"] = False
    control_panel.global_config["measured_saturation_flow"] = 1800.0
    control_panel.global_config["webster_splits"] = {
        NODE_A: {
            "cycle_time_sec": 60.0, "Y": 0.5, "y_ew": 0.3, "y_ns": 0.2,
            "EW_green_sec": 26.5, "NS_green_sec": 18.0, "oversaturated": False,
        },
    }

    dashboard._refresh_webster_timing_labels()

    assert "1,800" in dashboard.webster_status_lbl.values["text"]
    node_a = dashboard.webster_node_cards[NODE_A]
    assert node_a["card"].packed
    assert node_a["state"].values["text"] == "Optimal"
    assert "60 s" in node_a["cycle"].values["text"]
    assert "EW 0.30" in node_a["ratio"].values["text"]
    assert "Total 0.50" in node_a["ratio"].values["text"]
    assert node_a["note"].values.get("text", "") == ""
    # Node B has no calibrated split: the card still shows, reporting the
    # gap explicitly rather than timing that was never measured.
    node_b = dashboard.webster_node_cards[NODE_B]
    assert node_b["card"].packed
    assert node_b["state"].values["text"] == "Unavailable"
    assert "no timing data" in node_b["cycle"].values["text"]

    node_a_green = dashboard.node_a_canvas.green_time_label.values["text"]
    assert "26.5" in node_a_green
    assert "18.0" in node_a_green
    assert dashboard.node_b_canvas.green_time_label.values["text"] == "Green time: --"


def test_oversaturated_node_shows_capped_cycle_and_warning_note():
    class FakeLabel:
        def __init__(self):
            self.values = {}
            self.packed = False

        def config(self, **kwargs):
            self.values.update(kwargs)

        def pack(self, **_kwargs):
            self.packed = True

        def pack_forget(self):
            self.packed = False

        def winfo_manager(self):
            return "pack" if self.packed else ""

    class FakeCanvas:
        def __init__(self):
            self.green_time_label = FakeLabel()

        def winfo_width(self):
            return 170

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.node_a_canvas = FakeCanvas()
    dashboard.node_b_canvas = FakeCanvas()
    dashboard.webster_node_cards = {
        NODE_A: {
            "card": FakeLabel(), "state": FakeLabel(), "title": FakeLabel(),
            "cycle": FakeLabel(), "ratio": FakeLabel(), "note": FakeLabel(),
        }
    }

    control_panel.global_config["calibrating"] = False
    control_panel.global_config["measured_saturation_flow"] = 1500.0
    control_panel.global_config["webster_splits"] = {
        NODE_A: {
            "cycle_time_sec": 120.0, "Y": 1.2, "y_ew": 0.7, "y_ns": 0.5,
            "EW_green_sec": 60.0, "NS_green_sec": 40.0, "oversaturated": True,
        },
    }

    dashboard._refresh_webster_timing_labels()

    refs = dashboard.webster_node_cards[NODE_A]
    assert refs["state"].values["text"] == "Oversaturated"
    assert "capped" in refs["cycle"].values["text"]
    assert refs["note"].packed
    assert "Reduce demand" in refs["note"].values["text"]


def test_dbl_lamp_uses_distinct_pending_active_and_clearing_colors(monkeypatch):
    monkeypatch.setattr(canvas, "_get_dbl_flash_clock", lambda _paused: 0.0)
    surface = pygame.Surface((100, 100))
    x, y = 50, 50
    light_center = canvas.get_signal_light_center(x, y, "EASTBOUND")
    lamp_center = (light_center[0], light_center[1] - 16)
    expected_colors = {
        "INACTIVE": (0, 0, 0),
        "TRANSITIONING": (245, 158, 11),
        "ACTIVE": (0, 255, 120),
        "CLEARING": (0, 200, 220),
    }

    for state, expected in expected_colors.items():
        canvas.draw_dbl_signal(surface, x, y, "EASTBOUND", state, False)
        assert surface.get_at(lamp_center)[:3] == expected


def test_export_all_folds_in_snapshot_and_trends_when_available(
    tmp_path, monkeypatch
):
    """EXPORT ALL must carry the content the removed tab buttons produced.

    Session Trends in particular lives only in dashboard memory, so if the
    fold-in silently failed that data would be unrecoverable.
    """
    from openpyxl import load_workbook

    class FakeLabel:
        def config(self, **kwargs):
            self.values = kwargs

    monkeypatch.setattr(
        telemetry_dashboard_module,
        "AGENT_TURN_LOG_FILE",
        tmp_path / "missing_turns.jsonl",
    )
    monkeypatch.setattr(
        telemetry_dashboard_module,
        "TELEMETRY_LOG_FILE",
        tmp_path / "missing_telemetry.jsonl",
    )

    dashboard = TelemetryDashboard.__new__(TelemetryDashboard)
    dashboard.export_all_status_lbl = FakeLabel()
    dashboard.summary_export_status_lbl = FakeLabel()
    dashboard.trends_export_status_lbl = FakeLabel()
    dashboard.llm_samples = [
        {"turn": 1, "model": "model-a", "status": "OK", "latency_ms": 10.0}
    ]

    # Live snapshot, as the Summary tab would hold.
    dashboard.latest_telemetry = {
        "timestamp": time.time(),
        "frame_number": 120,
        "simulation_time_seconds": 2.0,
        "simulation_paused": False,
        "simulation_speed": 1.0,
        "network_summary": {"passenger_volume": 49},
        "network_throughput": {"passengers_per_minute": 100.0},
        "routes": {},
        "signal_state": {"nodes": {}},
    }

    # In-memory trend history, as the Session Trends tab would hold.
    dashboard.initialize_history_state()
    dashboard.draw_trend_charts = lambda: None
    dashboard.record_history_sample(
        dashboard_sample(60, 1.0, vehicles=8, buses=1, queued=2, pending=3)
    )
    dashboard.record_history_sample(
        dashboard_sample(120, 2.0, vehicles=12, buses=2, queued=6, pending=1)
    )

    destination = tmp_path / "combined_full.xlsx"
    assert dashboard.export_all(destination) == destination

    workbook = load_workbook(destination, data_only=True)
    try:
        names = workbook.sheetnames
        # The four session sheets plus the inputs that produced them.
        for expected in (
            "Decisions",
            "Telemetry",
            "LLM Performance",
            "LLM Summary",
            "Control Panel Inputs",
        ):
            assert expected in names, names
        # Folded-in snapshot sheets.
        for expected in (
            "Snapshot Overview",
            "Snapshot Queues & Demand",
            "Snapshot Routes",
            "Snapshot Signal Nodes",
        ):
            assert expected in names, names
        # Folded-in trends sheets, which exist nowhere else.
        assert "Session Trends" in names, names
        assert "Session Charts" in names, names
        assert "Trends Summary" in names, names
        trends = workbook["Session Trends"]
        assert trends.max_row == 3  # header plus the two recorded samples
        assert len(workbook["Session Charts"]._charts) == 3
    finally:
        workbook.close()
