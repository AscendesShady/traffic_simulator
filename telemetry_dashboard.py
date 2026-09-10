"""Read-only live telemetry dashboard with session-only trend history."""

from collections import deque
import json
import math
import os
from pathlib import Path
import subprocess
import time
import tkinter as tk
from tkinter import ttk

from guard import ROUTE_ORDER


COLOR_BG = "#1A1E29"
COLOR_CARD = "#222834"
COLOR_CARD_BORDER = "#333D50"
COLOR_ACCENT = "#2D8CFF"
COLOR_SUCCESS = "#2ECC71"
COLOR_WARNING = "#F59E0B"
COLOR_DANGER = "#EF4444"
COLOR_TEXT_PRIMARY = "#FFFFFF"
COLOR_TEXT_SECONDARY = "#94A3B8"
FONT_FAMILY = "Segoe UI"

BASE_DIR = Path(__file__).resolve().parent
TELEMETRY_FILE = BASE_DIR / "traffic_state_telemetry.json"
AGENT_TURN_LOG_FILE = BASE_DIR / "agent_turn_log.jsonl"
TELEMETRY_LOG_FILE = BASE_DIR / "telemetry_log.jsonl"
AI_CONTROL_FILE = BASE_DIR / "ai_control.json"
DASHBOARD_EXPORT_DIR = BASE_DIR / "excel_exports"
STALE_AFTER_SECONDS = 2.0
HISTORY_SAMPLE_SECONDS = 1.0
HISTORY_MAX_POINTS = 600
LLM_POLL_MILLISECONDS = 1000
WINDOW_DEFAULT_WIDTH = 900
WINDOW_DEFAULT_HEIGHT = 430
WINDOW_MIN_WIDTH = 480
WINDOW_MIN_HEIGHT = 320
WINDOW_INITIAL_MIN_HEIGHT = 360
WINDOW_SCREEN_MARGIN_X = 80
WINDOW_SCREEN_MARGIN_Y = 140
WINDOW_GEOMETRY_ENV = "TRAFFIC_TELEMETRY_GEOMETRY"

DECISION_EXPORT_HEADERS = [
    "turn",
    "timestamp",
    "model",
    "status",
    "reason",
    "pax_per_min_at_turn",
]
for _route_id in ROUTE_ORDER:
    DECISION_EXPORT_HEADERS.extend(
        (f"{_route_id}_tsp", f"{_route_id}_dbl")
    )
DECISION_EXPORT_HEADERS.extend(("locked_routes", "minimap", "raw_output"))

TELEMETRY_EXPORT_HEADERS = [
    "frame",
    "sim_time_s",
    "passengers_served_total",
    "passengers_served_bus",
    "passengers_served_car",
    "buses_served",
    "cars_served",
    "pax_per_min_cumulative",
    "pax_per_min_recent",
    "vehicles_in_network",
    "ai_armed",
    "ai_last_status",
    "queues_vehicles",
    "queues_passengers_est",
]

# Summary KPI strip: label shown after the value, in display order.
SUMMARY_KPIS = (
    ("VEH", "vehicles"),
    ("BUS", "buses"),
    ("PAX", "passengers"),
    ("Q", "queued"),
    ("AVGQ", "delay"),
    ("CONG", "congestion"),
    ("TSP", "tsp"),
    ("DBL", "dbl"),
    ("SIM", "timer"),
)

# Hover text for the deliberately compact KPI strip. Keep this keyed by the
# underlying metric rather than display order so every abbreviation has one
# clear, testable definition.
KPI_TOOLTIPS = {
    "vehicles": "Vehicles currently in the network",
    "buses": "Buses currently in the network",
    "passengers": "Passengers currently carried by network vehicles",
    "queued": "Queued road vehicles / pending arrivals",
    "delay": "Average queued vehicles over recent samples",
    "congestion": "Queued vehicles as a percentage of network traffic",
    "tsp": "Transit Signal Priority: active / pending",
    "dbl": "Dynamic Bus Lane: active / pending",
    "timer": "Elapsed simulation time",
}

LLM_PERFORMANCE_HEADERS = [
    "turn",
    "timestamp",
    "model",
    "status",
    "latency_ms",
    "input_tokens",
    "output_tokens",
    "tokens_per_sec",
    "tsp_on_count",
    "dbl_on_count",
    "reason",
    "vram_used_mb",
    "gpu_util_pct",
    "power_w",
    "temp_c",
    "pax_per_min_recent",
]

LLM_SUMMARY_HEADERS = [
    "model",
    "turns",
    "guard_ok_pct",
    "held_pct",
    "avg_latency_ms",
    "avg_tokens_per_sec",
    "avg_vram_used_mb",
    "peak_vram_used_mb",
    "avg_power_w",
    "total_output_tokens",
]


def read_jsonl(path):
    """Read complete JSON-object lines while tolerating missing/damaged logs."""
    rows = []
    try:
        with Path(path).open("r", encoding="utf-8") as source:
            for line in source:
                try:
                    value = json.loads(line)
                except (json.JSONDecodeError, TypeError, ValueError):
                    continue
                if isinstance(value, dict):
                    rows.append(value)
    except (OSError, UnicodeError):
        pass
    return rows


def _gpu_none():
    return {
        "vram_used_mb": None,
        "vram_total_mb": None,
        "gpu_util_pct": None,
        "power_w": None,
        "temp_c": None,
    }


def poll_gpu_stats():
    """Return current first-GPU statistics, or an all-None safe fallback."""
    try:
        result = subprocess.run(
            [
                "nvidia-smi",
                "--query-gpu=memory.used,memory.total,utilization.gpu,"
                "power.draw,temperature.gpu",
                "--format=csv,noheader,nounits",
            ],
            capture_output=True,
            text=True,
            timeout=2,
        )
        if result.returncode != 0:
            return _gpu_none()
        lines = result.stdout.strip().splitlines()
        if not lines:
            return _gpu_none()
        parts = [part.strip() for part in lines[0].split(",")]
        if len(parts) < 5:
            return _gpu_none()
        return {
            "vram_used_mb": float(parts[0]),
            "vram_total_mb": float(parts[1]),
            "gpu_util_pct": float(parts[2]),
            "power_w": float(parts[3]),
            "temp_c": float(parts[4]),
        }
    except Exception:
        return _gpu_none()


class HoverTooltip:
    """Small delayed tooltip shared by one widget region and its children."""

    def __init__(self, widget, text, delay_ms=350):
        self.widget = widget
        self.text = str(text)
        self.delay_ms = max(0, int(delay_ms))
        self.after_id = None
        self.hide_after_id = None
        self.tip_window = None
        self.targets = []
        self.add_target(widget)

    def add_target(self, widget):
        """Make another widget part of the same uninterrupted hover region."""
        if widget in self.targets:
            return
        self.targets.append(widget)
        widget.bind("<Enter>", self.schedule, add="+")
        widget.bind("<Leave>", self.schedule_hide, add="+")
        widget.bind("<ButtonPress>", self.hide, add="+")
        widget.bind("<Destroy>", self.hide, add="+")

    def schedule(self, _event=None):
        self.cancel_pending_hide()
        self.cancel_pending()
        if self.tip_window is not None:
            return
        self.after_id = self.widget.after(self.delay_ms, self.show)

    def schedule_hide(self, _event=None):
        self.cancel_pending()
        self.cancel_pending_hide()
        # Entering a child produces a leave event on its parent in Tk. A short
        # deferral lets the child's Enter cancel this hide, making the entire
        # KPI card behave as one continuous target without flicker.
        self.hide_after_id = self.widget.after(60, self.hide)

    def cancel_pending(self):
        if self.after_id is None:
            return
        try:
            self.widget.after_cancel(self.after_id)
        except tk.TclError:
            pass
        self.after_id = None

    def cancel_pending_hide(self):
        if self.hide_after_id is None:
            return
        try:
            self.widget.after_cancel(self.hide_after_id)
        except tk.TclError:
            pass
        self.hide_after_id = None

    def show(self):
        self.after_id = None
        if self.tip_window is not None or not self.text:
            return
        try:
            x_pos = self.widget.winfo_pointerx() + 12
            y_pos = self.widget.winfo_pointery() + 16
            tip = tk.Toplevel(self.widget)
            tip.wm_overrideredirect(True)
            tip.wm_geometry(f"+{x_pos}+{y_pos}")
            tk.Label(
                tip,
                text=self.text,
                font=(FONT_FAMILY, 9),
                bg="#111827",
                fg=COLOR_TEXT_PRIMARY,
                relief="solid",
                borderwidth=1,
                padx=8,
                pady=5,
            ).pack()
            self.tip_window = tip
        except tk.TclError:
            self.tip_window = None

    def hide(self, _event=None):
        self.cancel_pending()
        self.cancel_pending_hide()
        tip = self.tip_window
        self.tip_window = None
        if tip is not None:
            try:
                tip.destroy()
            except tk.TclError:
                pass


class TelemetryDashboard:
    def __init__(self, root):
        self.root = root
        self.root.title("Live Network Telemetry")
        window_width, window_height = self.initial_window_size(
            self.root.winfo_screenwidth(),
            self.root.winfo_screenheight(),
        )
        default_geometry = f"{window_width}x{window_height}"
        requested_geometry = os.environ.get(WINDOW_GEOMETRY_ENV, "").strip()
        try:
            self.root.geometry(requested_geometry or default_geometry)
        except tk.TclError:
            self.root.geometry(default_geometry)
        self.root.minsize(WINDOW_MIN_WIDTH, WINDOW_MIN_HEIGHT)
        self.root.resizable(True, True)
        self.root.configure(bg=COLOR_BG)
        self._resize_after_id = None
        self.initialize_history_state()
        self.initialize_llm_monitor_state()
        self.latest_telemetry = None
        self.last_read_error = None
        self.build_ui()
        self.root.bind("<Configure>", self.schedule_responsive_layout, add="+")
        self.root.after_idle(self.apply_responsive_layout)
        self.poll_telemetry()
        self.poll_llm_performance()

    @staticmethod
    def initial_window_size(screen_width, screen_height):
        """Fit the initial dashboard inside the screen while allowing resizing."""
        width = min(
            WINDOW_DEFAULT_WIDTH,
            max(WINDOW_MIN_WIDTH, int(screen_width) - WINDOW_SCREEN_MARGIN_X),
        )
        height = min(
            WINDOW_DEFAULT_HEIGHT,
            max(
                WINDOW_INITIAL_MIN_HEIGHT,
                int(screen_height) - WINDOW_SCREEN_MARGIN_Y,
            ),
        )
        return width, height

    def initialize_history_state(self):
        """Create bounded, process-local history containers.

        Nothing in this state is written to telemetry or another file. Closing
        the dashboard therefore discards the complete chart history.
        """
        self.history = {
            key: deque(maxlen=HISTORY_MAX_POINTS)
            for key in (
                "time",
                "vehicles",
                "buses",
                "road_queue",
                "pending_demand",
                "congestion",
            )
        }
        self.queue_history = deque(maxlen=20)
        self.last_seen_frame = None
        self.last_seen_simulation_time = None
        self.last_sample_time = None

    def initialize_llm_monitor_state(self):
        """Initialize process-local LLM and loosely-attributed GPU samples."""
        self.latest_gpu = _gpu_none()
        self.latest_ai_control = {"armed": False, "model": "None"}
        self.latest_llm_turn = None
        self.llm_samples = []
        self.llm_seen_turns = set()
        self._last_agent_log_size = None

    def build_ui(self):
        header = tk.Frame(self.root, bg=COLOR_BG)
        header.pack(fill="x", padx=14, pady=(8, 4))
        self.header_title_lbl = tk.Label(
            header,
            text="LIVE NETWORK TELEMETRY",
            font=(FONT_FAMILY, 13, "bold"),
            bg=COLOR_BG,
            fg=COLOR_TEXT_PRIMARY,
        )
        self.header_title_lbl.pack(side="left")
        self.status_lbl = tk.Label(
            header,
            text="WAITING FOR DATA",
            font=(FONT_FAMILY, 10, "bold"),
            bg=COLOR_BG,
            fg=COLOR_WARNING,
        )
        self.status_lbl.pack(side="right")

        style = ttk.Style(self.root)
        style.configure("Telemetry.TNotebook", background=COLOR_BG, borderwidth=0)
        style.configure(
            "Telemetry.TNotebook.Tab",
            font=(FONT_FAMILY, 9, "bold"),
            padding=(16, 7),
        )
        self.export_all_btn = tk.Button(
            header,
            text="EXPORT ALL",
            command=self.export_all,
            font=(FONT_FAMILY, 8, "bold"),
            bg=COLOR_ACCENT,
            fg=COLOR_TEXT_PRIMARY,
            activebackground=COLOR_CARD_BORDER,
            activeforeground=COLOR_TEXT_PRIMARY,
            relief="flat",
            padx=10,
            pady=3,
        )
        self.export_all_btn.pack(side="right", padx=(10, 0))
        self.export_all_status_lbl = tk.Label(
            header,
            text="",
            font=(FONT_FAMILY, 8),
            bg=COLOR_BG,
            fg=COLOR_SUCCESS,
        )
        self.export_all_status_lbl.pack(side="right", padx=(8, 0))
        self.notebook = ttk.Notebook(
            self.root,
            style="Telemetry.TNotebook",
        )
        self.notebook.pack(fill="both", expand=True, padx=10, pady=(0, 8))
        self.summary_tab, self.summary_content = self.create_responsive_tab()
        (
            self.trends_tab,
            self.trends_content,
            self.trends_scroll_canvas,
        ) = self.create_scrollable_tab()
        (
            self.llm_tab,
            self.llm_content,
            self.llm_scroll_canvas,
        ) = self.create_scrollable_tab()
        self.notebook.add(self.summary_tab, text="SUMMARY")
        self.notebook.add(self.trends_tab, text="SESSION TRENDS")
        self.notebook.add(self.llm_tab, text="LLM PERFORMANCE")

        summary_controls = tk.Frame(self.summary_content, bg=COLOR_BG)
        self.summary_export_hint_lbl = tk.Label(
            summary_controls,
            text="CURRENT AUTHORITATIVE SNAPSHOT",
            font=(FONT_FAMILY, 8, "bold"),
            bg=COLOR_BG,
            fg=COLOR_TEXT_SECONDARY,
        )
        self.summary_export_status_lbl = tk.Label(
            summary_controls,
            text="",
            font=(FONT_FAMILY, 8),
            bg=COLOR_BG,
            fg=COLOR_SUCCESS,
        )
        # Per-tab export buttons were removed: EXPORT ALL is the single
        # export. The status label stays unpacked so export_summary_snapshot,
        # which EXPORT ALL reuses, still has somewhere to report.

        # One dense KPI strip replaces the old 3x3 card grid: the same
        # numbers in roughly a fifth of the vertical space.
        self.metrics_frame = tk.Frame(
            self.summary_content,
            bg=COLOR_CARD,
            highlightbackground=COLOR_CARD_BORDER,
            highlightthickness=1,
        )
        self.metrics_frame.pack(fill="x", padx=10, pady=(4, 3))
        self.vars = {}
        self.metric_cards = []
        self.metric_title_labels = []
        self.metric_tooltips = []
        for index, (title, key) in enumerate(SUMMARY_KPIS):
            if index:
                separator = tk.Frame(
                    self.metrics_frame, bg=COLOR_CARD_BORDER, width=1
                )
                separator.pack(side="left", fill="y", pady=4)
            cell = tk.Frame(self.metrics_frame, bg=COLOR_CARD)
            cell.pack(side="left", fill="both", expand=True, padx=5, pady=3)
            value = tk.Label(
                cell,
                text="--",
                font=(FONT_FAMILY, 13, "bold"),
                bg=COLOR_CARD,
                fg=COLOR_ACCENT,
            )
            value.pack(side="left")
            title_label = tk.Label(
                cell,
                text=title,
                font=(FONT_FAMILY, 8, "bold"),
                bg=COLOR_CARD,
                fg=COLOR_TEXT_SECONDARY,
            )
            title_label.pack(side="left", padx=(3, 0))
            self.metric_cards.append(cell)
            self.metric_title_labels.append(title_label)
            tooltip = HoverTooltip(cell, KPI_TOOLTIPS[key])
            tooltip.add_target(value)
            tooltip.add_target(title_label)
            self.metric_tooltips.append(tooltip)
            self.vars[key] = value

        recovery_card = tk.Frame(
            self.summary_content,
            bg=COLOR_CARD,
            highlightbackground=COLOR_DANGER,
            highlightthickness=1,
        )
        recovery_card.pack(fill="x", padx=25, pady=(7, 3))
        self.recovery_card = recovery_card
        # Idle is the overwhelmingly common case, so the panel shows a single
        # line then and expands to full diagnostics only when recovery is
        # actually doing something.
        self.discharge_compact_lbl = tk.Label(
            recovery_card,
            text="STATUS: NORMAL CONTROL",
            font=(FONT_FAMILY, 8, "bold"),
            bg=COLOR_CARD,
            fg=COLOR_TEXT_SECONDARY,
            anchor="w",
        )
        self.discharge_detail_widgets = []
        self.discharge_expanded = True
        self.recovery_title_lbl = tk.Label(
            recovery_card,
            text="NETWORK GRIDLOCK RECOVERY",
            font=(FONT_FAMILY, 9, "bold"),
            bg=COLOR_CARD,
            fg=COLOR_DANGER,
        )
        self.recovery_title_lbl.pack(anchor="w", padx=10, pady=(6, 1))
        self.discharge_selected_lbl = tk.Label(
            recovery_card,
            text="Selected: Auto (Recommended)",
            font=(FONT_FAMILY, 8, "bold"),
            bg=COLOR_CARD,
            fg=COLOR_ACCENT,
            anchor="w",
        )
        self.discharge_selected_lbl.pack(fill="x", padx=10)
        self.discharge_status_lbl = tk.Label(
            recovery_card,
            text="Status: IDLE",
            font=(FONT_FAMILY, 8, "bold"),
            bg=COLOR_CARD,
            fg=COLOR_TEXT_SECONDARY,
            anchor="w",
        )
        self.discharge_status_lbl.pack(fill="x", padx=10)
        self.discharge_reason_lbl = tk.Label(
            recovery_card,
            text="Reason: Normal signal control is active",
            font=(FONT_FAMILY, 8),
            bg=COLOR_CARD,
            fg=COLOR_TEXT_PRIMARY,
            anchor="w",
            justify="left",
        )
        self.discharge_reason_lbl.pack(fill="x", padx=10)
        self.discharge_recommendation_lbl = tk.Label(
            recovery_card,
            text="Recommended first action: Select Auto or a corridor, then start discharge",
            font=(FONT_FAMILY, 8),
            bg=COLOR_CARD,
            fg=COLOR_WARNING,
            anchor="w",
            justify="left",
        )
        self.discharge_recommendation_lbl.pack(fill="x", padx=10, pady=(0, 6))
        self.discharge_detail_widgets = [
            self.recovery_title_lbl,
            self.discharge_selected_lbl,
            self.discharge_status_lbl,
            self.discharge_reason_lbl,
            self.discharge_recommendation_lbl,
        ]
        self.set_discharge_expanded(False)

        self.build_phase_cycle_ui()

        self.intersection_title_lbl = tk.Label(
            self.summary_content,
            text="INTERSECTION PHASE STATES",
            font=(FONT_FAMILY, 11, "bold"),
            bg=COLOR_BG,
            fg=COLOR_TEXT_PRIMARY,
        )
        self.intersection_title_lbl.pack(anchor="w", padx=20, pady=(10, 3))
        diagram_frame = tk.Frame(self.summary_content, bg=COLOR_BG)
        diagram_frame.pack(fill="both", expand=True, padx=20, pady=5)
        self.diagram_frame = diagram_frame
        self.diagram_frame.grid_rowconfigure(0, weight=1)
        self.diagram_frame.grid_columnconfigure(
            0, weight=1, uniform="intersection_node_columns"
        )
        self.diagram_frame.grid_columnconfigure(
            1, weight=1, uniform="intersection_node_columns"
        )
        self.node_title_labels = []
        self.node_a_canvas = self.create_node_canvas(diagram_frame, "NODE A (x=300)")
        self.node_a_canvas.master.grid(
            row=0, column=0, sticky="nsew", padx=(0, 5)
        )
        self.node_b_canvas = self.create_node_canvas(diagram_frame, "NODE B (x=700)")
        self.node_b_canvas.master.grid(
            row=0, column=1, sticky="nsew", padx=(5, 0)
        )
        self.node_a_canvas.bind(
            "<Configure>", lambda _event: self.draw_node_intersections()
        )
        self.node_b_canvas.bind(
            "<Configure>", lambda _event: self.draw_node_intersections()
        )
        self.build_trends_ui()
        self.build_llm_performance_ui()
        self.bind_tab_mousewheel(self.trends_content, self.trends_scroll_canvas)
        self.bind_tab_mousewheel(self.llm_content, self.llm_scroll_canvas)

    def create_responsive_tab(self):
        """Return a tab whose content reflows with the available window size."""
        tab = tk.Frame(self.notebook, bg=COLOR_BG)
        content = tk.Frame(tab, bg=COLOR_BG)
        content.pack(fill="both", expand=True)
        return tab, content

    def create_scrollable_tab(self):
        """Return a width-responsive tab with vertical overflow scrolling.

        The inner frame always tracks the viewport width, so shrinking the
        dashboard never introduces horizontal scrolling.  Only content that
        genuinely cannot remain legible at the available height scrolls.
        """
        tab = tk.Frame(self.notebook, bg=COLOR_BG)
        canvas = tk.Canvas(
            tab,
            bg=COLOR_BG,
            highlightthickness=0,
            borderwidth=0,
            takefocus=False,
        )
        scrollbar = ttk.Scrollbar(
            tab,
            orient="vertical",
            command=canvas.yview,
        )
        canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        content = tk.Frame(canvas, bg=COLOR_BG)
        content_window = canvas.create_window(
            (0, 0), window=content, anchor="nw"
        )

        def refresh_scroll_region(_event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def fit_content_width(event):
            canvas.itemconfigure(content_window, width=max(1, event.width))
            canvas.after_idle(refresh_scroll_region)

        content.bind("<Configure>", refresh_scroll_region, add="+")
        canvas.bind("<Configure>", fit_content_width, add="+")
        return tab, content, canvas

    @staticmethod
    def scroll_tab_with_wheel(event, canvas):
        """Scroll one tab with Windows/macOS wheel or Linux wheel buttons."""
        delta = getattr(event, "delta", 0)
        button = getattr(event, "num", None)
        if delta:
            units = -1 if delta > 0 else 1
        elif button == 4:
            units = -1
        elif button == 5:
            units = 1
        else:
            return None
        canvas.yview_scroll(units * 3, "units")
        return "break"

    def bind_tab_mousewheel(self, widget, canvas):
        """Bind wheel scrolling to a tab and every widget currently inside it."""
        widget.bind(
            "<MouseWheel>",
            lambda event: self.scroll_tab_with_wheel(event, canvas),
            add="+",
        )
        widget.bind(
            "<Button-4>",
            lambda event: self.scroll_tab_with_wheel(event, canvas),
            add="+",
        )
        widget.bind(
            "<Button-5>",
            lambda event: self.scroll_tab_with_wheel(event, canvas),
            add="+",
        )
        for child in widget.winfo_children():
            self.bind_tab_mousewheel(child, canvas)

    @staticmethod
    def responsive_profile(width, height):
        """Return bounded dimensions and fonts for the current client area."""
        width = max(WINDOW_MIN_WIDTH, int(width))
        height = max(WINDOW_MIN_HEIGHT, int(height))
        scale = max(0.60, min(1.0, min(width / 900.0, height / 430.0)))
        compact = width < 700 or height < 700
        return {
            "compact": compact,
            "header_font": max(11, round(16 * scale)),
            "status_font": max(8, round(10 * scale)),
            "section_font": max(8, round(11 * scale)),
            "metric_title_font": max(6, round(9 * scale)),
            "metric_value_font": max(11, min(16, round(18 * scale))),
            "detail_font": max(6, round(8 * scale)),
            "metric_row_height": max(26, min(40, round(height * 0.07))),
            "phase_height": max(35, min(85, round(height * 0.11))),
            "node_height": max(45, min(100, round(height * 0.13))),
            "chart_height": max(72, min(150, round((height - 120) / 3))),
            "wraplength": max(310, width - 92),
        }

    def schedule_responsive_layout(self, event=None):
        """Debounce resize work so dragging a window edge remains fluid."""
        if event is not None and event.widget is not self.root:
            return
        if self._resize_after_id is not None:
            self.root.after_cancel(self._resize_after_id)
        self._resize_after_id = self.root.after(35, self.apply_responsive_layout)

    def apply_responsive_layout(self, width=None, height=None):
        """Scale dashboard content to the window instead of exposing scrollbars."""
        self._resize_after_id = None
        width = self.root.winfo_width() if width is None else width
        height = self.root.winfo_height() if height is None else height
        profile = self.responsive_profile(width, height)
        compact = profile["compact"]

        self.header_title_lbl.config(
            font=(FONT_FAMILY, profile["header_font"], "bold")
        )
        self.status_lbl.config(
            font=(FONT_FAMILY, profile["status_font"], "bold")
        )
        self.header_title_lbl.master.pack_configure(
            padx=10 if compact else 20,
            pady=(6, 4) if compact else (15, 10),
        )
        self.notebook.pack_configure(
            padx=6 if compact else 14,
            pady=(0, 6 if compact else 12),
        )

        metric_pad = 1 if compact else 3
        self.metrics_frame.pack_configure(
            padx=8 if compact else 20,
            pady=2 if compact else 5,
        )
        for row_index in range(3):
            self.metrics_frame.grid_rowconfigure(
                row_index,
                weight=1,
                uniform="summary_metric_rows",
                minsize=profile["metric_row_height"],
            )
        # The KPI strip is packed, not gridded: each cell sits side by side.
        for card, title_label in zip(self.metric_cards, self.metric_title_labels):
            card.pack_configure(padx=metric_pad, pady=max(2, metric_pad - 2))
            title_label.config(
                font=(FONT_FAMILY, profile["metric_title_font"], "bold")
            )
            title_label.pack_configure(padx=(2 if compact else 3, 0), pady=0)
        for value_label in self.vars.values():
            value_label.config(
                font=(FONT_FAMILY, profile["metric_value_font"], "bold")
            )
            value_label.pack_configure(
                padx=5 if compact else 10,
                pady=(0, 1 if compact else 4),
            )

        detail_font = (FONT_FAMILY, profile["detail_font"])
        detail_bold_font = (FONT_FAMILY, profile["detail_font"], "bold")
        self.export_all_btn.config(font=detail_bold_font)
        self.export_all_status_lbl.config(font=detail_font)
        self.summary_export_hint_lbl.config(font=detail_bold_font)
        self.summary_export_status_lbl.config(font=detail_font)

        self.recovery_title_lbl.config(font=detail_bold_font)
        self.discharge_selected_lbl.config(font=detail_bold_font)
        self.discharge_status_lbl.config(font=detail_bold_font)
        self.discharge_reason_lbl.config(
            font=detail_font,
            wraplength=profile["wraplength"],
        )
        self.discharge_recommendation_lbl.config(
            font=detail_font,
            wraplength=profile["wraplength"],
        )
        self.recovery_title_lbl.master.pack_configure(
            padx=10 if compact else 25,
            pady=(3, 2) if compact else (7, 3),
        )
        # pack_configure re-packs a hidden widget, which would undo the
        # collapsed recovery panel, so only touch these while expanded.
        if getattr(self, "discharge_expanded", True):
            self.recovery_title_lbl.pack_configure(
                pady=(3 if compact else 6, 0)
            )
            self.discharge_recommendation_lbl.pack_configure(
                pady=(0, 3 if compact else 6)
            )

        self.phase_title_lbl.config(
            font=(FONT_FAMILY, max(8, profile["section_font"] - 1), "bold")
        )
        self.phase_hint_lbl.config(
            text=(
                "Nominal cycle | live dots"
                if compact
                else "Nominal plan | marker repeats | dots show live state"
            ),
            font=(FONT_FAMILY, profile["detail_font"]),
        )
        self.phase_status_lbl.config(font=detail_bold_font, width=16 if compact else 18)
        self.phase_heading.pack_configure(
            padx=6 if compact else 10,
            pady=(2, 0) if compact else (7, 0),
        )
        self.phase_cycle_canvas.config(height=profile["phase_height"])
        self.phase_cycle_canvas.master.pack_configure(
            padx=10 if compact else 25,
            pady=(3, 2) if compact else (10, 3),
        )

        self.intersection_title_lbl.config(
            font=(FONT_FAMILY, profile["section_font"], "bold")
        )
        self.intersection_title_lbl.pack_configure(
            padx=10 if compact else 20,
            pady=(3, 1) if compact else (10, 3),
        )
        self.diagram_frame.pack_configure(
            padx=10 if compact else 20,
            pady=2 if compact else 5,
        )
        for title_label in self.node_title_labels:
            title_label.config(
                font=(FONT_FAMILY, max(8, profile["section_font"] - 1), "bold")
            )
            title_label.pack_configure(pady=(3 if compact else 10, 0))
        self.node_a_canvas.config(height=profile["node_height"])
        self.node_b_canvas.config(height=profile["node_height"])

        self.trends_info_lbl.config(
            text=(
                "IN MEMORY ONLY  |  Session data clears on reset/close"
                if compact
                else "IN MEMORY ONLY  |  Up to 1 sample per simulated second  |  "
                f"Latest {HISTORY_MAX_POINTS} samples  |  Clears on reset/close"
            ),
            font=(FONT_FAMILY, profile["detail_font"]),
        )
        self.clear_history_btn.config(font=detail_bold_font)
        self.trends_export_status_lbl.config(font=detail_font)
        for chart in self.trend_charts:
            chart["title_label"].config(
                font=(FONT_FAMILY, profile["metric_title_font"], "bold")
            )
            chart["canvas"].config(height=profile["chart_height"])

        self.llm_idle_lbl.config(font=detail_bold_font)
        self.llm_export_status_lbl.config(
            font=detail_font,
            wraplength=max(100, profile["wraplength"] // 3),
        )
        self.llm_reason_lbl.config(
            font=detail_font,
            wraplength=profile["wraplength"],
        )
        self.llm_reason_lbl.pack_configure(
            padx=10 if compact else 24,
            pady=(0, 1 if compact else 2),
        )
        self.llm_idle_lbl.master.pack_configure(
            padx=8 if compact else 20,
            pady=(4, 2) if compact else (10, 4),
        )
        for label in self.llm_section_labels:
            label.config(
                font=(FONT_FAMILY, profile["metric_title_font"], "bold")
            )
            label.master.pack_configure(
                padx=8 if compact else 20,
                pady=(1, 1) if compact else (3, 2),
            )
        for card, title_label in zip(
            self.llm_metric_cards, self.llm_metric_title_labels
        ):
            card.grid_configure(
                padx=1 if compact else 3,
                pady=1 if compact else 3,
            )
            title_label.config(
                font=(FONT_FAMILY, profile["metric_title_font"], "bold")
            )
            title_label.pack_configure(
                padx=4 if compact else 8,
                pady=(2 if compact else 5, 0),
            )
        for value_label in self.llm_vars.values():
            value_label.config(
                font=(
                    FONT_FAMILY,
                    max(8, profile["metric_value_font"] - 3),
                    "bold",
                )
            )
            value_label.pack_configure(
                padx=4 if compact else 8,
                pady=(0, 2 if compact else 5),
            )

        self.draw_phase_cycle(self.latest_telemetry)
        self.draw_node_intersections()
        self.draw_trend_charts()

    DISCHARGE_QUIET_STATES = frozenset({"IDLE", "INACTIVE", "", "COMPLETED"})

    def set_discharge_expanded(self, expanded):
        """Show full recovery diagnostics only while they carry information."""
        if getattr(self, "discharge_expanded", None) == bool(expanded):
            return
        self.discharge_expanded = bool(expanded)
        if expanded:
            self.discharge_compact_lbl.pack_forget()
            for widget in self.discharge_detail_widgets:
                widget.pack(fill="x", padx=10)
            self.recovery_title_lbl.pack_configure(anchor="w", pady=(6, 1))
            self.discharge_recommendation_lbl.pack_configure(pady=(0, 6))
            self.recovery_card.config(highlightbackground=COLOR_DANGER)
        else:
            for widget in self.discharge_detail_widgets:
                widget.pack_forget()
            self.discharge_compact_lbl.pack(fill="x", padx=10, pady=3)
            self.recovery_card.config(highlightbackground=COLOR_CARD_BORDER)

    def build_phase_cycle_ui(self):
        card = tk.Frame(
            self.summary_content,
            bg=COLOR_CARD,
            highlightbackground=COLOR_CARD_BORDER,
            highlightthickness=1,
        )
        card.pack(fill="x", padx=25, pady=(10, 3))
        heading = tk.Frame(card, bg=COLOR_CARD)
        heading.pack(fill="x", padx=10, pady=(7, 0))
        self.phase_heading = heading
        self.phase_title_lbl = tk.Label(
            heading,
            text="SIGNAL PHASE CYCLE",
            font=(FONT_FAMILY, 10, "bold"),
            bg=COLOR_CARD,
            fg=COLOR_TEXT_PRIMARY,
        )
        self.phase_title_lbl.pack(side="left")
        self.phase_status_lbl = tk.Label(
            heading,
            text="WAITING",
            width=18,
            anchor="e",
            font=(FONT_FAMILY, 8, "bold"),
            bg=COLOR_CARD,
            fg=COLOR_WARNING,
        )
        self.phase_status_lbl.pack(side="right")
        self.phase_hint_lbl = tk.Label(
            heading,
            text="Nominal plan | marker repeats | dots show live state",
            font=(FONT_FAMILY, 8),
            bg=COLOR_CARD,
            fg=COLOR_TEXT_SECONDARY,
        )
        self.phase_hint_lbl.pack(side="right", padx=(10, 12))
        self.phase_cycle_canvas = tk.Canvas(
            card,
            bg=COLOR_CARD,
            highlightthickness=0,
            height=145,
        )
        self.phase_cycle_canvas.pack(fill="x", padx=6, pady=(0, 5))
        self.phase_cycle_canvas.bind(
            "<Configure>",
            lambda _event: self.draw_phase_cycle(self.latest_telemetry),
        )

    def build_trends_ui(self):
        controls = tk.Frame(self.trends_content, bg=COLOR_BG)
        controls.pack(fill="x", padx=20, pady=(12, 6))
        self.trends_info_lbl = tk.Label(
            controls,
            text=(
                "IN MEMORY ONLY  |  Up to 1 sample per simulated second  |  "
                f"Latest {HISTORY_MAX_POINTS} samples  |  Clears on reset/close"
            ),
            font=(FONT_FAMILY, 9),
            bg=COLOR_BG,
            fg=COLOR_TEXT_SECONDARY,
        )
        self.trends_info_lbl.pack(side="left")
        self.clear_history_btn = tk.Button(
            controls,
            text="CLEAR HISTORY",
            command=self.clear_history,
            font=(FONT_FAMILY, 8, "bold"),
            bg=COLOR_CARD,
            fg=COLOR_TEXT_PRIMARY,
            activebackground=COLOR_CARD_BORDER,
            activeforeground=COLOR_TEXT_PRIMARY,
            relief="flat",
            padx=10,
            pady=4,
        )
        self.clear_history_btn.pack(side="right")
        self.trends_export_status_lbl = tk.Label(
            controls,
            text="",
            font=(FONT_FAMILY, 8),
            bg=COLOR_BG,
            fg=COLOR_SUCCESS,
        )

        charts = tk.Frame(self.trends_content, bg=COLOR_BG)
        charts.pack(fill="both", expand=True, padx=20, pady=(0, 12))
        self.trend_charts = []
        self.create_trend_chart(
            charts,
            "NETWORK OCCUPANCY",
            (
                ("vehicles", "Vehicles", COLOR_ACCENT),
                ("buses", "Buses", COLOR_WARNING),
            ),
        )
        self.create_trend_chart(
            charts,
            "QUEUE PRESSURE",
            (
                ("road_queue", "Stopped road queue", COLOR_DANGER),
                ("pending_demand", "Waiting to enter", COLOR_WARNING),
            ),
        )
        self.create_trend_chart(
            charts,
            "CONGESTION  (stopped road queue / active vehicles)",
            (("congestion", "Congestion %", COLOR_SUCCESS),),
            fixed_max=100.0,
        )

    def build_llm_performance_ui(self):
        """Build the live model-performance instrument panel."""
        controls = tk.Frame(self.llm_content, bg=COLOR_BG)
        controls.pack(fill="x", padx=20, pady=(10, 4))
        self.llm_idle_lbl = tk.Label(
            controls,
            text="LLM idle — arm a model in the control panel",
            font=(FONT_FAMILY, 9, "bold"),
            bg=COLOR_BG,
            fg=COLOR_TEXT_SECONDARY,
            anchor="w",
        )
        self.llm_idle_lbl.pack(side="left", fill="x", expand=True)
        self.llm_export_status_lbl = tk.Label(
            controls,
            text="",
            font=(FONT_FAMILY, 8),
            bg=COLOR_BG,
            fg=COLOR_SUCCESS,
        )

        self.llm_vars = {}
        self.llm_metric_cards = []
        self.llm_metric_title_labels = []
        self.llm_section_labels = []
        self._build_llm_card_grid(
            "LATEST TURN",
            (
                ("Armed Model", "model"),
                ("Guard Status", "last_status"),
                ("Latency", "latency"),
                ("Input / Output", "tokens"),
                ("Generation Rate", "tokens_per_sec"),
                ("TSP / DBL Flags", "firing"),
            ),
            columns=3,
        )
        self.llm_reason_lbl = tk.Label(
            self.llm_content,
            text="Reason: --",
            font=(FONT_FAMILY, 8),
            bg=COLOR_BG,
            fg=COLOR_TEXT_PRIMARY,
            anchor="w",
            justify="left",
        )
        self.llm_reason_lbl.pack(fill="x", padx=24, pady=(0, 2))
        self._build_llm_card_grid(
            "CURRENT GPU  (one-second sample; loosely attributed)",
            (
                ("VRAM Used / Total", "vram"),
                ("GPU Utilization", "gpu_util"),
                ("Power Draw", "power"),
                ("Temperature", "temperature"),
            ),
            columns=4,
        )
        self.llm_vram_bar = ttk.Progressbar(
            self.llm_vars["vram"].master,
            orient="horizontal",
            mode="determinate",
            maximum=100,
            value=0,
        )
        self.llm_vram_bar.pack(fill="x", padx=8, pady=(0, 5))
        self._build_llm_card_grid(
            "DASHBOARD SESSION ROLLUP",
            (
                ("Turns Observed", "turns"),
                ("Guard OK Rate", "ok_rate"),
                ("HELD Rate", "held_rate"),
                ("Average Latency", "avg_latency"),
                ("Average Tok/s", "avg_tokens_per_sec"),
                ("Peak VRAM", "peak_vram"),
                ("Average Power", "avg_power"),
                ("Output Tokens", "total_output_tokens"),
            ),
            columns=4,
        )

    def _build_llm_card_grid(self, title, metrics, columns):
        section = tk.Frame(self.llm_content, bg=COLOR_BG)
        section.pack(fill="x", padx=20, pady=(3, 2))
        section_title = tk.Label(
            section,
            text=title,
            font=(FONT_FAMILY, 9, "bold"),
            bg=COLOR_BG,
            fg=COLOR_TEXT_SECONDARY,
            anchor="w",
        )
        section_title.pack(fill="x", pady=(0, 2))
        self.llm_section_labels.append(section_title)
        grid = tk.Frame(section, bg=COLOR_BG)
        grid.pack(fill="x")
        for column in range(columns):
            grid.grid_columnconfigure(
                column,
                weight=1,
                uniform=f"llm_{title}_columns",
            )
        for index, (metric_title, key) in enumerate(metrics):
            row, column = divmod(index, columns)
            card = tk.Frame(
                grid,
                bg=COLOR_CARD,
                highlightbackground=COLOR_CARD_BORDER,
                highlightthickness=1,
            )
            card.grid(
                row=row,
                column=column,
                sticky="nsew",
                padx=3,
                pady=3,
            )
            title_label = tk.Label(
                card,
                text=metric_title,
                font=(FONT_FAMILY, 8, "bold"),
                bg=COLOR_CARD,
                fg=COLOR_TEXT_SECONDARY,
                anchor="w",
            )
            title_label.pack(fill="x", padx=8, pady=(5, 0))
            value_label = tk.Label(
                card,
                text="--",
                font=(FONT_FAMILY, 12, "bold"),
                bg=COLOR_CARD,
                fg=COLOR_ACCENT,
                anchor="w",
            )
            value_label.pack(fill="x", padx=8, pady=(0, 5))
            self.llm_metric_cards.append(card)
            self.llm_metric_title_labels.append(title_label)
            self.llm_vars[key] = value_label

    def create_trend_chart(self, parent, title, series, fixed_max=None):
        card = tk.Frame(
            parent,
            bg=COLOR_CARD,
            highlightbackground=COLOR_CARD_BORDER,
            highlightthickness=1,
        )
        card.pack(fill="both", expand=True, pady=5)
        title_label = tk.Label(
            card,
            text=title,
            font=(FONT_FAMILY, 9, "bold"),
            bg=COLOR_CARD,
            fg=COLOR_TEXT_SECONDARY,
        )
        title_label.pack(anchor="w", padx=10, pady=(7, 0))
        canvas = tk.Canvas(
            card,
            bg=COLOR_CARD,
            highlightthickness=0,
            height=150,
        )
        canvas.pack(fill="both", expand=True, padx=4, pady=(0, 4))
        chart = {
            "canvas": canvas,
            "series": series,
            "fixed_max": fixed_max,
            "title_label": title_label,
        }
        self.trend_charts.append(chart)
        canvas.bind("<Configure>", lambda _event: self.draw_trend_charts())
        return canvas

    def create_node_canvas(self, parent, title):
        card = tk.Frame(
            parent,
            bg=COLOR_CARD,
            highlightbackground=COLOR_CARD_BORDER,
            highlightthickness=1,
        )
        title_label = tk.Label(
            card,
            text=title,
            font=(FONT_FAMILY, 10, "bold"),
            bg=COLOR_CARD,
            fg=COLOR_TEXT_SECONDARY,
        )
        title_label.pack(pady=(10, 0))
        self.node_title_labels.append(title_label)
        canvas = tk.Canvas(
            card,
            bg=COLOR_CARD,
            highlightthickness=0,
            width=1,
            height=135,
        )
        canvas.pack(fill="both", expand=True)
        return canvas

    @staticmethod
    def build_nominal_phase_segments(timing):
        """Return the normal six-phase plan used by the controller."""
        green = max(1, int(timing.get("green_frames", 240)))
        yellow = max(1, int(timing.get("yellow_frames", 60)))
        all_red = max(1, int(timing.get("all_red_frames", 60)))
        ns_start = green + yellow + all_red
        cycle = 2 * ns_start
        return cycle, {
            "EW": (
                (0, green, "GREEN"),
                (green, green + yellow, "YELLOW"),
                (green + yellow, cycle, "RED"),
            ),
            "NS_A": (
                (0, ns_start, "RED"),
                (ns_start, ns_start + green, "GREEN"),
                (ns_start + green, ns_start + green + yellow, "YELLOW"),
                (ns_start + green + yellow, cycle, "RED"),
            ),
            "NS_B": (
                (0, ns_start, "RED"),
                (ns_start, ns_start + green, "GREEN"),
                (ns_start + green, ns_start + green + yellow, "YELLOW"),
                (ns_start + green + yellow, cycle, "RED"),
            ),
        }

    @staticmethod
    def phase_marker_fraction(frame_number, cycle_frames):
        cycle_frames = max(1, int(cycle_frames))
        return int(frame_number) % cycle_frames / cycle_frames

    @staticmethod
    def aggregate_signal_state(states):
        normalized = {state for state in states if state in ("GREEN", "YELLOW", "RED")}
        if not normalized:
            return "RED"
        if len(normalized) == 1:
            return normalized.pop()
        return "MIXED"

    def draw_phase_cycle(self, data):
        """Draw the repeating nominal plan and authoritative live-state dots."""
        canvas = self.phase_cycle_canvas
        canvas.delete("all")
        width, height = canvas.winfo_width(), canvas.winfo_height()
        if width < 240 or height < 45:
            return
        if not data:
            if hasattr(self, "phase_status_lbl"):
                self.phase_status_lbl.config(text="WAITING", fg=COLOR_WARNING)
            canvas.create_text(
                width / 2,
                height / 2,
                text="Waiting for signal timing...",
                fill=COLOR_TEXT_SECONDARY,
                font=(FONT_FAMILY, 9),
            )
            return

        signal_state = data.get("signal_state", {})
        timing = signal_state.get("timing", {})
        cycle_frames, segment_map = self.build_nominal_phase_segments(timing)
        frames_per_second = max(1, int(timing.get("frames_per_second", 60)))
        compact_chart = width < 650 or height < 90
        left = 82 if compact_chart else 142
        right = 8 if compact_chart else 14
        top = 15 if compact_chart else 25
        bottom = 10 if compact_chart else 25
        plot_width = max(1, width - left - right)
        row_gap = 2 if compact_chart else 5
        row_height = max(7, (height - top - bottom - 2 * row_gap) / 3)
        rows = (
            ("EW", "E-W" if compact_chart else "EAST-WEST CORRIDOR"),
            ("NS_A", "N-S NODE A" if compact_chart else "NORTH-SOUTH NODE A"),
            ("NS_B", "N-S NODE B" if compact_chart else "NORTH-SOUTH NODE B"),
        )
        state_colors = {
            "GREEN": COLOR_SUCCESS,
            "YELLOW": COLOR_WARNING,
            "RED": COLOR_DANGER,
            "MIXED": COLOR_ACCENT,
        }

        def frame_x(frame):
            return left + frame / cycle_frames * plot_width

        row_centers = {}
        for row_index, (row_key, row_label) in enumerate(rows):
            y1 = top + row_index * (row_height + row_gap)
            y2 = y1 + row_height
            row_centers[row_key] = (y1 + y2) / 2
            canvas.create_text(
                left - 10,
                (y1 + y2) / 2,
                text=row_label,
                anchor="e",
                fill=COLOR_TEXT_SECONDARY,
                font=(FONT_FAMILY, 6 if compact_chart else 8, "bold"),
            )
            for start, end, state in segment_map[row_key]:
                canvas.create_rectangle(
                    frame_x(start),
                    y1,
                    frame_x(end),
                    y2,
                    fill=state_colors[state],
                    outline=COLOR_CARD,
                    width=1,
                )

        green = max(1, int(timing.get("green_frames", 240)))
        yellow = max(1, int(timing.get("yellow_frames", 60)))
        all_red = max(1, int(timing.get("all_red_frames", 60)))
        all_red_ranges = (
            (green + yellow, green + yellow + all_red),
            (cycle_frames - all_red, cycle_frames),
        )
        chart_bottom = row_centers["NS_B"] + row_height / 2
        for start, end in all_red_ranges:
            canvas.create_rectangle(
                frame_x(start),
                top - 2,
                frame_x(end),
                chart_bottom + 2,
                outline=COLOR_TEXT_PRIMARY,
                dash=(3, 2),
                width=1,
            )
            if not compact_chart:
                canvas.create_text(
                    (frame_x(start) + frame_x(end)) / 2,
                    9,
                    text="ALL RED",
                    fill=COLOR_TEXT_SECONDARY,
                    font=(FONT_FAMILY, 7, "bold"),
                )

        boundaries = sorted(
            {
                0,
                green,
                green + yellow,
                green + yellow + all_red,
                green + yellow + all_red + green,
                cycle_frames - all_red,
                cycle_frames,
            }
        )
        displayed_boundaries = (0, cycle_frames) if compact_chart else boundaries
        for boundary in displayed_boundaries:
            x = frame_x(boundary)
            canvas.create_line(x, chart_bottom + 2, x, chart_bottom + 5, fill=COLOR_TEXT_SECONDARY)
            canvas.create_text(
                x,
                height - 7,
                text=f"{boundary / frames_per_second:g}s",
                anchor="w" if boundary == 0 else ("e" if boundary == cycle_frames else "center"),
                fill=COLOR_TEXT_SECONDARY,
                font=(FONT_FAMILY, 7),
            )

        marker_fraction = self.phase_marker_fraction(
            data.get("frame_number", 0), cycle_frames
        )
        marker_x = left + marker_fraction * plot_width
        canvas.create_line(
            marker_x,
            top - 7,
            marker_x,
            chart_bottom + 4,
            fill=COLOR_TEXT_PRIMARY,
            width=2,
        )
        canvas.create_polygon(
            marker_x - 5,
            top - 8,
            marker_x + 5,
            top - 8,
            marker_x,
            top - 2,
            fill=COLOR_TEXT_PRIMARY,
            outline="",
        )

        nodes = signal_state.get("nodes", {})
        node_a_signals = nodes.get("300", {}).get("signals", {})
        node_b_signals = nodes.get("700", {}).get("signals", {})
        live_states = {
            "EW": self.aggregate_signal_state(
                (
                    node_a_signals.get("EB"),
                    node_a_signals.get("WB"),
                    node_b_signals.get("EB"),
                    node_b_signals.get("WB"),
                )
            ),
            "NS_A": self.aggregate_signal_state(
                (node_a_signals.get("NB"), node_a_signals.get("SB"))
            ),
            "NS_B": self.aggregate_signal_state(
                (node_b_signals.get("NB"), node_b_signals.get("SB"))
            ),
        }
        for row_key, live_state in live_states.items():
            center_y = row_centers[row_key]
            canvas.create_oval(
                marker_x - 5,
                center_y - 5,
                marker_x + 5,
                center_y + 5,
                fill=state_colors[live_state],
                outline=COLOR_TEXT_PRIMARY,
                width=1,
            )

        discharge_active = bool(data.get("network_discharge", {}).get("active"))
        priority_active = any(
            node.get("priority_state", "NORMAL") != "NORMAL"
            for node in nodes.values()
            if isinstance(node, dict)
        )
        all_signals = [
            signal
            for node in nodes.values()
            if isinstance(node, dict)
            for signal in node.get("signals", {}).values()
        ]
        if discharge_active:
            status_text = "NETWORK DISCHARGE"
            status_color = COLOR_WARNING
        elif all_signals and all(signal == "RED" for signal in all_signals):
            status_text = "ALL RED ACTIVE"
            status_color = COLOR_DANGER
        elif priority_active:
            status_text = "PRIORITY OVERRIDE"
            status_color = COLOR_ACCENT
        else:
            status_text = "NORMAL PLAN"
            status_color = COLOR_SUCCESS
        if hasattr(self, "phase_status_lbl"):
            self.phase_status_lbl.config(text=status_text, fg=status_color)

    def draw_node_intersections(self):
        """Redraw both live node diagrams after telemetry or geometry changes."""
        data = self.latest_telemetry or {}
        signal_state = data.get("signal_state", {})
        nodes = signal_state.get("nodes", {})
        fallback = signal_state.get("current_phase", "UNKNOWN")
        node_a = nodes.get("300", {})
        node_b = nodes.get("700", {})
        self.draw_intersection(
            self.node_a_canvas,
            "A",
            node_a.get("phase", fallback),
            node_a.get("signals"),
        )
        self.draw_intersection(
            self.node_b_canvas,
            "B",
            node_b.get("phase", fallback),
            node_b.get("signals"),
        )

    def draw_intersection(self, canvas, node_key, phase, signals=None):
        canvas.delete("all")
        width, height = canvas.winfo_width(), canvas.winfo_height()
        if width < 10 or height < 10:
            return
        center_x, center_y = width // 2, height // 2
        shortest_side = min(width, height)
        road_w = max(14, min(40, int(shortest_side * 0.38)))
        signal_offset = max(9, min(40, int(shortest_side * 0.32)))
        canvas.create_rectangle(0, center_y - road_w // 2, width, center_y + road_w // 2, fill="#333D50", outline="")
        canvas.create_rectangle(center_x - road_w // 2, 0, center_x + road_w // 2, height, fill="#333D50", outline="")
        ew_color, ns_color = COLOR_DANGER, COLOR_DANGER
        if phase == "EW_GREEN":
            ew_color = COLOR_SUCCESS
        elif phase in ("EW_YELLOW", "CONFLICT_YELLOW"):
            ew_color = COLOR_WARNING
            ns_color = COLOR_WARNING
        elif phase == "NS_GREEN":
            ns_color = COLOR_SUCCESS
        elif phase == "NS_YELLOW":
            ns_color = COLOR_WARNING
        elif phase in ("PRIORITY_ACTIVE", "PRIORITY_CLEARING"):
            pass
        color_by_state = {
            "RED": COLOR_DANGER,
            "YELLOW": COLOR_WARNING,
            "GREEN": COLOR_SUCCESS,
        }
        signals = signals or {}
        eb_color = color_by_state.get(signals.get("EB"), ew_color)
        wb_color = color_by_state.get(signals.get("WB"), ew_color)
        nb_color = color_by_state.get(signals.get("NB"), ns_color)
        sb_color = color_by_state.get(signals.get("SB"), ns_color)
        radius = max(4, min(8, int(shortest_side * 0.08)))
        canvas.create_oval(center_x - signal_offset - radius, center_y - radius, center_x - signal_offset + radius, center_y + radius, fill=eb_color)
        canvas.create_oval(center_x + signal_offset - radius, center_y - radius, center_x + signal_offset + radius, center_y + radius, fill=wb_color)
        canvas.create_oval(center_x - radius, center_y - signal_offset - radius, center_x + radius, center_y - signal_offset + radius, fill=nb_color)
        canvas.create_oval(center_x - radius, center_y + signal_offset - radius, center_x + radius, center_y + signal_offset + radius, fill=sb_color)
        if height >= 55:
            canvas.create_text(center_x, height - 12, text=phase.replace("_", " "), fill=COLOR_TEXT_PRIMARY, font=(FONT_FAMILY, 8 if height < 90 else 9, "bold"))

    def safe_read_telemetry(self):
        self.last_read_error = None
        if not TELEMETRY_FILE.exists():
            return None
        try:
            with TELEMETRY_FILE.open("r", encoding="utf-8") as telemetry_file:
                data = json.load(telemetry_file)
            if not isinstance(data, dict):
                raise ValueError("Telemetry root must be an object")
            if not isinstance(data.get("timestamp"), (int, float)):
                raise ValueError("Telemetry timestamp is missing or invalid")
            if not isinstance(data.get("network_summary"), dict):
                raise ValueError("Telemetry network_summary is missing or invalid")
            return data
        except (OSError, json.JSONDecodeError, TypeError, ValueError) as exc:
            self.last_read_error = str(exc)
            return None

    @staticmethod
    def classify_status(data, now=None, stale_after=STALE_AFTER_SECONDS):
        if not data:
            return "CONNECTING"
        now = time.time() if now is None else now
        age = max(0.0, now - float(data.get("timestamp", 0.0)))
        if age > stale_after:
            return "STALE"
        if data.get("simulation_paused", False):
            return "PAUSED"
        return "LIVE"

    @staticmethod
    def _numeric(value):
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            return None
        number = float(value)
        return number if math.isfinite(number) else None

    @classmethod
    def _average(cls, values):
        numbers = [cls._numeric(value) for value in values]
        numbers = [number for number in numbers if number is not None]
        return sum(numbers) / len(numbers) if numbers else None

    def safe_read_ai_control(self, path=None):
        """Return current armed/model state without letting file races escape."""
        source = AI_CONTROL_FILE if path is None else Path(path)
        try:
            with source.open("r", encoding="utf-8") as control_file:
                payload = json.load(control_file)
            if not isinstance(payload, dict):
                raise ValueError("AI control root must be an object")
            return {
                "armed": payload.get("armed") is True,
                "model": str(payload.get("model", "None")),
            }
        except (OSError, json.JSONDecodeError, TypeError, ValueError):
            return {"armed": False, "model": "None"}

    def safe_read_agent_turns(self, path=None):
        """Read valid JSONL turns and skip malformed or partially-written lines."""
        source = AGENT_TURN_LOG_FILE if path is None else Path(path)
        try:
            size = source.stat().st_size
            records = []
            with source.open("r", encoding="utf-8") as turn_log:
                for line in turn_log:
                    try:
                        record = json.loads(line)
                    except (json.JSONDecodeError, TypeError, ValueError):
                        continue
                    turn = record.get("turn") if isinstance(record, dict) else None
                    if (
                        isinstance(turn, bool)
                        or not isinstance(turn, (int, float))
                        or not float(turn).is_integer()
                    ):
                        continue
                    records.append(record)
            return records, size
        except (OSError, TypeError, ValueError):
            return [], None

    @classmethod
    def build_llm_sample(cls, turn_record, gpu_stats):
        """Combine one guarded agent turn with the current dashboard GPU sample."""
        record = turn_record if isinstance(turn_record, dict) else {}
        gpu = gpu_stats if isinstance(gpu_stats, dict) else _gpu_none()
        flags = record.get("flags", {})
        if not isinstance(flags, dict):
            flags = {}
        tsp_on = 0
        dbl_on = 0
        for route_flags in flags.values():
            if not isinstance(route_flags, dict):
                continue
            tsp_on += route_flags.get("tsp") is True
            dbl_on += route_flags.get("dbl") is True
        return {
            "turn": int(record.get("turn", 0)),
            "timestamp": record.get("timestamp"),
            "model": str(record.get("model", "None")),
            "status": str(record.get("status", "UNKNOWN")),
            "latency_ms": record.get("latency_ms"),
            "input_tokens": record.get("input_tokens"),
            "output_tokens": record.get("output_tokens"),
            "tokens_per_sec": record.get("tokens_per_sec"),
            "tsp_on_count": tsp_on,
            "dbl_on_count": dbl_on,
            "reason": str(record.get("reason", "") or ""),
            "vram_used_mb": gpu.get("vram_used_mb"),
            "vram_total_mb": gpu.get("vram_total_mb"),
            "gpu_util_pct": gpu.get("gpu_util_pct"),
            "power_w": gpu.get("power_w"),
            "temp_c": gpu.get("temp_c"),
            "pax_per_min_recent": record.get("pax_per_min_recent"),
        }

    @classmethod
    def summarize_llm_samples(cls, samples):
        """Return reconciled dashboard-session aggregates for valid samples."""
        rows = [sample for sample in samples if isinstance(sample, dict)]
        turns = len(rows)
        ok_turns = sum(row.get("status") == "OK" for row in rows)
        held_turns = sum(row.get("status") == "HELD_ALL_OFF" for row in rows)
        vram_values = [cls._numeric(row.get("vram_used_mb")) for row in rows]
        vram_values = [value for value in vram_values if value is not None]
        output_values = [cls._numeric(row.get("output_tokens")) for row in rows]
        output_values = [value for value in output_values if value is not None]
        return {
            "turns": turns,
            "ok_rate_pct": (ok_turns / turns * 100.0) if turns else 0.0,
            "held_rate_pct": (held_turns / turns * 100.0) if turns else 0.0,
            "avg_latency_ms": cls._average(
                row.get("latency_ms") for row in rows
            ),
            "avg_tokens_per_sec": cls._average(
                row.get("tokens_per_sec") for row in rows
            ),
            "avg_vram_used_mb": cls._average(
                row.get("vram_used_mb") for row in rows
            ),
            "peak_vram_used_mb": max(vram_values) if vram_values else None,
            "avg_power_w": cls._average(row.get("power_w") for row in rows),
            "total_output_tokens": sum(
                output_values
            ) if output_values else 0.0,
        }

    @staticmethod
    def _format_number(value, suffix="", decimals=1):
        number = TelemetryDashboard._numeric(value)
        if number is None:
            return "N/A"
        return f"{number:,.{decimals}f}{suffix}"

    def update_llm_performance_display(self):
        control = self.latest_ai_control
        latest = self.latest_llm_turn
        gpu = self.latest_gpu
        armed = bool(control.get("armed", False))
        model = str(control.get("model", "None"))
        if not armed:
            self.llm_idle_lbl.config(
                text="LLM idle — arm a model in the control panel",
                fg=COLOR_TEXT_SECONDARY,
            )
        elif latest is None:
            self.llm_idle_lbl.config(
                text=f"{model} armed — waiting for the first completed turn",
                fg=COLOR_WARNING,
            )
        else:
            self.llm_idle_lbl.config(
                text=f"{model} armed — live turn monitoring",
                fg=COLOR_SUCCESS,
            )

        self.llm_vars["model"].config(
            text=model,
            fg=COLOR_SUCCESS if armed else COLOR_TEXT_SECONDARY,
        )
        status = str(latest.get("status", "--")) if latest else "--"
        status_color = {
            "OK": COLOR_SUCCESS,
            "HELD_ALL_OFF": COLOR_DANGER,
            "INVALID": COLOR_DANGER,
            "STANDDOWN_DISCHARGE": COLOR_WARNING,
        }.get(status, COLOR_TEXT_SECONDARY)
        self.llm_vars["last_status"].config(text=status, fg=status_color)
        if latest:
            self.llm_vars["latency"].config(
                text=self._format_number(latest.get("latency_ms"), " ms")
            )
            input_tokens = latest.get("input_tokens")
            output_tokens = latest.get("output_tokens")
            token_text = (
                f"{int(input_tokens)} / {int(output_tokens)}"
                if self._numeric(input_tokens) is not None
                and self._numeric(output_tokens) is not None
                else "N/A"
            )
            self.llm_vars["tokens"].config(text=token_text)
            self.llm_vars["tokens_per_sec"].config(
                text=self._format_number(latest.get("tokens_per_sec"), " tok/s", 2)
            )
            sample = self.build_llm_sample(latest, gpu)
            self.llm_vars["firing"].config(
                text=(
                    f"TSP {sample['tsp_on_count']} / "
                    f"DBL {sample['dbl_on_count']}"
                ),
                fg=(
                    COLOR_SUCCESS
                    if sample["tsp_on_count"] or sample["dbl_on_count"]
                    else COLOR_ACCENT
                ),
            )
            self.llm_reason_lbl.config(
                text=f"Reason: {str(latest.get('reason', '') or '--')}"
            )
        else:
            for key in ("latency", "tokens", "tokens_per_sec", "firing"):
                self.llm_vars[key].config(text="--", fg=COLOR_ACCENT)
            self.llm_reason_lbl.config(text="Reason: --")

        used = self._numeric(gpu.get("vram_used_mb"))
        total = self._numeric(gpu.get("vram_total_mb"))
        if used is None or total is None or total <= 0:
            self.llm_vars["vram"].config(text="N/A")
            self.llm_vram_bar["value"] = 0
            gpu_available = False
        else:
            self.llm_vars["vram"].config(text=f"{used:,.0f} / {total:,.0f} MB")
            self.llm_vram_bar["value"] = min(100.0, used / total * 100.0)
            gpu_available = True
        self.llm_vars["gpu_util"].config(
            text=self._format_number(gpu.get("gpu_util_pct"), "%", 0)
        )
        self.llm_vars["power"].config(
            text=self._format_number(gpu.get("power_w"), " W")
        )
        self.llm_vars["temperature"].config(
            text=self._format_number(gpu.get("temp_c"), " °C", 0)
        )
        if not gpu_available:
            self.llm_section_labels[1].config(
                text="CURRENT GPU  — N/A (no nvidia-smi)"
            )
        else:
            self.llm_section_labels[1].config(
                text="CURRENT GPU  (one-second sample; loosely attributed)"
            )

        rollup = self.summarize_llm_samples(self.llm_samples)
        self.llm_vars["turns"].config(text=str(rollup["turns"]))
        self.llm_vars["ok_rate"].config(text=f"{rollup['ok_rate_pct']:.1f}%")
        self.llm_vars["held_rate"].config(
            text=f"{rollup['held_rate_pct']:.1f}%"
        )
        self.llm_vars["avg_latency"].config(
            text=self._format_number(rollup["avg_latency_ms"], " ms")
        )
        self.llm_vars["avg_tokens_per_sec"].config(
            text=self._format_number(rollup["avg_tokens_per_sec"], " tok/s", 2)
        )
        self.llm_vars["peak_vram"].config(
            text=self._format_number(rollup["peak_vram_used_mb"], " MB", 0)
        )
        self.llm_vars["avg_power"].config(
            text=self._format_number(rollup["avg_power_w"], " W")
        )
        self.llm_vars["total_output_tokens"].config(
            text=f"{rollup['total_output_tokens']:,.0f}"
        )

    def poll_llm_performance(self):
        """Poll model logs and current GPU state on a separate one-second clock."""
        try:
            self.latest_gpu = poll_gpu_stats()
            self.latest_ai_control = self.safe_read_ai_control()
            records, log_size = self.safe_read_agent_turns()
            if (
                log_size is not None
                and self._last_agent_log_size is not None
                and log_size < self._last_agent_log_size
            ):
                self.llm_samples.clear()
                self.llm_seen_turns.clear()
                self.latest_llm_turn = None
            if log_size is not None:
                self._last_agent_log_size = log_size
            for record in records:
                turn = int(record["turn"])
                if turn in self.llm_seen_turns:
                    continue
                self.llm_samples.append(
                    self.build_llm_sample(record, self.latest_gpu)
                )
                self.llm_seen_turns.add(turn)
            if records:
                self.latest_llm_turn = records[-1]
            self.update_llm_performance_display()
        except Exception:
            # A malformed log, unavailable GPU tool, or widget race must never
            # stop either the LLM monitor or the main telemetry dashboard.
            pass
        finally:
            self.root.after(LLM_POLL_MILLISECONDS, self.poll_llm_performance)

    @staticmethod
    def _write_decision_rows(sheet, decisions):
        sheet.append(DECISION_EXPORT_HEADERS)
        for decision in decisions:
            flags = decision.get("flags", {})
            if not isinstance(flags, dict):
                flags = {}
            row = [
                decision.get("turn"),
                decision.get("timestamp"),
                decision.get("model"),
                decision.get("status"),
                decision.get("reason", ""),
                decision.get("pax_per_min_recent"),
            ]
            for route_id in ROUTE_ORDER:
                route_flags = flags.get(route_id, {})
                if not isinstance(route_flags, dict):
                    route_flags = {}
                row.extend(
                    (
                        bool(route_flags.get("tsp", False)),
                        bool(route_flags.get("dbl", False)),
                    )
                )
            locked_routes = decision.get("locked_routes", [])
            if not isinstance(locked_routes, (list, tuple, set, frozenset)):
                locked_routes = []
            row.extend(
                (
                    ",".join(str(route_id) for route_id in locked_routes),
                    str(decision.get("minimap", ""))[:32767],
                    str(decision.get("raw_output", ""))[:32767],
                )
            )
            sheet.append(row)

    @staticmethod
    def _write_telemetry_rows(sheet, telemetry_rows):
        sheet.append(TELEMETRY_EXPORT_HEADERS)
        for telemetry in telemetry_rows:
            sheet.append(
                [
                    telemetry.get("frame"),
                    telemetry.get("sim_time_s"),
                    telemetry.get("passengers_served_total"),
                    telemetry.get("passengers_served_bus"),
                    telemetry.get("passengers_served_car"),
                    telemetry.get("buses_served"),
                    telemetry.get("cars_served"),
                    telemetry.get("pax_per_min_cumulative"),
                    telemetry.get("pax_per_min_recent"),
                    telemetry.get("vehicles_in_network"),
                    telemetry.get("ai_armed"),
                    telemetry.get("ai_last_status"),
                    json.dumps(telemetry.get("queues_vehicles")),
                    json.dumps(telemetry.get("queues_passengers_est")),
                ]
            )

    @staticmethod
    def _write_llm_performance_rows(sheet, samples):
        sheet.append(LLM_PERFORMANCE_HEADERS)
        for sample in samples:
            sheet.append(
                [sample.get(header) for header in LLM_PERFORMANCE_HEADERS]
            )

    def _write_llm_summary_rows(self, sheet, samples):
        sheet.append(LLM_SUMMARY_HEADERS)
        models = sorted(
            {str(sample.get("model", "None")) for sample in samples}
        )
        for model in models:
            model_samples = [
                sample
                for sample in samples
                if str(sample.get("model", "None")) == model
            ]
            summary = self.summarize_llm_samples(model_samples)
            sheet.append(
                [
                    model,
                    summary["turns"],
                    round(summary["ok_rate_pct"], 2),
                    round(summary["held_rate_pct"], 2),
                    summary["avg_latency_ms"],
                    summary["avg_tokens_per_sec"],
                    summary["avg_vram_used_mb"],
                    summary["peak_vram_used_mb"],
                    summary["avg_power_w"],
                    summary["total_output_tokens"],
                ]
            )

    def export_all(self, destination=None):
        """Export session logs and dashboard LLM samples into one workbook."""
        decisions = read_jsonl(AGENT_TURN_LOG_FILE)
        telemetry_rows = read_jsonl(TELEMETRY_LOG_FILE)
        samples = [
            sample
            for sample in getattr(self, "llm_samples", [])
            if isinstance(sample, dict)
        ]
        if not decisions and not telemetry_rows and not samples:
            self.export_all_status_lbl.config(
                text="Nothing to export yet",
                fg=COLOR_WARNING,
            )
            return None

        try:
            from openpyxl import Workbook

            if destination is None:
                DASHBOARD_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
                destination = DASHBOARD_EXPORT_DIR / (
                    f"combined_export_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
                )
            else:
                destination = Path(destination)
                destination.parent.mkdir(parents=True, exist_ok=True)

            workbook = Workbook()
            decisions_sheet = workbook.active
            decisions_sheet.title = "Decisions"
            telemetry_sheet = workbook.create_sheet("Telemetry")
            performance_sheet = workbook.create_sheet("LLM Performance")
            summary_sheet = workbook.create_sheet("LLM Summary")

            self._write_decision_rows(decisions_sheet, decisions)
            self._write_telemetry_rows(telemetry_sheet, telemetry_rows)
            self._write_llm_performance_rows(performance_sheet, samples)
            self._write_llm_summary_rows(summary_sheet, samples)

            # Record the operator inputs that produced this session, so the
            # workbook is self-describing without the control panel.
            inputs_sheet = workbook.create_sheet("Control Panel Inputs")
            try:
                import main

                main.write_control_panel_inputs_sheet(inputs_sheet)
            except Exception:
                inputs_sheet.append(["section", "parameter", "value"])

            # Fold in the content the removed per-tab buttons used to produce,
            # so EXPORT ALL is genuinely the only export needed. Session
            # Trends in particular exists only in this process's memory.
            self._style_excel_sheets(
                (
                    decisions_sheet,
                    telemetry_sheet,
                    performance_sheet,
                    summary_sheet,
                    inputs_sheet,
                )
            )
            try:
                self.export_summary_snapshot(
                    workbook=workbook, prefix="Snapshot "
                )
            except Exception:
                pass
            try:
                self.export_session_trends(
                    workbook=workbook, prefix="Trends "
                )
            except Exception:
                pass

            workbook.save(destination)
            workbook.close()
            message = f"Exported combined workbook: {destination.name}"
            self.export_all_status_lbl.config(text=message, fg=COLOR_SUCCESS)
            print(f"Exported combined workbook: {destination}")
            return destination
        except Exception as exc:
            self.export_all_status_lbl.config(
                text=f"Export failed: {str(exc)[:100]}",
                fg=COLOR_DANGER,
            )
            return None

    def export_llm_performance(self, destination=None):
        """Export the in-memory turn/GPU samples without touching sim exports."""
        if not self.llm_samples:
            self.llm_export_status_lbl.config(
                text="No samples yet",
                fg=COLOR_WARNING,
            )
            return None
        try:
            from openpyxl import Workbook

            if destination is None:
                DASHBOARD_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
                destination = DASHBOARD_EXPORT_DIR / (
                    f"llm_perf_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
                )
            else:
                destination = Path(destination)
                destination.parent.mkdir(parents=True, exist_ok=True)

            workbook = Workbook()
            samples_sheet = workbook.active
            samples_sheet.title = "Samples"
            summary_sheet = workbook.create_sheet("Summary")
            self._write_llm_performance_rows(samples_sheet, self.llm_samples)
            self._write_llm_summary_rows(summary_sheet, self.llm_samples)
            self._style_excel_sheets((samples_sheet, summary_sheet))

            workbook.save(destination)
            workbook.close()
            self.llm_export_status_lbl.config(
                text=(
                    f"Exported {len(self.llm_samples)} samples to "
                    f"{destination.name}"
                ),
                fg=COLOR_SUCCESS,
            )
            return destination
        except Exception as exc:
            self.llm_export_status_lbl.config(
                text=f"Export failed: {str(exc)[:100]}",
                fg=COLOR_DANGER,
            )
            return None

    @staticmethod
    def _style_excel_sheets(sheets):
        from openpyxl.styles import Font

        for sheet in sheets:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column_cells in sheet.columns:
                width = max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in column_cells
                )
                sheet.column_dimensions[column_cells[0].column_letter].width = min(
                    max(10, width + 2), 42
                )

    @staticmethod
    def _add_session_trend_charts(workbook, trends_sheet, sample_count):
        """Add editable Excel charts backed by the exported trend table."""
        from openpyxl.chart import LineChart, Reference
        from openpyxl.chart.series import SeriesLabel

        charts_sheet = workbook.create_sheet("Session Charts")
        charts_sheet.sheet_view.showGridLines = False

        categories = Reference(
            trends_sheet,
            min_col=1,
            min_row=2,
            max_row=sample_count + 1,
        )
        chart_specs = (
            {
                "title": "Network Occupancy Over Time",
                "y_title": "Vehicles in network",
                "columns": (2, 3),
                "series": (
                    ("Vehicles", "2D8CFF"),
                    ("Buses", "F59E0B"),
                ),
                "anchor": "A1",
                "number_format": "0",
            },
            {
                "title": "Queue Pressure Over Time",
                "y_title": "Vehicles and pending arrivals",
                "columns": (4, 5),
                "series": (
                    ("Road queue", "EF4444"),
                    ("Pending demand", "F59E0B"),
                ),
                "anchor": "A16",
                "number_format": "0",
            },
            {
                "title": "Network Congestion Over Time",
                "y_title": "Congestion (%)",
                "columns": (6, 6),
                "series": (("Congestion", "2ECC71"),),
                "anchor": "A31",
                "number_format": "0.0",
            },
        )

        for spec in chart_specs:
            chart = LineChart()
            chart.title = spec["title"]
            chart.x_axis.title = "Simulation time (seconds)"
            chart.y_axis.title = spec["y_title"]
            chart.y_axis.numFmt = spec["number_format"]
            chart.legend.position = "t"
            chart.height = 7.2
            chart.width = 14.5
            chart.style = 13

            data = Reference(
                trends_sheet,
                min_col=spec["columns"][0],
                max_col=spec["columns"][1],
                min_row=1,
                max_row=sample_count + 1,
            )
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            for series, (label, color) in zip(chart.series, spec["series"]):
                series.tx = SeriesLabel(v=label)
                series.graphicalProperties.line.solidFill = color
                series.graphicalProperties.line.width = 24000

            charts_sheet.add_chart(chart, spec["anchor"])

        return charts_sheet

    def export_summary_snapshot(
        self, destination=None, workbook=None, prefix=""
    ):
        """Export the latest authoritative telemetry snapshot.

        With `workbook` supplied the sheets are appended to it and returned
        instead of being saved, so EXPORT ALL can fold this content in
        without duplicating the row building.
        """
        data = self.latest_telemetry
        if not isinstance(data, dict):
            self.summary_export_status_lbl.config(
                text="No live snapshot yet",
                fg=COLOR_WARNING,
            )
            return None
        try:
            from openpyxl import Workbook

            if destination is None:
                DASHBOARD_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
                destination = DASHBOARD_EXPORT_DIR / (
                    f"telemetry_summary_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
                )
            else:
                destination = Path(destination)
                destination.parent.mkdir(parents=True, exist_ok=True)

            shared = workbook is not None
            if shared:
                overview = workbook.create_sheet(f"{prefix}Overview")
            else:
                workbook = Workbook()
                overview = workbook.active
                overview.title = "Overview"
            overview.append(["metric", "value"])
            overview_rows = [
                ("telemetry_status", self.classify_status(data)),
                ("timestamp", data.get("timestamp")),
                ("frame_number", data.get("frame_number")),
                ("simulation_time_seconds", data.get("simulation_time_seconds")),
                ("simulation_paused", data.get("simulation_paused")),
                ("simulation_speed", data.get("simulation_speed")),
            ]
            for section_name in (
                "network_summary",
                "network_throughput",
                "network_discharge",
            ):
                section = data.get(section_name, {})
                if not isinstance(section, dict):
                    continue
                for key, value in section.items():
                    if isinstance(value, (dict, list)):
                        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
                    overview_rows.append((f"{section_name}.{key}", value))
            for row in overview_rows:
                overview.append(row)

            network_summary = data.get("network_summary", {})
            if not isinstance(network_summary, dict):
                network_summary = {}
            queues = network_summary.get("queues", {})
            passenger_queues = network_summary.get("queues_passengers_est", {})
            demand = data.get("demand_generation", {})
            if not isinstance(queues, dict):
                queues = {}
            if not isinstance(passenger_queues, dict):
                passenger_queues = {}
            if not isinstance(demand, dict):
                demand = {}
            queue_sheet = workbook.create_sheet(f"{prefix}Queues & Demand")
            queue_headers = [
                "approach",
                "queued_vehicles",
                "queued_passengers_est",
                "demand_active",
                "demand_model",
                "configured_rate_vpm",
                "effective_rate_vpm",
                "pending_arrivals",
                "peak_active",
            ]
            queue_sheet.append(queue_headers)
            approaches = sorted(set(queues) | set(passenger_queues) | set(demand))
            for approach in approaches:
                demand_row = demand.get(approach, {})
                if not isinstance(demand_row, dict):
                    demand_row = {}
                queue_sheet.append(
                    [
                        approach,
                        queues.get(approach),
                        passenger_queues.get(approach),
                        demand_row.get("active"),
                        demand_row.get("model"),
                        demand_row.get("configured_rate_vpm"),
                        demand_row.get("effective_rate_vpm"),
                        demand_row.get("pending_arrivals"),
                        demand_row.get("peak_active"),
                    ]
                )

            routes_sheet = workbook.create_sheet(f"{prefix}Routes")
            route_headers = [
                "route_id",
                "route_name",
                "active",
                "tsp_enabled",
                "dbl_enabled",
                "buses_on_route",
                "route_passengers_total",
                "nearest_bus_id",
                "nearest_bus_eta_sec",
                "nearest_bus_target_node_x",
                "nearest_bus_priority_granted",
                "nearest_bus_priority_pending",
            ]
            routes_sheet.append(route_headers)
            routes = data.get("routes", {})
            if not isinstance(routes, dict):
                routes = {}
            for route_id, route in sorted(routes.items()):
                route = route if isinstance(route, dict) else {}
                routes_sheet.append(
                    [route_id] + [route.get(header) for header in route_headers[1:]]
                )

            nodes_sheet = workbook.create_sheet(f"{prefix}Signal Nodes")
            node_headers = [
                "node",
                "node_x",
                "phase",
                "phase_index",
                "phase_timer_frames",
                "priority_state",
                "priority_timer_frames",
                "signal_eb",
                "signal_wb",
                "signal_nb",
                "signal_sb",
                "active_request_id",
                "active_request_route",
                "active_request_state",
                "queued_request_count",
                "reservation_count",
            ]
            nodes_sheet.append(node_headers)
            signal_state = data.get("signal_state", {})
            nodes = signal_state.get("nodes", {}) if isinstance(signal_state, dict) else {}
            if not isinstance(nodes, dict):
                nodes = {}
            for node_key, node in sorted(nodes.items()):
                node = node if isinstance(node, dict) else {}
                signals = node.get("signals", {})
                request = node.get("active_request")
                signals = signals if isinstance(signals, dict) else {}
                request = request if isinstance(request, dict) else {}
                queued_requests = node.get("queued_requests", [])
                nodes_sheet.append(
                    [
                        node_key,
                        node.get("node_x"),
                        node.get("phase"),
                        node.get("phase_index"),
                        node.get("phase_timer_frames"),
                        node.get("priority_state"),
                        node.get("priority_timer_frames"),
                        signals.get("EB"),
                        signals.get("WB"),
                        signals.get("NB"),
                        signals.get("SB"),
                        request.get("request_id"),
                        request.get("route_id"),
                        request.get("state"),
                        len(queued_requests) if isinstance(queued_requests, list) else None,
                        node.get("reservation_count"),
                    ]
                )

            sheets = (overview, queue_sheet, routes_sheet, nodes_sheet)
            self._style_excel_sheets(sheets)
            if shared:
                return sheets
            workbook.save(destination)
            workbook.close()
            self.summary_export_status_lbl.config(
                text=f"Exported {destination.name}",
                fg=COLOR_SUCCESS,
            )
            return destination
        except Exception as exc:
            self.summary_export_status_lbl.config(
                text=f"Export failed: {str(exc)[:100]}",
                fg=COLOR_DANGER,
            )
            return None

    def export_session_trends(self, destination=None, workbook=None, prefix=""):
        """Export only the bounded, in-memory Session Trends samples.

        With `workbook` supplied the sheets are appended to it and returned
        rather than saved, so EXPORT ALL keeps this dashboard-only history.
        """
        sample_count = len(self.history.get("time", ()))
        if sample_count == 0:
            self.trends_export_status_lbl.config(
                text="No trend samples yet",
                fg=COLOR_WARNING,
            )
            return None
        try:
            from openpyxl import Workbook

            if destination is None:
                DASHBOARD_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
                destination = DASHBOARD_EXPORT_DIR / (
                    f"session_trends_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
                )
            else:
                destination = Path(destination)
                destination.parent.mkdir(parents=True, exist_ok=True)

            headers = [
                "time_seconds",
                "vehicles",
                "buses",
                "road_queue",
                "pending_demand",
                "congestion_pct",
            ]
            keys = [
                "time",
                "vehicles",
                "buses",
                "road_queue",
                "pending_demand",
                "congestion",
            ]
            rows = list(zip(*(list(self.history[key]) for key in keys)))
            shared = workbook is not None
            if shared:
                # Already unique in a combined workbook, so no prefix needed.
                trends_sheet = workbook.create_sheet("Session Trends")
            else:
                workbook = Workbook()
                trends_sheet = workbook.active
                trends_sheet.title = "Session Trends"
            trends_sheet.append(headers)
            for row in rows:
                trends_sheet.append(row)

            charts_sheet = self._add_session_trend_charts(
                workbook,
                trends_sheet,
                len(rows),
            )

            summary_sheet = workbook.create_sheet(f"{prefix}Summary")
            summary_sheet.append(["metric", "value"])
            times = list(self.history["time"])
            summary_rows = [
                ("samples", len(rows)),
                ("start_time_seconds", times[0]),
                ("end_time_seconds", times[-1]),
                ("duration_seconds", times[-1] - times[0]),
            ]
            for key, label in (
                ("vehicles", "vehicles"),
                ("buses", "buses"),
                ("road_queue", "road_queue"),
                ("pending_demand", "pending_demand"),
                ("congestion", "congestion_pct"),
            ):
                values = list(self.history[key])
                summary_rows.extend(
                    [
                        (f"average_{label}", self._average(values)),
                        (f"peak_{label}", max(values) if values else None),
                    ]
                )
            for row in summary_rows:
                summary_sheet.append(row)

            self._style_excel_sheets((trends_sheet, summary_sheet))
            if shared:
                return (trends_sheet, charts_sheet, summary_sheet)
            workbook.save(destination)
            workbook.close()
            self.trends_export_status_lbl.config(
                text=f"Exported {len(rows)} samples to {destination.name}",
                fg=COLOR_SUCCESS,
            )
            return destination
        except Exception as exc:
            self.trends_export_status_lbl.config(
                text=f"Export failed: {str(exc)[:100]}",
                fg=COLOR_DANGER,
            )
            return None

    def poll_telemetry(self):
        try:
            data = self.safe_read_telemetry()
            status = self.classify_status(data)
            if status == "LIVE":
                self.status_lbl.config(text="LIVE", fg=COLOR_SUCCESS)
                self.record_history_sample(data)
                self.update_metrics(data)
            elif status == "PAUSED":
                self.status_lbl.config(text="PAUSED", fg=COLOR_WARNING)
                self.update_metrics(data)
            elif status == "STALE":
                self.status_lbl.config(text="STALE DATA", fg=COLOR_DANGER)
                self.update_metrics(data)
            elif self.last_read_error:
                self.status_lbl.config(text="TELEMETRY ERROR", fg=COLOR_DANGER)
            else:
                self.status_lbl.config(text="CONNECTING...", fg=COLOR_WARNING)
        except Exception as exc:
            self.last_read_error = str(exc)
            self.status_lbl.config(text="TELEMETRY ERROR", fg=COLOR_DANGER)
        finally:
            self.root.after(250, self.poll_telemetry)

    def clear_history(self):
        """Clear this dashboard process's samples without touching telemetry."""
        for values in self.history.values():
            values.clear()
        self.queue_history.clear()
        self.last_seen_frame = None
        self.last_seen_simulation_time = None
        self.last_sample_time = None
        if hasattr(self, "trends_export_status_lbl"):
            self.trends_export_status_lbl.config(text="")
        self.draw_trend_charts()

    def record_history_sample(self, data):
        """Record one bounded sample when simulated time advances sufficiently.

        Repeated telemetry reads at the same simulation frame are ignored. A
        backwards frame/time jump is treated as a simulation reset and starts a
        new session history rather than joining unrelated runs on one chart.
        """
        frame = int(data.get("frame_number", 0))
        simulation_time = float(data.get("simulation_time_seconds", frame / 60.0))
        reset_detected = (
            self.last_seen_frame is not None
            and (
                frame < self.last_seen_frame
                or simulation_time < self.last_seen_simulation_time
            )
        )
        if reset_detected:
            self.clear_history()
        elif self.last_seen_frame is not None and frame <= self.last_seen_frame:
            return False

        self.last_seen_frame = frame
        self.last_seen_simulation_time = simulation_time
        if (
            self.last_sample_time is not None
            and simulation_time - self.last_sample_time < HISTORY_SAMPLE_SECONDS
        ):
            return False

        summary = data.get("network_summary", {})
        total_vehicles = int(summary.get("total_vehicles", 0))
        total_buses = int(summary.get("total_buses", 0))
        total_queued = sum(summary.get("queues", {}).values())
        pending_demand = int(summary.get("pending_demand", 0))
        congestion = total_queued / total_vehicles * 100 if total_vehicles else 0.0
        sample = {
            "time": simulation_time,
            "vehicles": total_vehicles,
            "buses": total_buses,
            "road_queue": total_queued,
            "pending_demand": pending_demand,
            "congestion": congestion,
        }
        for key, value in sample.items():
            self.history[key].append(value)
        self.queue_history.append(total_queued)
        self.last_sample_time = simulation_time
        self.draw_trend_charts()
        return True

    @staticmethod
    def _nice_ceiling(value):
        if value <= 1:
            return 1.0
        magnitude = 10 ** math.floor(math.log10(value))
        normalized = value / magnitude
        for step in (1, 2, 5, 10):
            if normalized <= step:
                return float(step * magnitude)
        return float(10 * magnitude)

    def draw_trend_charts(self):
        for chart in getattr(self, "trend_charts", []):
            self.draw_line_chart(
                chart["canvas"],
                chart["series"],
                fixed_max=chart["fixed_max"],
            )

    def draw_line_chart(self, canvas, series, fixed_max=None):
        """Render a compact multi-series time chart on a standard Tk canvas."""
        canvas.delete("all")
        width = canvas.winfo_width()
        height = canvas.winfo_height()
        if width < 80 or height < 70:
            return

        left, right, top, bottom = 48, 14, 25, 24
        plot_width = max(1, width - left - right)
        plot_height = max(1, height - top - bottom)
        times = list(self.history["time"])
        if not times:
            canvas.create_text(
                width / 2,
                height / 2,
                text="Waiting for live samples...",
                fill=COLOR_TEXT_SECONDARY,
                font=(FONT_FAMILY, 9),
            )
            return

        all_values = [
            float(value)
            for key, _label, _color in series
            for value in self.history[key]
        ]
        y_max = (
            float(fixed_max)
            if fixed_max is not None
            else self._nice_ceiling(max(all_values, default=1.0))
        )
        y_max = max(1.0, y_max)
        time_min, time_max = times[0], times[-1]
        time_span = max(HISTORY_SAMPLE_SECONDS, time_max - time_min)

        for index in range(5):
            fraction = index / 4
            y = top + fraction * plot_height
            value = y_max * (1.0 - fraction)
            canvas.create_line(
                left,
                y,
                width - right,
                y,
                fill=COLOR_CARD_BORDER,
                width=1,
            )
            label = f"{value:.0f}" if y_max >= 10 else f"{value:.1f}"
            canvas.create_text(
                left - 6,
                y,
                text=label,
                anchor="e",
                fill=COLOR_TEXT_SECONDARY,
                font=(FONT_FAMILY, 7),
            )

        canvas.create_text(
            left,
            height - 7,
            text=f"{time_min:.1f}s",
            anchor="w",
            fill=COLOR_TEXT_SECONDARY,
            font=(FONT_FAMILY, 7),
        )
        canvas.create_text(
            width - right,
            height - 7,
            text=f"{time_max:.1f}s",
            anchor="e",
            fill=COLOR_TEXT_SECONDARY,
            font=(FONT_FAMILY, 7),
        )

        legend_x = left
        for key, label, color in series:
            values = list(self.history[key])
            canvas.create_line(
                legend_x,
                11,
                legend_x + 12,
                11,
                fill=color,
                width=3,
            )
            legend_text = f"{label}: {values[-1]:.1f}" if values else label
            legend_item = canvas.create_text(
                legend_x + 17,
                11,
                text=legend_text,
                anchor="w",
                fill=COLOR_TEXT_PRIMARY,
                font=(FONT_FAMILY, 8),
            )
            bounds = canvas.bbox(legend_item)
            legend_x = (bounds[2] + 22) if bounds else legend_x + 130

            points = []
            for sample_time, value in zip(times, values):
                x = left + (sample_time - time_min) / time_span * plot_width
                y = top + (1.0 - min(max(float(value), 0.0), y_max) / y_max) * plot_height
                points.extend((x, y))
            if len(points) >= 4:
                canvas.create_line(
                    *points,
                    fill=color,
                    width=2,
                    smooth=True,
                    splinesteps=12,
                )
            elif len(points) == 2:
                x, y = points
                canvas.create_oval(x - 2, y - 2, x + 2, y + 2, fill=color, outline="")

    def update_metrics(self, data):
        frame = int(data.get("frame_number", 0))
        summary = data.get("network_summary", {})
        total_vehicles = int(summary.get("total_vehicles", 0))
        total_buses = int(summary.get("total_buses", 0))
        total_queued = sum(summary.get("queues", {}).values())
        pending_demand = int(summary.get("pending_demand", 0))
        passengers = int(summary.get("passenger_volume", 0))
        congestion = total_queued / total_vehicles * 100 if total_vehicles else 0.0
        average_delay = (
            sum(self.queue_history) / len(self.queue_history)
            if self.queue_history
            else 0.0
        )
        buses = data.get("active_buses", data.get("approaching_buses", []))
        active_tsp = sum(bool(bus.get("tsp_active_triggered")) for bus in buses)
        active_dbl = sum(bool(bus.get("dbl_active_triggered")) for bus in buses)
        pending_tsp = sum(bool(bus.get("tsp_priority_pending")) for bus in buses)
        pending_dbl = sum(bool(bus.get("dbl_priority_pending")) for bus in buses)

        self.vars["vehicles"].config(text=str(total_vehicles))
        self.vars["buses"].config(text=str(total_buses))
        self.vars["passengers"].config(text=str(passengers))
        self.vars["queued"].config(text=f"{total_queued}/{pending_demand}")
        self.vars["delay"].config(text=f"{average_delay:.1f} v")
        self.vars["congestion"].config(text=f"{congestion:.1f}%")
        tsp_color = COLOR_SUCCESS if active_tsp else (COLOR_WARNING if pending_tsp else COLOR_ACCENT)
        dbl_color = COLOR_SUCCESS if active_dbl else (COLOR_WARNING if pending_dbl else COLOR_ACCENT)
        self.vars["tsp"].config(text=f"{active_tsp}/{pending_tsp}", fg=tsp_color)
        self.vars["dbl"].config(text=f"{active_dbl}/{pending_dbl}", fg=dbl_color)
        simulation_time = float(data.get("simulation_time_seconds", frame / 60.0))
        self.vars["timer"].config(text=f"{simulation_time:.1f} s")

        if hasattr(self, "discharge_status_lbl"):
            display = self.format_discharge_status(data.get("network_discharge", {}))
            self.discharge_selected_lbl.config(text=display["selected"])
            self.discharge_status_lbl.config(
                text=display["status"], fg=display["status_color"]
            )
            self.discharge_reason_lbl.config(text=display["reason"])
            self.discharge_recommendation_lbl.config(
                text=display["recommendation"]
            )
            runtime = data.get("network_discharge", {}) or {}
            raw_status = str(runtime.get("status", "IDLE")).upper()
            quiet = (
                not runtime.get("active", False)
                and raw_status in self.DISCHARGE_QUIET_STATES
            )
            self.set_discharge_expanded(not quiet)
            if quiet:
                self.discharge_compact_lbl.config(
                    text=f"STATUS: NORMAL CONTROL  |  {display['selected']}"
                )

        self.latest_telemetry = data
        self.root.update_idletasks()
        if hasattr(self, "phase_cycle_canvas"):
            self.draw_phase_cycle(data)
        self.draw_node_intersections()

    @staticmethod
    def format_discharge_status(status):
        status = status if isinstance(status, dict) else {}
        selected = status.get("selected", "Auto (Recommended)")
        state = status.get("status", "IDLE")
        stage = status.get("stage", "")
        discharged = int(status.get("vehicles_discharged", 0))
        reason = status.get("reason", "Normal signal control is active")
        recommendation = status.get(
            "recommendation", "Select Auto or a corridor, then start discharge"
        )
        status_line = f"Status: {state}"
        if stage:
            status_line += f"  |  Stage: {stage}"
        if discharged:
            status_line += f"  |  Discharged: {discharged}"
        status_color = {
            "DISCHARGING": COLOR_SUCCESS,
            "WAITING": COLOR_WARNING,
            "REQUESTED": COLOR_WARNING,
            "TRANSITIONING": COLOR_WARNING,
            "STOPPING": COLOR_WARNING,
            "COMPLETED": COLOR_SUCCESS,
        }.get(state, COLOR_TEXT_SECONDARY)
        return {
            "selected": f"Selected: {selected}",
            "status": status_line,
            "status_color": status_color,
            "reason": f"Reason: {reason}",
            "recommendation": f"Recommended first action: {recommendation}",
        }


if __name__ == "__main__":
    root = tk.Tk()
    TelemetryDashboard(root)
    root.mainloop()
