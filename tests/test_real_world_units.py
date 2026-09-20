"""Real-world unit conversions anchored on saturation flow.

Display and export only: nothing here touches the simulation. The tests
pin the anchoring arithmetic (k = S_real / S_sim, metres per pixel from
HCM jam spacing), the exactness of the time conversion, the honesty tags
on every row, and that the export records the conversion basis.
"""
import pytest

import src.ui.canvas_gemini as canvas
import src.ui.control_panel as control_panel
import src.core.main as main
import src.telemetry.real_world_units as units
from src.core.vehicle import Vehicle
from tests.helpers import NODE_A, NODE_B


def test_time_conversion_exact():
    assert units.frames_to_seconds(60) == 1.0
    assert units.frames_to_seconds(0) == 0.0
    assert units.frames_to_seconds(90) == 1.5
    assert units.FRAMES_PER_SECOND == 60


def test_saturation_anchor():
    assert units.saturation_scale(1291) == pytest.approx(1900 / 1291)
    assert units.saturation_scale(1900) == 1.0
    assert units.saturation_scale(1291, real_saturation_flow=1800) == pytest.approx(1800 / 1291)
    # No measurement yet: no scale, and nothing raises.
    assert units.saturation_scale(None) is None
    assert units.saturation_scale(0) is None
    assert units.saturation_headway_s(1900) == pytest.approx(3600 / 1900)
    assert units.REAL_SATURATION_FLOW_VEH_HR == 1900


def test_distance_derived():
    # The sim's queued spacing is the car length plus its standing gap.
    assert units.SIM_QUEUE_SPACING_PX == Vehicle(0, 0, "EB").length + 12
    mpp = units.meters_per_pixel()
    assert mpp == pytest.approx(units.REAL_JAM_SPACING_M / units.SIM_QUEUE_SPACING_PX)
    assert units.px_to_m(units.SIM_QUEUE_SPACING_PX) == pytest.approx(7.5)
    # Lane width converts with the same factor, and a changed anchor moves it.
    assert units.px_to_m(canvas.LANE) == pytest.approx(canvas.LANE * mpp)
    assert units.px_to_m(canvas.LANE, units.meters_per_pixel(jam_spacing_m=6.0)) == (
        pytest.approx(canvas.LANE * 6.0 / units.SIM_QUEUE_SPACING_PX)
    )
    # Speed rides on the same scale: px/frame -> m/s -> km/h.
    assert units.px_per_frame_to_kmh(1.0) == pytest.approx(60 * mpp * 3.6)


def test_flow_and_los_derived():
    k = units.saturation_scale(1291)
    assert units.demand_to_real_veh_hr(12, k) == pytest.approx(12 * 60 * k)
    assert units.demand_to_real_veh_hr(12, None) is None
    capacity = units.approach_capacity_veh_hr(0.5)
    assert capacity == pytest.approx(1900 * 0.5)          # one (critical) lane
    assert units.volume_to_capacity(475, capacity) == pytest.approx(0.5)
    assert units.approach_capacity_veh_hr(0.5, lanes=3) == pytest.approx(1900 * 1.5)
    assert units.volume_to_capacity(None, capacity) is None
    assert [units.hcm_level_of_service(d) for d in (5, 10, 15, 30, 50, 70, 100)] == [
        "A", "A", "B", "C", "D", "E", "F"
    ]
    assert units.hcm_level_of_service(None) is None


def _table(config_overrides=None, telemetry=None):
    config = dict(control_panel.global_config)
    config.update(config_overrides or {})
    return units.build_conversion_table(
        config, control_panel.approach_configs, control_panel.APPROACH_NAMES, telemetry
    )


def test_labels_present():
    sections = _table({"measured_saturation_flow": 1291.0})
    titles = [s["title"] for s in sections]
    assert titles == ["Time", "Capacity", "Distance", "Speed", "Flow", "Delay / LOS"]
    for section in sections:
        assert section["rows"], section["title"]
        for row in section["rows"]:
            assert row["tag"] in units.ALL_TAGS, row
            assert set(row) >= {"quantity", "sim", "real", "tag", "note"}
    by_title = {s["title"]: s["rows"] for s in sections}
    assert {r["tag"] for r in by_title["Time"]} == {units.TAG_EXACT}
    assert {r["tag"] for r in by_title["Capacity"]} == {units.TAG_ANCHOR}
    assert {r["tag"] for r in by_title["Distance"]} == {units.TAG_DERIVED}
    assert {r["tag"] for r in by_title["Speed"]} == {units.TAG_APPROX}
    assert {r["tag"] for r in by_title["Flow"]} == {units.TAG_DERIVED}
    assert {r["tag"] for r in by_title["Delay / LOS"]} == {units.TAG_DERIVED}
    speed_row = by_title["Speed"][0]
    assert "not calibrated to real speeds" in speed_row["note"]
    statement = units.anchor_statement(1291.0)
    assert statement.startswith("Real-world units anchored on saturation flow: sim S=1,291 veh/hr")
    assert "real 1,900 veh/hr/lane (HCM)" in statement


def test_table_values_follow_the_anchor():
    telemetry = {
        "frame_number": 600,
        "simulation_time_seconds": 10.0,
        "network_summary": {"mean_speed_px_per_frame": 0.4},
        "network_throughput": {
            "mean_stopped_delay_sec_per_vehicle": 25.0, "vehicles_served_total": 40,
        },
    }
    sections = {s["title"]: {r["quantity"]: r for r in s["rows"]}
                for s in _table({"measured_saturation_flow": 1291.0,
                                 "priority_eligibility_px": 500,
                                 "_active_vehicle_speed_scale": 0.5}, telemetry)}
    assert sections["Time"]["Simulation clock"]["real"] == "10.0 s"
    assert sections["Capacity"]["Scale factor k"]["real"] == f"{1900 / 1291:.2f}×"
    mpp = units.meters_per_pixel()
    assert sections["Distance"]["Lane width"]["real"] == f"{canvas.LANE * mpp:.2f} m"
    assert sections["Distance"]["TSP eligibility zone"]["real"] == f"{500 * mpp:.0f} m"
    assert sections["Speed"]["Free-flow speed (cars)"]["real"] == (
        f"{0.5 * 60 * mpp * 3.6:.1f} km/h"
    )
    assert sections["Speed"]["Current mean vehicle speed"]["real"] == (
        f"{0.4 * 60 * mpp * 3.6:.1f} km/h"
    )
    assert sections["Delay / LOS"]["Mean stopped delay per vehicle"]["real"] == "25.0 s/veh"
    assert sections["Delay / LOS"]["Level of service (HCM signalized)"]["real"] == "LOS C"


def test_units_vc_agrees_with_webster_by_construction():
    """Both views evaluate the busiest lane: v/c on the units tab equals
    Webster's y / (g/C) for the same approach, independent of the anchor."""
    import src.core.webster as webster

    config = dict(control_panel.global_config)
    config["measured_saturation_flow"] = 1291.0
    config["webster_splits"] = webster.compute_all_nodes(
        control_panel.approach_configs, s=1291.0, lost_time_sec=4.0
    )
    sections = units.build_conversion_table(
        config, control_panel.approach_configs, control_panel.APPROACH_NAMES
    )
    flow_rows = {r["quantity"]: r for s in sections if s["title"] == "Flow" for r in s["rows"]}
    split = config["webster_splits"][NODE_A]
    g_over_c = split["EW_green_sec"] / split["cycle_time_sec"]
    expected = split["y_ew"] / g_over_c
    note = flow_rows["EB Corridor demand"]["note"]
    assert "critical lane 40% of flow" in note
    assert f"v/c = {expected:.2f}" in note


def test_table_survives_missing_inputs():
    sections = _table({"measured_saturation_flow": None, "webster_splits": {}}, None)
    flat = {r["quantity"]: r for s in sections for r in s["rows"]}
    assert flat["Scale factor k"]["real"] == "--"
    assert flat["Level of service (HCM signalized)"]["real"] == "--"
    assert "needs Webster" in flat["EB Corridor demand"]["note"]


def test_conversions_in_export(tmp_path, monkeypatch):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    monkeypatch.setitem(control_panel.global_config, "measured_saturation_flow", 1291.0)
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", tmp_path / "telemetry.jsonl")
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", tmp_path / "turns.jsonl")
    monkeypatch.setattr(main, "BUS_EVENTS_LOG_PATH", tmp_path / "bus_events.jsonl")
    monkeypatch.setattr(main, "TELEMETRY_PATH", tmp_path / "missing_telemetry.json")
    destination = tmp_path / "export.xlsx"

    assert main.export_test_workbook("None", 60, 1, destination) == destination

    workbook = load_workbook(destination, data_only=True)
    try:
        assert "Unit Conversions" in workbook.sheetnames
        sheet = workbook["Unit Conversions"]
        assert [c.value for c in sheet[1]] == units.UNIT_CONVERSION_HEADERS
        rows = {
            (r[0], r[1]): r for r in sheet.iter_rows(min_row=2, values_only=True)
        }
        assert rows[("Anchor", "REAL_SATURATION_FLOW_VEH_HR")][3] == 1900
        assert rows[("Anchor", "REAL_JAM_SPACING_M")][3] == 7.5
        assert rows[("Anchor", "sim_measured_saturation_flow_veh_hr")][2] == 1291.0
        assert rows[("Anchor", "scale_factor_k")][3] == pytest.approx(1900 / 1291, abs=1e-4)
        assert rows[("Anchor", "meters_per_pixel")][3] == pytest.approx(0.25)
        assert rows[("Anchor", "statement")][3].startswith("Real-world units anchored")
        # Every table row is present with its tag.
        assert rows[("Distance", "Lane width")][4] == units.TAG_DERIVED
        assert rows[("Speed", "Free-flow speed (cars)")][4] == units.TAG_APPROX
        assert {r[4] for r in rows.values()} <= set(units.ALL_TAGS)
    finally:
        workbook.close()


def test_dashboard_has_units_tab_and_updates_it():
    import tkinter as tk
    from src.ui.telemetry_dashboard import TelemetryDashboard

    host = tk.Tk()
    try:
        dashboard = TelemetryDashboard(tk.Frame(host))
        assert len(dashboard.notebook.tabs()) == 4
        assert dashboard.notebook.tab(dashboard.units_tab, "text") == "Units"
        key = ("Time", "Simulation clock")
        # The first poll may already have read a telemetry file from disk;
        # what matters is that a new sample repaints the row in place.
        dashboard.update_units_display({"frame_number": 120, "simulation_time_seconds": 2.0})
        assert dashboard.units_row_widgets[key]["real"].cget("text") == "2.0 s"
        assert dashboard.units_row_widgets[key]["tag"].cget("text") == "[EXACT]"
    finally:
        host.destroy()
