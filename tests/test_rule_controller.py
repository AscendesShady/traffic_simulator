"""Rule-based comparator: same mechanism as the LLM, a rule decides.

The rule reads the same telemetry, emits the same positional flag object,
and flows through the same guard, decision.json, merge, turn log and
exports. These tests pin the rule's conditional behaviour (cross-queue
check, per-node cap, DBL obstruction), that selecting it bypasses every
model call, and that a rule run produces the same artefacts an LLM run does.
"""
import json
from types import SimpleNamespace

import pytest

import src.agents.agent as agent
import src.ui.control_panel as control_panel
import src.core.guard as guard
import src.core.main as main
import src.agents.rule_controller as rc
from src.telemetry.bus_event_log import BusEventTracker
from src.ui.canvas_gemini import H_Y, HEIGHT, INT_X, LANE, ROAD_W, STOP, WIDTH
from src.core.signal_controller import TSP_ACTION_EXTENDING, SignalController
from tests.helpers import make_bus_for_leg, NODE_A, NODE_B


@pytest.fixture(autouse=True)
def isolate_agent_files(tmp_path, monkeypatch):
    monkeypatch.setattr(agent, "TURN_LOG_PATH", tmp_path / "agent_turn_log.jsonl")
    monkeypatch.setattr(agent, "DECISION_PATH", tmp_path / "decision.json")
    monkeypatch.setattr(guard, "REJECT_LOG_PATH", tmp_path / "agent_rejects.log")


def bus(route_id, node_x, approach, eta_sec, passengers=45):
    return {
        "route_id": route_id,
        "bus_id": f"BUS_{route_id}",
        "route_leg": {"node_x": node_x, "approach": approach},
        "leg_state": "APPROACHING",
        "direction": approach,
        "distance_to_stop_bar_px": 200,
        "eta_to_stop_bar_sec_freeflow": eta_sec,
        "passengers": passengers,
        "priority_granted": False,
        "priority_clearing": False,
    }


def telemetry(buses, node_queues=None, obstructed=()):
    """A minimal snapshot in the exporter's shape, with per-node queues."""
    node_queues = node_queues or {}
    zero = {"EB": 0, "WB": 0, "NB": 0, "SB": 0}
    nodes = {
        str(node_x): {"queues_passengers_est": {**zero, **node_queues.get(str(node_x), {})}}
        for node_x in INT_X
    }
    return {
        "active_buses": buses,
        "signal_state": {"nodes": nodes},
        "routes": {
            route_id: {"dbl_lane_obstructed": route_id in obstructed}
            for route_id in guard.ROUTE_ORDER
        },
        "network_throughput": {"passengers_per_minute_recent": 40.0},
    }


def flags_by_route(decision):
    return {
        route_id: {"tsp": decision["tsp"][i], "dbl": decision["dbl"][i]}
        for i, route_id in enumerate(guard.ROUTE_ORDER)
    }


# --- the rule itself ----------------------------------------------------------

def test_rule_grants_tsp_for_uncontested_bus():
    snapshot = telemetry(
        [bus("R1_EB_A_NB", NODE_A, "EB", 20)],
        node_queues={str(NODE_A): {"NB": 8, "SB": 4}},
    )
    decision = rc.rule_based_decision(snapshot, decision_lag_sec=0.0)
    flags = flags_by_route(decision)

    assert flags["R1_EB_A_NB"] == {"tsp": True, "dbl": True}
    assert all(not v["tsp"] and not v["dbl"] for r, v in flags.items() if r != "R1_EB_A_NB")
    assert decision["reason"].startswith("rule: ")
    assert f"TSP R1_EB_A_NB@{NODE_A} 45pax vs cross 12<45" in decision["reason"]


def test_rule_withholds_tsp_for_bus_arriving_on_green():
    """A bus that reaches the bar inside the residual green gains nothing
    from priority; the rule spends no cross-street time on it. DBL is a
    lane reservation, not a signal grant, so it is unaffected."""
    snapshot = telemetry(
        [bus("R1_EB_A_NB", NODE_A, "EB", 20)],
        node_queues={str(NODE_A): {"NB": 8, "SB": 4}},
    )
    snapshot["active_buses"][0]["would_have_stopped"] = False
    flags = flags_by_route(rc.rule_based_decision(snapshot, decision_lag_sec=0.0))

    assert flags["R1_EB_A_NB"] == {"tsp": False, "dbl": True}
    reason = rc.rule_based_decision(snapshot, decision_lag_sec=0.0)["reason"]
    assert f"TSP R1_EB_A_NB@{NODE_A} arrives on green" in reason


def test_rule_withholds_tsp_for_contested_bus():
    """The conditional part: cross-street load at/above the threshold blocks
    TSP even for an otherwise actionable bus. DBL is unaffected."""
    snapshot = telemetry(
        [bus("R1_EB_A_NB", NODE_A, "EB", 20)],
        node_queues={str(NODE_A): {"NB": 40, "SB": 20}},
    )
    decision = rc.rule_based_decision(snapshot, decision_lag_sec=0.0)
    flags = flags_by_route(decision)

    assert flags["R1_EB_A_NB"]["tsp"] is False
    assert flags["R1_EB_A_NB"]["dbl"] is True
    assert f"withheld: TSP R1_EB_A_NB@{NODE_A} cross 60>=45" in decision["reason"]

    # Exactly at the threshold is withheld; one below is granted.
    at = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", 20)], {str(NODE_A): {"NB": 45}})
    below = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", 20)], {str(NODE_A): {"NB": 44}})
    assert flags_by_route(rc.rule_based_decision(at, 0.0))["R1_EB_A_NB"]["tsp"] is False
    assert flags_by_route(rc.rule_based_decision(below, 0.0))["R1_EB_A_NB"]["tsp"] is True

    # Only the conflicting phase counts: a queue on the bus's own axis is
    # not cross traffic.
    same_axis = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", 20)], {str(NODE_A): {"WB": 200}})
    assert flags_by_route(rc.rule_based_decision(same_axis, 0.0))["R1_EB_A_NB"]["tsp"] is True


def test_rule_threshold_is_tunable():
    snapshot = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", 20)], {str(NODE_A): {"NB": 30}})
    assert rc.RULE_CROSS_QUEUE_THRESHOLD_PAX == 45
    assert flags_by_route(rc.rule_based_decision(snapshot, 0.0))["R1_EB_A_NB"]["tsp"] is True
    strict = rc.rule_based_decision(snapshot, 0.0, cross_queue_threshold_pax=30)
    assert flags_by_route(strict)["R1_EB_A_NB"]["tsp"] is False
    assert "cross 30>=30" in strict["reason"]


def test_rule_caps_grants_per_node():
    """Two qualifying buses at Node A: only one TSP grant, the higher net
    passenger benefit; the other is named as capped. Node B is judged
    independently, so a bus there still gets its own grant."""
    snapshot = telemetry(
        [
            bus("R1_EB_A_NB", NODE_A, "EB", 20, passengers=45),
            bus("R4_WB_A_SB", NODE_A, "WB", 25, passengers=45),
            bus("R2_EB_B_NB", NODE_B, "EB", 30, passengers=45),
        ],
        node_queues={str(NODE_A): {"NB": 8, "SB": 4}},
    )
    decision = rc.rule_based_decision(snapshot, decision_lag_sec=0.0)
    flags = flags_by_route(decision)

    tsp_at_node_a = [r for r in ("R1_EB_A_NB", "R4_WB_A_SB") if flags[r]["tsp"]]
    assert len(tsp_at_node_a) == rc.MAX_TSP_GRANTS_PER_NODE == 1
    # Equal benefit (45-12 each): the earlier ETA wins, deterministically.
    assert tsp_at_node_a == ["R1_EB_A_NB"]
    assert f"TSP R4_WB_A_SB@{NODE_A} node cap 1" in decision["reason"]
    assert flags["R2_EB_B_NB"]["tsp"] is True
    # DBL is a lane reservation, not capped.
    assert flags["R1_EB_A_NB"]["dbl"] and flags["R4_WB_A_SB"]["dbl"]

    # Highest benefit wins when loads differ, regardless of ETA order.
    unequal = telemetry(
        [
            bus("R1_EB_A_NB", NODE_A, "EB", 20, passengers=20),
            bus("R4_WB_A_SB", NODE_A, "WB", 25, passengers=45),
        ],
        node_queues={str(NODE_A): {"NB": 8, "SB": 4}},
    )
    assert flags_by_route(rc.rule_based_decision(unequal, 0.0))["R4_WB_A_SB"]["tsp"] is True
    assert flags_by_route(rc.rule_based_decision(unequal, 0.0))["R1_EB_A_NB"]["tsp"] is False

    # The cap is a parameter for reporting sensitivity.
    two = rc.rule_based_decision(snapshot, 0.0, max_tsp_grants_per_node=2)
    assert flags_by_route(two)["R4_WB_A_SB"]["tsp"] is True


def test_rule_dbl_respects_obstruction():
    snapshot = telemetry(
        [bus("R1_EB_A_NB", NODE_A, "EB", 20)],
        obstructed=("R1_EB_A_NB",),
    )
    decision = rc.rule_based_decision(snapshot, decision_lag_sec=0.0)
    flags = flags_by_route(decision)

    assert flags["R1_EB_A_NB"]["dbl"] is False
    assert flags["R1_EB_A_NB"]["tsp"] is True  # TSP judged separately
    assert "DBL R1_EB_A_NB lane obstructed or queued" in decision["reason"]

    # The explicit queue count is independently fail-safe if a producer ever
    # emits an inconsistent combined obstruction boolean.
    snapshot["routes"]["R1_EB_A_NB"]["dbl_lane_obstructed"] = False
    snapshot["routes"]["R1_EB_A_NB"]["dbl_lane_queue_ahead"] = 1
    decision = rc.rule_based_decision(snapshot, decision_lag_sec=0.0)
    assert flags_by_route(decision)["R1_EB_A_NB"]["dbl"] is False


def test_rule_uses_the_agents_actionable_definition():
    """A bus that will already have crossed by the time the decision lands
    is not actionable for TSP (the agent's own horizon), but it still
    counts as approaching for DBL."""
    late = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", 5)])
    decision = rc.rule_based_decision(late, decision_lag_sec=8.0)
    flags = flags_by_route(decision)
    assert agent.is_actionable(agent.eta_at_decision_land(5, 8.0)) is False
    assert flags["R1_EB_A_NB"]["tsp"] is False
    assert flags["R1_EB_A_NB"]["dbl"] is True

    beyond = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", agent.ACTIONABLE_HORIZON_SEC + 10)])
    assert flags_by_route(rc.rule_based_decision(beyond, 0.0))["R1_EB_A_NB"]["tsp"] is False

    # Buses not on an unfinished leg are ignored entirely.
    finished = telemetry([{**bus("R1_EB_A_NB", NODE_A, "EB", 20), "route_leg": None}])
    assert rc.rule_based_decision(finished, 0.0)["reason"] == "rule: no grant"


def test_rule_is_deterministic_and_safe_on_empty_input():
    snapshot = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", 20), bus("R5_WB_B_SB", NODE_B, "WB", 12)])
    first = rc.rule_based_decision(snapshot, 0.0)
    assert all(rc.rule_based_decision(snapshot, 0.0) == first for _ in range(5))
    for empty in ({}, None, {"active_buses": "garbage"}, {"active_buses": [None, 3]}):
        decision = rc.rule_based_decision(empty, 0.0)
        assert decision["tsp"] == [False] * 6 and decision["dbl"] == [False] * 6
        assert decision["reason"] == "rule: no grant"


# --- same structure, same guard ------------------------------------------------

def test_rule_output_passes_guard():
    snapshot = telemetry(
        [bus("R1_EB_A_NB", NODE_A, "EB", 20), bus("R5_WB_B_SB", NODE_B, "WB", 12)],
        node_queues={str(NODE_B): {"EB": 50}},
    )
    decision = rc.rule_based_decision(snapshot, decision_lag_sec=0.0)

    # Exactly the positional surface the model is asked for.
    assert set(decision) == {"tsp", "dbl", "reason"}
    assert len(decision["tsp"]) == len(decision["dbl"]) == len(guard.ROUTE_ORDER)
    assert all(isinstance(v, bool) for v in decision["tsp"] + decision["dbl"])
    assert len(decision["reason"]) <= guard.MAX_REASON_LEN

    raw = json.dumps(decision)
    assert guard.validate_flags_positional(json.loads(raw)) == flags_by_route(decision)
    safe = guard.safe_decision(raw, 7, rc.RULE_MODEL_NAME)
    assert safe["status"] == "OK"
    assert safe["model"] == "rule-based"
    assert safe["flags"] == flags_by_route(decision)
    assert safe["reason"] == decision["reason"]
    # The rule never trips the reject log.
    assert not guard.REJECT_LOG_PATH.exists()


def test_rule_mode_skips_llm(monkeypatch):
    """With model="rule-based" the agent makes no Ollama or Gemini call; the
    turn still reports latency (near zero) and no tokens."""
    calls = []

    def fake_chat(**kwargs):
        calls.append(("ollama", kwargs))
        raise AssertionError("ollama must not be called in rule mode")

    def fake_gemini(*args, **kwargs):
        calls.append(("gemini", args))
        raise AssertionError("gemini must not be called in rule mode")

    monkeypatch.setattr(agent, "ollama", SimpleNamespace(chat=fake_chat))
    monkeypatch.setattr(agent, "_call_gemini", fake_gemini)
    monkeypatch.setattr(agent, "_call_ollama", lambda *a, **k: fake_chat())

    snapshot = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", 20)])
    state = {
        "telemetry": snapshot, "minimap": "unused by the rule",
        "locked_routes": set(), "raw_output": "", "call_metrics": {},
        "decision": {}, "status": "OK", "recent_decisions": [],
        "turn": 3, "model": control_panel.RULE_BASED_MODEL, "decision_lag_sec": 0.0,
    }
    update = agent.ai_turn(state)

    assert calls == []
    assert update["status"] == "OK"
    metrics = update["call_metrics"]
    assert metrics["input_tokens"] is None and metrics["output_tokens"] is None
    assert 0 <= metrics["latency_ms"] < 1000
    parsed = json.loads(update["raw_output"])
    assert parsed == rc.rule_based_decision(snapshot, decision_lag_sec=0.0)

    guarded = agent.anti_cheat({**state, **update})
    assert guarded["decision"]["status"] == "OK"
    assert guarded["decision"]["flags"]["R1_EB_A_NB"] == {"tsp": True, "dbl": True}
    assert rc.is_rule_model("rule-based") and rc.is_rule_model(" Rule-Based ")
    assert not rc.is_rule_model("gemma4:12b") and not rc.is_rule_model(None)


def test_rule_mode_plans_with_near_zero_lag():
    """The rule's horizon is ~0 by construction from turn one, and a rule
    turn's (rounded-to-zero) latency never feeds the model carry-forward."""
    assert rc.RULE_DECISION_LAG_SEC == 0.0
    assert agent.turn_decision_lag("rule-based", 8.0) == 0.0
    assert agent.turn_decision_lag("rule-based", 12.5) == 0.0
    assert agent.turn_decision_lag("gemma4:12b", 12.5) == 12.5
    assert agent.turn_decision_lag("None", agent.DEFAULT_DECISION_LAG_SEC) == 8.0
    # The generic carry-forward still treats 0 as unmeasured for models,
    # which is exactly why the rule needs its own constant.
    assert agent.decision_lag_seconds(0.0) == agent.DEFAULT_DECISION_LAG_SEC

    # A bus 5 s out: not actionable under the model default, actionable for
    # the rule -- the comparator must not be handicapped by a lag it lacks.
    close = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", 5)])
    assert flags_by_route(rc.rule_based_decision(close, agent.DEFAULT_DECISION_LAG_SEC))["R1_EB_A_NB"]["tsp"] is False
    assert flags_by_route(rc.rule_based_decision(close))["R1_EB_A_NB"]["tsp"] is True
    assert flags_by_route(rc.rule_based_decision(close, rc.RULE_DECISION_LAG_SEC))["R1_EB_A_NB"]["tsp"] is True
    # The loop wiring is the real thing, not a re-statement in the test.
    source = __import__("inspect").getsource(agent.run_forever)
    assert "turn_decision_lag(" in source
    assert "not rule_controller.is_rule_model(" in source


def test_rule_respects_locked_routes_like_a_model():
    """The locked-route overlay in anti_cheat applies to rule turns too: a
    route already granted keeps its live flags whatever the rule says."""
    snapshot = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", 20)], {str(NODE_A): {"NB": 100}})
    snapshot["routes"]["R1_EB_A_NB"].update({"tsp_enabled": True, "dbl_enabled": False})
    state = {
        "telemetry": snapshot, "minimap": "", "locked_routes": {"R1_EB_A_NB"},
        "raw_output": "", "call_metrics": {}, "decision": {}, "status": "OK",
        "recent_decisions": [], "turn": 1, "model": "rule-based", "decision_lag_sec": 0.0,
    }
    guarded = agent.anti_cheat({**state, **agent.ai_turn(state)})
    # Rule would withhold (cross 100) and grant DBL; the lock keeps live state.
    assert guarded["decision"]["flags"]["R1_EB_A_NB"] == {"tsp": True, "dbl": False}


def test_rule_selectable_in_control_panel(monkeypatch):
    monkeypatch.setattr(
        control_panel.subprocess, "run",
        lambda *a, **k: SimpleNamespace(
            returncode=0, stdout="NAME ID SIZE MODIFIED\nmodel-a:latest abc 1GB now\n"
        ),
    )
    assert control_panel.get_decision_sources() == [
        "None", "rule-based", control_panel.MAX_PRESSURE_MODEL, "model-a:latest"
    ]
    assert control_panel.RULE_BASED_MODEL == rc.RULE_MODEL_NAME
    # Persisted to ai_control.json through the ordinary path.
    control_panel.set_active_ai_model("rule-based", persist=False)
    assert control_panel.global_config["ai_runtime"]["model"] == "rule-based"


# --- same artefacts as an LLM run ----------------------------------------------

def _drive_bus_with_tsp(tracker, controller):
    """Run one EB bus through Node A under a TSP extension (the same
    scenario the bus-event tests use), logging it via the tracker."""
    control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] = True
    controller.phase = 0
    controller.timer = 97
    vehicle = make_bus_for_leg("R1_EB_A_NB", NODE_A, "RULE_BUS")
    vehicles = [vehicle]
    for frame in range(1, 3000):
        controller.update(vehicles)
        signals = controller.get_all_signals(INT_X)
        for v in list(vehicles):
            v.update(signals, INT_X, H_Y, ROAD_W, STOP, LANE, vehicles, controller)
            if (v.x < -60 or v.x > WIDTH + 60 or v.y < -60 or v.y > HEIGHT + 60) and v.passed_nodes:
                record = tracker.complete(v, frame, controller)
                vehicles.remove(v)
                return record
        tracker.observe(vehicles, frame, controller)
    raise AssertionError("bus never completed")


def test_rule_decision_logged_and_exported(tmp_path, monkeypatch):
    """A rule turn goes decision.json -> merge -> live flags, appears in the
    turn log with model 'rule-based' and a latency, and the export then has
    Decisions, LLM Performance and Bus Events rows exactly as an LLM run would."""
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    now = 1_900_000_000.0
    monkeypatch.setattr(agent.time, "time", lambda: now)
    monkeypatch.setattr(main.time, "time", lambda: now)
    runtime = {"armed": True, "model": "rule-based", "tick_seconds": 5, "last_status": "WAITING_FOR_DECISION"}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    monkeypatch.setitem(control_panel.global_config, "_run_uuid", "run-rule")

    # 1. The agent's turn, exactly as the graph runs it, with the rule deciding.
    snapshot = telemetry([bus("R1_EB_A_NB", NODE_A, "EB", 20)], {str(NODE_A): {"NB": 8, "SB": 4}})
    state = {
        "telemetry": snapshot, "minimap": "minimap text", "locked_routes": set(),
        "raw_output": "", "call_metrics": {}, "decision": {}, "status": "OK",
        "recent_decisions": [], "turn": 1, "model": "rule-based", "decision_lag_sec": 0.0,
        "run_uuid": "run-rule",
    }
    state.update(agent.ai_turn(state))
    state.update(agent.anti_cheat(state))
    agent.write_decision(state)

    decision_on_disk = json.loads(agent.DECISION_PATH.read_text(encoding="utf-8"))
    assert decision_on_disk["model"] == "rule-based"
    assert decision_on_disk["status"] == "OK"
    assert decision_on_disk["flags"]["R1_EB_A_NB"] == {"tsp": True, "dbl": True}
    assert decision_on_disk["reason"].startswith(f"rule: TSP R1_EB_A_NB@{NODE_A}")

    logged = [json.loads(line) for line in agent.TURN_LOG_PATH.read_text(encoding="utf-8").splitlines()]
    assert len(logged) == 1
    assert logged[0]["model"] == "rule-based"
    assert logged[0]["status"] == "OK"
    assert logged[0]["flags"]["R1_EB_A_NB"] == {"tsp": True, "dbl": True}
    assert logged[0]["reason"] == decision_on_disk["reason"]
    assert isinstance(logged[0]["latency_ms"], (int, float)) and logged[0]["latency_ms"] < 1000
    assert logged[0]["input_tokens"] is None and logged[0]["tokens_per_sec"] is None
    assert logged[0]["pax_per_min_recent"] == 40.0
    assert logged[0]["telemetry_snapshot"] == snapshot

    # 2. main.py merges it identically to a model decision.
    for route in control_panel.bus_routes_config.values():
        route["tsp_enabled"] = route["dbl_enabled"] = False
    assert main.merge_ai_decision(agent.DECISION_PATH) is True
    assert runtime["last_status"] == "OK" and runtime["last_turn"] == 1
    assert control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] is True
    assert control_panel.bus_routes_config["R1_EB_A_NB"]["dbl_enabled"] is True
    assert not any(
        r["tsp_enabled"] for k, r in control_panel.bus_routes_config.items() if k != "R1_EB_A_NB"
    )

    # 3. The granted bus is served by the same bounded TSP and logged per bus.
    bus_log = tmp_path / "bus_events.jsonl"
    controller = SignalController({"green_time": 100}, yellow_time=2, red_clearance_time=2, min_green_frames=30)
    record = _drive_bus_with_tsp(BusEventTracker(bus_log), controller)
    assert record["nodes"][0]["tsp_treated"] is True
    assert record["nodes"][0]["tsp_action"] == TSP_ACTION_EXTENDING

    # 4. The combined export carries all of it in the usual sheets.
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", agent.TURN_LOG_PATH)
    monkeypatch.setattr(main, "BUS_EVENTS_LOG_PATH", bus_log)
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", tmp_path / "telemetry.jsonl")
    monkeypatch.setattr(main, "TELEMETRY_PATH", tmp_path / "no_snapshot.json")
    destination = tmp_path / "rule_run.xlsx"
    assert main.export_test_workbook("rule-based", 60, 1, destination) == destination
    assert main.build_test_export_filename(
        "rule-based", 300, 7, "20260914_112746"
    ) == "rule-based_5min_7seed_14092026_112746.xlsx"

    workbook = load_workbook(destination, data_only=True)
    try:
        assert workbook.sheetnames == [
            "Decisions", "Telemetry", "AI Decision Audit",
            "LLM Performance", "LLM Summary",
            "Control Panel Inputs", "Control Panel Inputs (start)",
            "Bus Events", "Unit Conversions", "Experiment Summary",
        ]
        decisions = workbook["Decisions"]
        header = [c.value for c in decisions[1]]
        row = dict(zip(header, [c.value for c in decisions[2]]))
        assert row["model"] == "rule-based" and row["status"] == "OK"
        assert row["R1_EB_A_NB_tsp"] is True and row["R1_EB_A_NB_dbl"] is True
        assert row["reason"].startswith("rule: ")

        audit = workbook["AI Decision Audit"]
        audit_row = dict(
            zip([c.value for c in audit[1]], [c.value for c in audit[2]])
        )
        assert audit_row["model"] == "rule-based"
        assert audit_row["requested_tsp_routes"] == "R1_EB_A_NB"
        assert audit_row["requested_dbl_routes"] == "R1_EB_A_NB"
        assert audit_row["observation_minimap"] == "minimap text"

        perf = workbook["LLM Performance"]
        perf_row = dict(zip([c.value for c in perf[1]], [c.value for c in perf[2]]))
        assert perf_row["model"] == "rule-based"
        assert perf_row["tsp_on_count"] == 1 and perf_row["dbl_on_count"] == 1
        assert perf_row["latency_ms"] is not None and perf_row["input_tokens"] is None

        events = workbook["Bus Events"]
        event_row = dict(zip([c.value for c in events[1]], [c.value for c in events[2]]))
        assert event_row["bus_id"] == "RULE_BUS"
        assert event_row["tsp_treated_any"] is True
        assert event_row["node1_tsp_action"] == TSP_ACTION_EXTENDING
        assert event_row["passengers"] == 45
    finally:
        workbook.close()
