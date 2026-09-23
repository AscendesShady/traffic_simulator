import json
from pathlib import Path
import re
import threading
import time
from types import SimpleNamespace

import pytest

import src.agents.agent as agent
import src.ui.control_panel as control_panel
import src.core.guard as guard
import src.core.main as main
from src.core.signal_controller import SignalController
from tests.helpers import make_bus_for_leg, NODE_A, NODE_B


@pytest.fixture(autouse=True)
def isolate_agent_turn_log(tmp_path, monkeypatch):
    monkeypatch.setattr(agent, "TURN_LOG_PATH", tmp_path / "agent_turn_log.jsonl")


RUN_ID = "run-under-test"


@pytest.fixture(autouse=True)
def _decisions_belong_to_this_run(monkeypatch):
    monkeypatch.setitem(control_panel.global_config, "_run_uuid", RUN_ID)


def valid_flags(tsp_route=None, dbl_route=None):
    flags = guard.all_off_flags()
    if tsp_route:
        flags[tsp_route]["tsp"] = True
    if dbl_route:
        flags[dbl_route]["dbl"] = True
    return flags


def positional_output(tsp_route=None, dbl_route=None, reason=""):
    return {
        "reason": reason,
        "tsp": [route_id == tsp_route for route_id in guard.ROUTE_ORDER],
        "dbl": [route_id == dbl_route for route_id in guard.ROUTE_ORDER],
    }


def agent_state(**overrides):
    state = {
        "telemetry": {},
        "minimap": "",
        "locked_routes": set(),
        "raw_output": "",
        "call_metrics": {},
        "decision": {},
        "status": "OK",
        "recent_decisions": [],
        "turn": 1,
        "model": "test-model",
    }
    state.update(overrides)
    return state


def test_guard_extracts_reasoning_output_and_enforces_exact_boolean_surface():
    flags = valid_flags(tsp_route="R1_EB_A_NB")
    raw = (
        "<think>ignore {nested reasoning}</think>\n```json\n"
        + json.dumps(positional_output(tsp_route="R1_EB_A_NB"))
        + "\n``` trailing text"
    )

    extracted = guard.extract_json(raw)
    assert guard.validate_flags_positional(extracted) == flags

    invalid = {"flags": valid_flags()}
    invalid["flags"]["R1_EB_A_NB"]["tsp"] = 1
    assert guard.validate_flags(invalid) is None

    extra_route = {"flags": valid_flags()}
    extra_route["flags"]["HALLUCINATED"] = {"tsp": False, "dbl": False}
    assert guard.validate_flags(extra_route) is None


def test_validate_flags_positional_maps_canonical_route_order():
    tsp = [False] * len(guard.ROUTE_ORDER)
    dbl = [False] * len(guard.ROUTE_ORDER)
    tsp[0] = True
    dbl[-1] = True

    mapped = guard.validate_flags_positional({"tsp": tsp, "dbl": dbl})

    assert list(mapped) == guard.ROUTE_ORDER
    assert mapped[guard.ROUTE_ORDER[0]] == {"tsp": True, "dbl": False}
    assert mapped[guard.ROUTE_ORDER[-1]] == {"tsp": False, "dbl": True}
    assert guard.ROUTE_ORDER == sorted(control_panel.bus_routes_config.keys())


def test_position_one_tsp_reaches_merge_and_signal_controller(
    tmp_path, monkeypatch
):
    route_id = guard.ROUTE_ORDER[0]
    assert route_id == "R1_EB_A_NB"
    decision = guard.safe_decision(
        json.dumps(
            positional_output(
                tsp_route=route_id,
                reason="Prioritize the approaching loaded bus.",
            )
        ),
        turn=7,
        model="nemotron-3-nano:4b",
    )
    assert decision["status"] == "OK"
    assert decision["flags"][route_id] == {"tsp": True, "dbl": False}

    decision_path = tmp_path / "decision.json"
    decision_path.write_text(json.dumps(dict(decision, run_uuid=RUN_ID)), encoding="utf-8")
    runtime = {"tick_seconds": 5, "last_status": "INACTIVE"}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)

    assert main.merge_ai_decision(decision_path) is True
    assert control_panel.bus_routes_config[route_id]["tsp_enabled"] is True

    bus = make_bus_for_leg(route_id, node_x=NODE_A, bus_id="POSITIONAL_TSP")
    controller = SignalController(control_panel.global_config)
    controller.update([bus])
    request = controller.get_node_status(NODE_A)["active_request"]

    assert request is not None
    assert request["route_id"] == route_id
    assert request["tsp_requested"] is True
    assert request["dbl_requested"] is False


@pytest.mark.parametrize(
    "payload",
    [
        {"tsp": [False] * 5, "dbl": [False] * 6},
        {"tsp": [False] * 6, "dbl": [False] * 7},
        {"tsp": [False, False, 1, False, False, False], "dbl": [False] * 6},
        {"tsp": [False] * 6, "dbl": [False, False, "false", False, False, False]},
        {"dbl": [False] * 6},
        {"tsp": [False] * 6},
    ],
)
def test_validate_flags_positional_rejects_malformed_arrays(payload):
    assert guard.validate_flags_positional(payload) is None


def test_safe_decision_does_not_accept_old_keyed_model_contract(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(guard, "REJECT_LOG_PATH", tmp_path / "rejects.log")

    decision = guard.safe_decision(
        json.dumps({"reason": "Old contract.", "flags": valid_flags()}),
        turn=1,
        model="test-model",
    )

    assert decision["status"] == "HELD_ALL_OFF"
    assert decision["flags"] == guard.all_off_flags()


def test_guard_captures_optional_reason_without_weakening_flags(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(guard, "REJECT_LOG_PATH", tmp_path / "rejects.log")
    flags = valid_flags(tsp_route="R1_EB_A_NB")

    valid = guard.safe_decision(
        json.dumps(
            positional_output(
                tsp_route="R1_EB_A_NB",
                reason="Prioritize the loaded bus.",
            )
        ),
        turn=1,
        model="test-model",
    )
    assert valid["status"] == "OK"
    assert valid["flags"] == flags
    assert valid["reason"] == "Prioritize the loaded bus."

    for optional_reason in (None, 42, ["not", "text"], ""):
        payload = positional_output(tsp_route="R1_EB_A_NB")
        if optional_reason is not None:
            payload["reason"] = optional_reason
        decision = guard.safe_decision(
            json.dumps(payload), turn=2, model="test-model"
        )
        assert decision["status"] == "OK"
        assert decision["flags"] == flags
        assert decision["reason"] == ""

    long_reason = guard.safe_decision(
        json.dumps(
            positional_output(
                tsp_route="R1_EB_A_NB", reason="x" * 5000
            )
        ),
        turn=3,
        model="test-model",
    )
    assert long_reason["status"] == "OK"
    assert long_reason["reason"] == "x" * guard.MAX_REASON_LEN

    punctuation = 'Release Node B because queue } exceeds the "normal" load.'
    punctuated = guard.safe_decision(
        json.dumps(
            positional_output(tsp_route="R1_EB_A_NB", reason=punctuation)
        ),
        turn=4,
        model="test-model",
    )
    assert punctuated["status"] == "OK"
    assert punctuated["reason"] == punctuation

    malformed_output = positional_output()
    malformed_output["tsp"][0] = 1
    held = guard.safe_decision(
        json.dumps({**malformed_output, "reason": "This must not survive."}),
        turn=5,
        model="test-model",
    )
    assert held["status"] == "HELD_ALL_OFF"
    assert held["flags"] == guard.all_off_flags()
    assert held["reason"] == ""


def test_guard_failure_is_all_off_logged_and_never_raises(tmp_path, monkeypatch):
    reject_log = tmp_path / "agent_rejects.log"
    monkeypatch.setattr(guard, "REJECT_LOG_PATH", reject_log)

    decision = guard.safe_decision("not JSON", turn=7, model="reasoning-model")

    assert decision["status"] == "HELD_ALL_OFF"
    assert decision["flags"] == guard.all_off_flags()
    assert not any(
        route_flags[flag]
        for route_flags in decision["flags"].values()
        for flag in ("tsp", "dbl")
    )
    logged = json.loads(reject_log.read_text(encoding="utf-8").splitlines()[0])
    assert logged == {"turn": 7, "model": "reasoning-model", "raw": "not JSON"}

    monkeypatch.setattr(guard, "REJECT_LOG_PATH", tmp_path / "missing" / "reject.log")
    assert guard.safe_decision(None, object(), None)["status"] == "HELD_ALL_OFF"


def test_control_panel_writes_atomic_agent_control_and_discovers_models(
    tmp_path, monkeypatch
):
    runtime = {
        "armed": True,
        "model": "local-model:latest",
        "tick_seconds": 9,
        "last_status": "OK",
        "last_turn": 3,
    }
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    destination = tmp_path / "ai_control.json"

    assert control_panel.write_ai_control(destination)
    assert json.loads(destination.read_text(encoding="utf-8")) == {
        "armed": True,
        "model": "local-model:latest",
        "tick_seconds": 9,
        "simulation_running": False,
        "control_mode": "assisted",
    }
    assert list(tmp_path.iterdir()) == [destination]


def test_batch_run_has_its_own_decision_interval(tmp_path, monkeypatch):
    """The Batch Run card's slider governs Benchmark Test/Batch Benchmark
    runs; the Single Run card's slider is untouched and used otherwise."""
    monkeypatch.setitem(
        control_panel.global_config, "ai_runtime",
        {"armed": True, "model": "local-model:latest", "tick_seconds": 5},
    )
    monkeypatch.setitem(
        control_panel.global_config, "batch_runtime",
        dict(control_panel.DEFAULT_BATCH_RUNTIME, tick_seconds=12),
    )
    destination = tmp_path / "ai_control.json"

    monkeypatch.setitem(control_panel.global_config, "test_running", False)
    control_panel.write_ai_control(destination)
    assert json.loads(destination.read_text(encoding="utf-8"))["tick_seconds"] == 5

    monkeypatch.setitem(control_panel.global_config, "test_running", True)
    control_panel.write_ai_control(destination)
    assert json.loads(destination.read_text(encoding="utf-8"))["tick_seconds"] == 12
    # The Single Run slider's own value is never mutated by a batch run.
    assert control_panel.global_config["ai_runtime"]["tick_seconds"] == 5

    monkeypatch.setattr(
        control_panel.subprocess,
        "run",
        lambda *args, **kwargs: SimpleNamespace(
            returncode=0,
            stdout="NAME ID SIZE MODIFIED\nmodel-a:latest abc 1GB now\nmodel-b:8b def 2GB now\n",
        ),
    )
    assert control_panel.get_ollama_models() == [
        "None",
        "model-a:latest",
        "model-b:8b",
    ]


def test_api_dropdown_hidden_without_key(monkeypatch):
    for environment_variable in control_panel.API_MODEL_REGISTRY:
        monkeypatch.delenv(environment_variable, raising=False)

    assert control_panel.get_api_models() == ["None"]


def _clear_api_keys(monkeypatch):
    for environment_variable in control_panel.API_MODEL_REGISTRY:
        monkeypatch.delenv(environment_variable, raising=False)


def test_api_dropdown_lists_models_with_key(monkeypatch):
    _clear_api_keys(monkeypatch)
    monkeypatch.setenv("GEMINI_API_KEY", "configured-for-test")

    assert control_panel.get_api_models() == [
        "None",
        *control_panel.API_MODEL_REGISTRY["GEMINI_API_KEY"],
    ]


def test_api_dropdown_lists_openai_models_with_key(monkeypatch):
    _clear_api_keys(monkeypatch)
    monkeypatch.setenv("OPENAI_API_KEY", "configured-for-test")

    assert control_panel.get_api_models() == [
        "None",
        *control_panel.API_MODEL_REGISTRY["OPENAI_API_KEY"],
    ]


def test_api_dropdown_lists_grok_models_with_key(monkeypatch):
    _clear_api_keys(monkeypatch)
    monkeypatch.setenv("GROK_API_KEY", "configured-for-test")

    assert control_panel.get_api_models() == [
        "None",
        *control_panel.API_MODEL_REGISTRY["GROK_API_KEY"],
    ]


def test_api_dropdown_lists_every_configured_provider(monkeypatch):
    _clear_api_keys(monkeypatch)
    monkeypatch.setenv("GEMINI_API_KEY", "configured-for-test")
    monkeypatch.setenv("OPENAI_API_KEY", "configured-for-test")
    monkeypatch.setenv("GROK_API_KEY", "configured-for-test")

    assert control_panel.get_api_models() == [
        "None",
        *control_panel.API_MODEL_REGISTRY["GEMINI_API_KEY"],
        *control_panel.API_MODEL_REGISTRY["OPENAI_API_KEY"],
        *control_panel.API_MODEL_REGISTRY["GROK_API_KEY"],
    ]


def test_single_active_model(tmp_path, monkeypatch):
    class FakeSelector:
        def __init__(self, value):
            self.value = value

        def set(self, value):
            self.value = value

    runtime = {
        "armed": True,
        "model": "local-model:latest",
        "tick_seconds": 5,
        "last_status": "OK",
    }
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    monkeypatch.setattr(control_panel, "AI_CONTROL_PATH", tmp_path / "ai_control.json")

    local_selector = FakeSelector("local-model:latest")
    api_selector = FakeSelector("gemini-2.5-flash")
    selected = control_panel.set_active_ai_model(
        api_selector.value, local_selector
    )
    assert selected == "gemini-2.5-flash"
    assert runtime["model"] == "gemini-2.5-flash"
    assert local_selector.value == "None"
    assert runtime["last_status"] == "MODEL_CHANGED_WAITING"

    local_selector.value = "local-model:latest"
    selected = control_panel.set_active_ai_model(
        local_selector.value, api_selector
    )
    assert selected == "local-model:latest"
    assert runtime["model"] == "local-model:latest"
    assert api_selector.value == "None"
    assert json.loads(control_panel.AI_CONTROL_PATH.read_text(encoding="utf-8"))[
        "model"
    ] == "local-model:latest"


def test_route_flag_button_repaint_reflects_external_decision(monkeypatch):
    class FakeButton:
        def __init__(self):
            self.options = {}

        def config(self, **options):
            self.options.update(options)

    route_id = "R2_EB_B_NB"
    route = control_panel.bus_routes_config[route_id]
    monkeypatch.setitem(route, "tsp_enabled", True)
    monkeypatch.setitem(route, "dbl_enabled", True)
    buttons = {"tsp": FakeButton(), "dbl": FakeButton()}
    monkeypatch.setattr(
        control_panel, "route_flag_buttons", {route_id: buttons}
    )

    control_panel.repaint_route_flag_buttons()

    # An enabled flag is a filled chip; the text carries the state too, so
    # it never depends on colour alone.
    assert buttons["tsp"].options["text"] == "TSP ON"
    assert buttons["tsp"].options["bg"] == control_panel.COLOR_SUCCESS
    assert buttons["dbl"].options["text"] == "DBL ON"
    assert buttons["dbl"].options["bg"] == control_panel.COLOR_ACCENT
    assert buttons["dbl"].options["fg"] == control_panel.COLOR_TEXT_PRIMARY

    monkeypatch.setitem(route, "tsp_enabled", False)
    monkeypatch.setitem(route, "dbl_enabled", False)
    control_panel.repaint_route_flag_buttons()
    for key in ("tsp", "dbl"):
        assert buttons[key].options["text"] == f"{key.upper()} OFF"
        assert buttons[key].options["bg"] == control_panel.COLOR_TOGGLE_OFF
        assert buttons[key].options["fg"] == control_panel.COLOR_TEXT_PRIMARY


def test_main_merge_applies_valid_flags_and_rejects_invalid_flags(
    tmp_path, monkeypatch
):
    runtime = {
        "armed": True,
        "model": "test-model",
        "tick_seconds": 5,
        "last_status": "WAITING_FOR_DECISION",
        "last_turn": 0,
    }
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    for route in control_panel.bus_routes_config.values():
        monkeypatch.setitem(route, "tsp_enabled", False)
        monkeypatch.setitem(route, "dbl_enabled", False)

    decision_path = tmp_path / "decision.json"
    decision_path.write_text(
        json.dumps(
            {"run_uuid": RUN_ID, 
                "schema_version": 1,
                "turn": 42,
                "model": "test-model",
                "timestamp": main.time.time(),
                "status": "OK",
                "flags": valid_flags(
                    tsp_route="R1_EB_A_NB",
                    dbl_route="R4_WB_A_SB",
                ),
            }
        ),
        encoding="utf-8",
    )

    assert main.merge_ai_decision(decision_path)
    assert control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] is True
    assert control_panel.bus_routes_config["R4_WB_A_SB"]["dbl_enabled"] is True
    assert runtime["last_status"] == "OK"
    assert runtime["last_turn"] == 42

    invalid = {"run_uuid": RUN_ID, "model": "test-model", "turn": 43,
               "timestamp": main.time.time(), "flags": valid_flags()}
    del invalid["flags"]["R6_WB_ONLY"]
    decision_path.write_text(json.dumps(invalid), encoding="utf-8")
    assert not main.merge_ai_decision(decision_path)
    assert runtime["last_status"] == "INVALID_DECISION"
    assert all(
        route[flag] is False
        for route in control_panel.bus_routes_config.values()
        for flag in ("tsp_enabled", "dbl_enabled")
    )


def test_missing_decision_leaves_manual_flags_unchanged(tmp_path, monkeypatch):
    runtime = {"last_status": "INACTIVE", "last_turn": 0}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    route = control_panel.bus_routes_config["R2_EB_B_NB"]
    monkeypatch.setitem(route, "tsp_enabled", True)

    assert not main.merge_ai_decision(tmp_path / "missing.json")
    assert route["tsp_enabled"] is True
    assert runtime["last_status"] == "WAITING_FOR_DECISION"


def test_fresh_decision_applies(tmp_path, monkeypatch):
    now = 1_900_000_000.0
    monkeypatch.setattr(main.time, "time", lambda: now)
    runtime = {"tick_seconds": 5, "last_status": "INACTIVE"}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    decision_path = tmp_path / "decision.json"
    flags = valid_flags(tsp_route="R1_EB_A_NB")
    decision_path.write_text(
        json.dumps({"run_uuid": RUN_ID, "timestamp": now, "status": "OK", "flags": flags}),
        encoding="utf-8",
    )

    assert main.merge_ai_decision(decision_path)
    assert runtime["last_status"] == "OK"
    assert control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] is True


def test_control_delay_is_measured_on_the_sim_clock(tmp_path, monkeypatch):
    """Snapshot frame -> first frame the decision takes effect, in sim
    seconds: not wall latency x pace, and not re-stamped by later merges."""
    now = 1_900_000_000.0
    monkeypatch.setattr(main.time, "time", lambda: now)
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", {"tick_seconds": 5})
    monkeypatch.setitem(main.run_metrics, "decision_applied", {})
    decision_path = tmp_path / "decision.json"
    decision_path.write_text(json.dumps({
        "run_uuid": RUN_ID, "timestamp": now, "status": "OK", "turn": 3,
        "telemetry_frame": 940, "flags": valid_flags(tsp_route="R1_EB_A_NB"),
    }), encoding="utf-8")

    monkeypatch.setitem(main._current_frame, "n", 1000)
    assert main.merge_ai_decision(decision_path)
    monkeypatch.setitem(main._current_frame, "n", 1030)
    assert main.merge_ai_decision(decision_path)     # same turn, merged again

    assert main.run_metrics["decision_applied"] == {3: (940, 1000)}
    assert main._control_delay_columns() == {
        "decisions_applied": 1,
        "control_delay_sim_sec_median": 1.0,
        "control_delay_sim_sec_p95": 1.0,
    }
    assert main._applied_columns(3) == [1000, 1.0]


def test_stale_decision_held_all_off(tmp_path, monkeypatch):
    now = 1_900_000_000.0
    monkeypatch.setattr(main.time, "time", lambda: now)
    runtime = {"tick_seconds": 5, "last_status": "INACTIVE"}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    decision_path = tmp_path / "decision.json"
    decision_path.write_text(
        json.dumps(
            {"run_uuid": RUN_ID, 
                "timestamp": now - 3600,
                "status": "OK",
                "flags": valid_flags(dbl_route="R4_WB_A_SB"),
            }
        ),
        encoding="utf-8",
    )

    assert not main.merge_ai_decision(decision_path)
    assert runtime["last_status"] == "STALE_DECISION"
    assert all(
        route[flag] is False
        for route in control_panel.bus_routes_config.values()
        for flag in ("tsp_enabled", "dbl_enabled")
    )


def test_missing_timestamp_treated_stale(tmp_path, monkeypatch):
    runtime = {"tick_seconds": 5, "last_status": "INACTIVE"}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    decision_path = tmp_path / "decision.json"
    decision_path.write_text(
        json.dumps({"run_uuid": RUN_ID, "status": "OK", "flags": valid_flags(tsp_route="R2_EB_B_NB")}),
        encoding="utf-8",
    )

    assert not main.merge_ai_decision(decision_path)
    assert runtime["last_status"] == "STALE_DECISION"
    assert all(
        route[flag] is False
        for route in control_panel.bus_routes_config.values()
        for flag in ("tsp_enabled", "dbl_enabled")
    )


def test_invalid_timestamp_treated_stale(tmp_path, monkeypatch):
    runtime = {"tick_seconds": 5, "last_status": "INACTIVE"}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    decision_path = tmp_path / "decision.json"

    for invalid_timestamp in ("not-a-time", float("nan"), float("inf")):
        decision_path.write_text(
            json.dumps(
                {"run_uuid": RUN_ID, 
                    "timestamp": invalid_timestamp,
                    "status": "OK",
                    "flags": valid_flags(tsp_route="R6_WB_ONLY"),
                }
            ),
            encoding="utf-8",
        )
        assert not main.merge_ai_decision(decision_path)
        assert runtime["last_status"] == "STALE_DECISION"
        assert all(
            route[flag] is False
            for route in control_panel.bus_routes_config.values()
            for flag in ("tsp_enabled", "dbl_enabled")
        )


def test_stale_threshold_scales_with_tick(tmp_path, monkeypatch):
    now = 1_900_000_000.0
    monkeypatch.setattr(main.time, "time", lambda: now)
    runtime = {"tick_seconds": 15, "last_status": "INACTIVE"}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    decision_path = tmp_path / "decision.json"
    flags = valid_flags(tsp_route="R3_EB_ONLY")

    decision_path.write_text(
        json.dumps({"run_uuid": RUN_ID, "timestamp": now - 20, "status": "OK", "flags": flags}),
        encoding="utf-8",
    )
    assert main.merge_ai_decision(decision_path)
    assert control_panel.bus_routes_config["R3_EB_ONLY"]["tsp_enabled"] is True

    decision_path.write_text(
        json.dumps({"run_uuid": RUN_ID, "timestamp": now - 50, "status": "OK", "flags": flags}),
        encoding="utf-8",
    )
    assert not main.merge_ai_decision(decision_path)
    assert runtime["last_status"] == "STALE_DECISION"


def test_agent_death_self_heals(tmp_path, monkeypatch):
    clock = {"now": 1_900_000_000.0}
    monkeypatch.setattr(main.time, "time", lambda: clock["now"])
    runtime = {"tick_seconds": 5, "last_status": "INACTIVE"}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    decision_path = tmp_path / "decision.json"
    decision_path.write_text(
        json.dumps(
            {"run_uuid": RUN_ID, 
                "timestamp": clock["now"],
                "status": "OK",
                "flags": valid_flags(dbl_route="R5_WB_B_SB"),
            }
        ),
        encoding="utf-8",
    )

    assert main.merge_ai_decision(decision_path)
    assert control_panel.bus_routes_config["R5_WB_B_SB"]["dbl_enabled"] is True

    clock["now"] += 16
    assert not main.merge_ai_decision(decision_path)
    assert runtime["last_status"] == "STALE_DECISION"
    assert all(
        route[flag] is False
        for route in control_panel.bus_routes_config.values()
        for flag in ("tsp_enabled", "dbl_enabled")
    )


def test_minimap_uses_only_approaching_unfinished_bus_legs():
    actionable = {
        "bus_id": "ACTIONABLE",
        "route_id": "R1_EB_A_NB",
        "direction": "NB",
        "route_leg": {"node_x": NODE_A, "movement": "LEFT"},
        "leg_state": "APPROACHING",
        "distance_to_stop_bar_px": 120.0,
        "eta_to_stop_bar_sec_freeflow": 2.0,
        "passengers": 45,
        "priority_granted": True,
    }
    completed = {
        **actionable,
        "bus_id": "COMPLETE",
        "route_leg": None,
        "leg_state": "COMPLETE",
        "distance_to_stop_bar_px": 0.0,
        "eta_to_stop_bar_sec_freeflow": 0.0,
    }
    past_bar = {
        **actionable,
        "bus_id": "PAST_BAR",
        "distance_to_stop_bar_px": -5.0,
        "eta_to_stop_bar_sec_freeflow": 0.0,
    }
    telemetry = {
        "simulation_time_seconds": 10.0,
        "routes": {
            route_id: {
                "active": True,
                "tsp_enabled": route_id == "R1_EB_A_NB",
                "dbl_enabled": False,
            }
            for route_id in guard.VALID_ROUTES
        },
        "network_throughput": {
            "passengers_per_minute": 100.0,
            "passengers_per_minute_recent": 120.0,
        },
        "network_summary": {"queues_passengers_est": {"EB": 20}},
        "signal_state": {"nodes": {str(NODE_A): {"phase": "EW_GREEN"}}},
        "active_buses": [completed, past_bar, actionable],
    }
    state = agent_state(telemetry=telemetry)

    approaching = agent.actionable_buses(telemetry)
    assert [bus["bus_id"] for bus in approaching["R1_EB_A_NB"]] == [
        "ACTIONABLE"
    ]

    minimap_update = agent.read_minimap(state)
    locked_update = agent.check_locked({**state, **minimap_update})
    assert "approaching_buses=1" in minimap_update["minimap"]
    assert "direction=NB" not in minimap_update["minimap"]
    assert "COMPLETE" not in minimap_update["minimap"]
    assert locked_update["locked_routes"] == {"R1_EB_A_NB"}
    assert "LOCKED_ROUTES=R1_EB_A_NB" in locked_update["minimap"]
    for position, route_id in enumerate(guard.ROUTE_ORDER, start=1):
        assert f"{position}) {route_id}:" in minimap_update["minimap"]


def test_minimap_has_no_recent_decisions():
    old_flags = valid_flags(tsp_route="R1_EB_A_NB", dbl_route="R4_WB_A_SB")
    state = agent_state(
        telemetry={"routes": {}, "active_buses": []},
        recent_decisions=[
            {"turn": 7, "status": "OK", "flags": old_flags},
        ],
    )

    minimap = agent.read_minimap(state)["minimap"]

    assert "RECENT_DECISIONS" not in minimap
    assert '"R1_EB_A_NB"' not in minimap
    assert '"R4_WB_A_SB"' not in minimap


def test_minimap_still_has_routes_and_nodes():
    telemetry = {
        "simulation_time_seconds": 42.0,
        "routes": {
            route_id: {
                "active": route_id == "R1_EB_A_NB",
                "tsp_enabled": False,
                "dbl_enabled": False,
            }
            for route_id in guard.ROUTE_ORDER
        },
        "network_throughput": {
            "passengers_per_minute": 80.0,
            "passengers_per_minute_recent": 95.0,
        },
        "network_summary": {"queues_passengers_est": {"EB": 20, "WB": 12}},
        "signal_state": {
            "nodes": {
                str(NODE_A): {"phase": "EW_GREEN", "signals": {"EB": "GREEN"}},
                "780": {"phase": "NS_GREEN", "signals": {"NB": "GREEN"}},
            }
        },
        "active_buses": [],
    }

    minimap = agent.read_minimap(agent_state(telemetry=telemetry))["minimap"]

    assert "simulation_time_seconds=42.0" in minimap
    assert "passengers_per_minute=80.0" in minimap
    assert "passengers_per_minute_recent=95.0" in minimap
    assert 'queues_passengers_est={"EB": 20, "WB": 12}' in minimap
    assert "NODE_SUMMARY:" in minimap
    assert f"NODE {NODE_A}: phase=EW_GREEN" in minimap
    assert "NODE 780: phase=NS_GREEN" in minimap
    assert "ROUTES:" in minimap
    for position, route_id in enumerate(guard.ROUTE_ORDER, start=1):
        assert f"{position}) {route_id}:" in minimap
    assert "approaching_buses=0 none approaching" in minimap


def test_minimap_node_summary_has_approach_queues_and_actionable_bus():
    telemetry = {
        "simulation_time_seconds": 20.0,
        "routes": {
            route_id: {
                "active": True,
                "tsp_enabled": False,
                "dbl_enabled": False,
            }
            for route_id in guard.ROUTE_ORDER
        },
        "network_summary": {"queues_passengers_est": {}},
        "signal_state": {
            "nodes": {
                str(NODE_A): {
                    "phase": "NS_GREEN",
                    "signals": {
                        "EB": "RED", "WB": "RED", "NB": "GREEN", "SB": "GREEN"
                    },
                    "queues_passengers_est": {
                        "EB": 48, "WB": 32, "NB": 12, "SB": 8
                    },
                },
                str(NODE_B): {
                    "phase": "EW_GREEN",
                    "signals": {
                        "EB": "GREEN", "WB": "GREEN", "NB": "RED", "SB": "RED"
                    },
                    "queues_passengers_est": {
                        "EB": 20, "WB": 60, "NB": 40, "SB": 16
                    },
                },
            }
        },
        "active_buses": [
            {
                "bus_id": "STRATEGY_BUS",
                "route_id": "R1_EB_A_NB",
                "direction": "EB",
                "route_leg": {
                    "node_x": NODE_A,
                    "approach": "EB",
                    "movement": "LEFT",
                },
                "leg_state": "APPROACHING",
                "distance_to_stop_bar_px": 400.0,
                "eta_to_stop_bar_sec_freeflow": 16.0,
                "passengers": 45,
                "priority_granted": False,
            }
        ],
    }

    minimap = agent.read_minimap(
        agent_state(telemetry=telemetry, decision_lag_sec=8.0)
    )["minimap"]
    node_300 = next(
        line for line in minimap.splitlines() if line.startswith(f"- NODE {NODE_A}:")
    )
    node_700 = next(
        line for line in minimap.splitlines() if line.startswith(f"- NODE {NODE_B}:")
    )

    assert "waiting_pax: EB=48 WB=32 NB=12 SB=8 (total=100)" in node_300
    assert "actionable_bus: EB(45pax/1bus)" in node_300
    assert "exclusive GREEN and sets all others RED" in node_300
    assert "waiting_pax: EB=20 WB=60 NB=40 SB=16 (total=136)" in node_700
    assert "actionable_bus: none" in node_700


def test_minimap_node_summary_shows_queue_length_and_downstream_space():
    telemetry = {
        "simulation_time_seconds": 20.0,
        "routes": {
            route_id: {
                "active": True,
                "tsp_enabled": False,
                "dbl_enabled": False,
            }
            for route_id in guard.ROUTE_ORDER
        },
        "network_summary": {"queues_passengers_est": {}},
        "signal_state": {
            "nodes": {
                str(NODE_A): {
                    "phase": "NS_GREEN",
                    "signals": {
                        "EB": "RED", "WB": "RED", "NB": "GREEN", "SB": "GREEN"
                    },
                    "queues_passengers_est": {
                        "EB": 48, "WB": 32, "NB": 12, "SB": 8
                    },
                    "queue_length_m": {
                        "EB": 18.8, "WB": 7.5, "NB": 0.0, "SB": 3.75
                    },
                    "downstream_space_m": {
                        "EB": 67.0, "WB": 58.5, "NB": 40.2, "SB": 3.0
                    },
                    "downstream_blocked": {
                        "EB": False, "WB": False, "NB": False, "SB": True
                    },
                },
                str(NODE_B): {
                    "phase": "EW_GREEN",
                    "signals": {
                        "EB": "GREEN", "WB": "GREEN", "NB": "RED", "SB": "RED"
                    },
                    "queues_passengers_est": {
                        "EB": 20, "WB": 60, "NB": 40, "SB": 16
                    },
                    "queue_length_m": {
                        "EB": 0.0, "WB": 45.0, "NB": 12.0, "SB": 0.0
                    },
                    "downstream_space_m": {
                        "EB": 58.5, "WB": 0.0, "NB": 58.5, "SB": 58.5
                    },
                    "downstream_blocked": {
                        "EB": False, "WB": True, "NB": False, "SB": False
                    },
                },
            }
        },
        "active_buses": [],
    }

    minimap = agent.read_minimap(
        agent_state(telemetry=telemetry, decision_lag_sec=8.0)
    )["minimap"]
    node_lines = {
        node_x: next(
            line for line in minimap.splitlines()
            if line.startswith(f"- NODE {node_x}:")
        )
        for node_x in (str(NODE_A), str(NODE_B))
    }

    for line in node_lines.values():
        assert "queue_len_m:" in line
        assert "downstream_free_m:" in line
        # Both new fields follow waiting_pax, so the model reads passengers,
        # then extent, then room to move, in that order.
        assert line.index("waiting_pax:") < line.index("queue_len_m:")
        assert line.index("queue_len_m:") < line.index("downstream_free_m:")

    assert "queue_len_m: EB=18.8 WB=7.5 NB=0.0 SB=3.8" in node_lines[str(NODE_A)]
    assert (
        "downstream_free_m: EB=67.0 WB=58.5 NB=40.2 SB=3.0(BLOCKED)"
        in node_lines[str(NODE_A)]
    )
    assert "queue_len_m: EB=0.0 WB=45.0 NB=12.0 SB=0.0" in node_lines[str(NODE_B)]
    assert (
        "downstream_free_m: EB=58.5 WB=0.0(BLOCKED) NB=58.5 SB=58.5"
        in node_lines[str(NODE_B)]
    )
    # Unblocked approaches never carry the marker.
    assert "EB=67.0(BLOCKED)" not in node_lines[str(NODE_A)]
    # The existing per-node content is untouched.
    assert "waiting_pax: EB=48 WB=32 NB=12 SB=8 (total=100)" in node_lines[str(NODE_A)]
    assert "actionable_bus: none" in node_lines[str(NODE_A)]

    # Telemetry written before these fields existed still renders a NODE
    # line, with the new values marked unknown rather than crashing.
    for node in telemetry["signal_state"]["nodes"].values():
        for key in ("queue_length_m", "downstream_space_m", "downstream_blocked"):
            del node[key]
    legacy = agent.read_minimap(
        agent_state(telemetry=telemetry, decision_lag_sec=8.0)
    )["minimap"]
    legacy_300 = next(
        line for line in legacy.splitlines() if line.startswith(f"- NODE {NODE_A}:")
    )
    assert "queue_len_m: EB=? WB=? NB=? SB=?" in legacy_300
    assert "downstream_free_m: EB=? WB=? NB=? SB=?" in legacy_300
    assert "BLOCKED" not in legacy_300


def test_prompt_teaches_queue_length_and_downstream_space_rules():
    prompt = " ".join(agent.SYSTEM_PROMPT.split())

    assert "queue_len_m is how far back each approach is backed up" in prompt
    assert "downstream_free_m is the room beyond the node" in prompt
    assert "giving that approach green will NOT help" in prompt
    assert "never grant priority into it" in prompt
    assert "grant nothing there and let the Webster timing work" in prompt
    # The rules name the minimap fields the model will actually see.
    assert "queue_len_m" in prompt
    assert "downstream_free_m" in prompt
    assert "BLOCKED" in prompt
    # The decision tests come before the output schema, so the model reads
    # what to weigh before it reads how to answer.
    assert prompt.index("THE TSP TEST") < prompt.index("NETWORK:")
    assert prompt.index("NETWORK:") < prompt.index(agent.OUTPUT_SCHEMA.split()[0])


def test_output_contract_unchanged_by_spatial_fields():
    # The prompt still asks for the same positional schema: six booleans per
    # array, no route keys, and the guard still accepts exactly that.
    schema = json.loads(agent.OUTPUT_SCHEMA)
    assert set(schema) == {"reason", "tsp", "dbl"}
    assert len(guard.ROUTE_ORDER) == 6
    assert schema["tsp"] == [False] * 6
    assert schema["dbl"] == [False] * 6
    assert agent.OUTPUT_SCHEMA in agent.SYSTEM_PROMPT

    output = positional_output(tsp_route="R1_EB_A_NB", dbl_route="R4_WB_A_SB")
    assert len(output["tsp"]) == 6 and len(output["dbl"]) == 6
    assert all(isinstance(flag, bool) for flag in output["tsp"] + output["dbl"])
    mapped = guard.validate_flags_positional(output)
    assert mapped == valid_flags(tsp_route="R1_EB_A_NB", dbl_route="R4_WB_A_SB")

    decision = guard.safe_decision(json.dumps(output), turn=1, model="test-model")
    assert decision["flags"] == mapped
    assert decision["status"] != "HELD_ALL_OFF"

    # Extra spatial keys in the model's output are not part of the contract:
    # the guard ignores them without weakening the flags.
    with_extras = dict(output, queue_len_m={"EB": 1.0}, downstream_free_m=5.0)
    assert guard.validate_flags_positional(with_extras) == mapped


def test_prompt_frames_objective_and_arrival_test():
    """The prompt states the DV (person-hours of delay for everyone) and a
    computable grant test on fields the ROUTES line carries, instead of a
    restraint slogan the model cannot check against anything."""
    prompt = " ".join(agent.SYSTEM_PROMPT.split())

    assert "already running Webster-optimal timing" in prompt
    assert "fewest person-hours of delay" in prompt
    assert "bus riders and cross-street drivers alike" in prompt
    # The TSP test names its three inputs and the arrival fact that decides
    # whether priority can help at all.
    assert "would_stop=true AND passengers > cross_pax AND actionable=true" in prompt
    assert "would_stop=false: the bus reaches the stop bar inside the residual green" in prompt
    assert "residual_green_s" in prompt
    assert "If cross_pax >= passengers, set tsp=false" in prompt
    assert "Grant at most ONE route per node per turn" in prompt
    assert "approaching_buses=0 must have tsp=false and dbl=false" in prompt
    # DBL has a positive trigger, not only vetoes.
    assert "grant dbl when dbl_lane_queue_ahead=0 AND dbl_lane_obstructed=false" in prompt
    # The positional schema is untouched by the reframing: same keys, same
    # array lengths, same types.
    schema = json.loads(agent.OUTPUT_SCHEMA)
    assert set(schema) == {"reason", "tsp", "dbl"}
    assert schema["tsp"] == [False] * len(guard.ROUTE_ORDER)
    assert schema["dbl"] == [False] * len(guard.ROUTE_ORDER)
    assert isinstance(schema["reason"], str)

    # Throughput-maximising language licenses over-granting, so it must not
    # survive anywhere the model reads -- including the interpolated schema.
    assert "maximize" not in agent.SYSTEM_PROMPT
    assert "maximize" not in agent.OUTPUT_SCHEMA
    assert agent.OUTPUT_SCHEMA in agent.SYSTEM_PROMPT

    # The latency, DBL-obstruction and per-node rules are unchanged.
    assert "eta_at_decision_land_sec" in prompt
    assert "dbl_lane_obstructed=true" in prompt
    assert "DECISION_LAG_SEC" in prompt


def test_minimap_route_line_carries_arrival_test_fields():
    telemetry = {
        "simulation_time_seconds": 10.0,
        "routes": {"R1_EB_A_NB": {"active": True, "tsp_enabled": True}},
        "signal_state": {"nodes": {str(NODE_A): {"phase": "EW_GREEN", "signals": {}}}},
        "active_buses": [{
            "bus_id": "B1", "route_id": "R1_EB_A_NB", "leg_state": "APPROACHING",
            "route_leg": {"node_x": NODE_A, "approach": "EB"},
            "distance_to_stop_bar_px": 120.0, "eta_to_stop_bar_sec_freeflow": 12.0,
            "passengers": 45, "signal_colour_ahead": "GREEN", "residual_green_sec": 4.5,
            "would_have_stopped": True, "cross_traffic_pax": 17,
        }],
    }
    minimap = agent.read_minimap(agent_state(telemetry=telemetry, decision_lag_sec=1.0))["minimap"]
    line = next(l for l in minimap.splitlines() if l.startswith("1) R1_EB_A_NB"))
    assert "signal_ahead=GREEN" in line
    assert "residual_green_s=4.5" in line
    assert "would_stop=True" in line
    assert "cross_pax=17" in line


def test_granted_route_stays_locked_through_clearing():
    clearing = {
        "bus_id": "CLEARING",
        "route_id": "R2_EB_B_NB",
        "route_leg": {"node_x": 780, "movement": "LEFT"},
        "leg_state": "IN_INTERSECTION",
        "distance_to_stop_bar_px": 0.0,
        "eta_to_stop_bar_sec_freeflow": 0.0,
        "priority_granted": False,
        "priority_clearing": True,
    }
    state = agent_state(telemetry={"active_buses": [clearing]})

    locked_update = agent.check_locked(state)

    assert locked_update["locked_routes"] == {"R2_EB_B_NB"}
    assert "LOCKED_ROUTES=R2_EB_B_NB" in locked_update["minimap"]


def test_clearing_bus_not_shown_as_approaching():
    route_id = "R2_EB_B_NB"
    clearing = {
        "bus_id": "CLEARING",
        "route_id": route_id,
        "route_leg": {"node_x": 780, "movement": "LEFT"},
        "leg_state": "IN_INTERSECTION",
        "distance_to_stop_bar_px": 0.0,
        "eta_to_stop_bar_sec_freeflow": 0.0,
        "passengers": 60,
        "priority_granted": False,
        "priority_clearing": True,
    }
    telemetry = {
        "routes": {route_id: {"active": True, "tsp_enabled": True}},
        "active_buses": [clearing],
    }
    state = agent_state(telemetry=telemetry)

    minimap_update = agent.read_minimap(state)
    locked_update = agent.check_locked({**state, **minimap_update})

    route_line = next(
        line
        for line in minimap_update["minimap"].splitlines()
        if line.startswith(
            f"{guard.ROUTE_ORDER.index(route_id) + 1}) {route_id}:"
        )
    )
    assert "approaching_buses=0 none approaching" in route_line
    assert "nearest_eta_sec" not in route_line
    assert locked_update["locked_routes"] == {route_id}


def test_approaching_grant_still_locks():
    bus = {
        "bus_id": "APPROACHING",
        "route_id": "R4_WB_A_SB",
        "route_leg": {"node_x": NODE_A, "movement": "LEFT"},
        "leg_state": "APPROACHING",
        "distance_to_stop_bar_px": 80.0,
        "eta_to_stop_bar_sec_freeflow": 1.5,
        "priority_granted": True,
        "priority_clearing": False,
    }

    locked_update = agent.check_locked(
        agent_state(telemetry={"active_buses": [bus]})
    )

    assert locked_update["locked_routes"] == {"R4_WB_A_SB"}


def test_agent_call_is_single_network_call_and_guarded(monkeypatch):
    flags = valid_flags(tsp_route="R2_EB_B_NB")
    reason = "Prioritize the approaching high-occupancy bus."
    calls = []

    def fake_chat(**kwargs):
        calls.append(kwargs)
        return {
            "message": {
                "content": json.dumps(
                    positional_output(
                        tsp_route="R2_EB_B_NB", reason=reason
                    )
                )
            }
        }

    monkeypatch.setattr(agent, "ollama", SimpleNamespace(chat=fake_chat))
    state = agent_state(minimap="whole network", turn=8)
    raw_update = agent.ai_turn(state)
    guarded = agent.anti_cheat({**state, **raw_update})

    assert len(calls) == 1
    assert calls[0]["model"] == "test-model"
    # num_thread leaves SIM_RESERVED_CORES physical cores to the simulator.
    assert calls[0]["options"] == {
        "temperature": 0.2, "num_thread": agent.OLLAMA_NUM_THREAD, "num_ctx": agent.OLLAMA_NUM_CTX,
    }
    assert calls[0]["format"] == agent.OLLAMA_OUTPUT_FORMAT
    assert calls[0]["messages"][1]["content"] == "whole network"
    assert guarded["decision"]["status"] == "OK"
    assert guarded["decision"]["flags"] == flags
    assert guarded["decision"]["reason"] == reason
    assert "one sentence under about 40 words" in agent.SYSTEM_PROMPT
    assert "Never write route IDs as JSON keys" in agent.SYSTEM_PROMPT
    assert agent.OUTPUT_SCHEMA.index('"reason"') < agent.OUTPUT_SCHEMA.index('"tsp"')


def test_ollama_timeout_holds_all_off(monkeypatch):
    started = threading.Event()
    release = threading.Event()
    finished = threading.Event()

    def slow_chat(**kwargs):
        started.set()
        release.wait(timeout=1.0)
        finished.set()
        return {
            "message": {
                "content": json.dumps(positional_output(reason="Too late."))
            }
        }

    monkeypatch.setattr(agent, "ollama", SimpleNamespace(chat=slow_chat))
    monkeypatch.setattr(agent, "_OLLAMA_CLIENT", None)
    monkeypatch.setattr(agent, "_OLLAMA_CALL_LOCK", threading.Lock())
    monkeypatch.setattr(agent, "OLLAMA_TIMEOUT_SECONDS", 0.02)

    state = agent_state()
    started_at = time.perf_counter()
    failed = agent.ai_turn(state)
    elapsed = time.perf_counter() - started_at

    assert started.is_set()
    assert elapsed < 0.15
    guarded = agent.anti_cheat({**state, **failed})

    assert failed["status"] == "INVALID"
    assert "TimeoutError" in failed["raw_output"]
    assert failed["call_metrics"]["latency_ms"] >= 10.0
    assert failed["call_metrics"]["input_tokens"] is None
    assert failed["call_metrics"]["output_tokens"] is None
    assert guarded["decision"]["status"] == "HELD_ALL_OFF"
    assert guarded["decision"]["flags"] == guard.all_off_flags()

    agent.log_turn({**state, **failed}, guarded["decision"])
    logged = json.loads(agent.TURN_LOG_PATH.read_text(encoding="utf-8"))
    assert logged["latency_ms"] is not None
    assert logged["input_tokens"] is None
    assert logged["output_tokens"] is None
    assert logged["tokens_per_sec"] is None

    release.set()
    assert finished.wait(timeout=0.2)


def test_ollama_normal_call_unaffected(monkeypatch):
    expected = json.dumps(
        positional_output(
            tsp_route="R1_EB_A_NB",
            reason="Prioritize the approaching bus.",
        )
    )
    calls = []

    class FakeClient:
        def __init__(self, **kwargs):
            calls.append(("init", kwargs))

        @staticmethod
        def chat(**kwargs):
            calls.append(("chat", kwargs))
            return {
                "message": {"content": expected},
                "prompt_eval_count": 120,
                "eval_count": 30,
                "eval_duration": 2_000_000_000,
                "total_duration": 3_000_000_000,
            }

    monkeypatch.setattr(agent, "ollama", SimpleNamespace(Client=FakeClient))
    monkeypatch.setattr(agent, "_OLLAMA_CLIENT", None)
    monkeypatch.setattr(agent, "_OLLAMA_CALL_LOCK", threading.Lock())

    raw, metrics = agent._call_ollama("test-model", "whole network")
    decision = guard.safe_decision(raw, turn=1, model="test-model")

    assert raw == expected
    assert calls[0] == ("init", {"timeout": agent.OLLAMA_TIMEOUT_SECONDS})
    assert calls[1][0] == "chat"
    assert calls[1][1]["format"] == agent.OLLAMA_OUTPUT_FORMAT
    assert metrics == {
        "input_tokens": 120,
        "output_tokens": 30,
        "eval_duration_ns": 2_000_000_000,
        "total_duration_ns": 3_000_000_000,
    }
    assert decision["status"] == "OK"
    assert decision["flags"]["R1_EB_A_NB"]["tsp"] is True

    agent.log_turn(
        agent_state(
            raw_output=raw,
            call_metrics={**metrics, "latency_ms": 3100.0},
        ),
        decision,
    )
    logged = json.loads(agent.TURN_LOG_PATH.read_text(encoding="utf-8"))
    assert logged["latency_ms"] == 3100.0
    assert logged["input_tokens"] == 120
    assert logged["output_tokens"] == 30
    assert logged["tokens_per_sec"] == 15.0


def test_ollama_format_fallback_still_works_under_timeout(
    monkeypatch,
):
    calls = []

    def fake_chat(**kwargs):
        calls.append(kwargs)
        if "format" in kwargs:
            raise TypeError("unexpected keyword argument 'format'")
        return {
            "message": {
                "content": json.dumps(
                    positional_output(reason="Fallback remained guarded.")
                )
            }
        }

    monkeypatch.setattr(agent, "ollama", SimpleNamespace(chat=fake_chat))
    monkeypatch.setattr(agent, "_OLLAMA_CLIENT", None)
    monkeypatch.setattr(agent, "_OLLAMA_CALL_LOCK", threading.Lock())

    raw = agent.ai_turn(agent_state(minimap="whole network"))
    guarded = agent.anti_cheat({**agent_state(), **raw})

    assert raw["status"] == "OK"
    assert guarded["decision"]["status"] == "OK"
    assert guarded["decision"]["flags"] == guard.all_off_flags()
    assert len(calls) == 2
    assert calls[0]["format"] == agent.OLLAMA_OUTPUT_FORMAT
    assert "format" not in calls[1]


def test_gemini_routing(monkeypatch):
    assert agent._is_gemini("gemini-2.5-flash") is True
    assert agent._is_gemini("llama3.1:8b") is False

    calls = []
    monkeypatch.setattr(
        agent,
        "_call_gemini",
        lambda model, system_prompt, minimap: (
            calls.append((model, system_prompt, minimap))
            or (
                json.dumps(positional_output(reason="Cloud route.")),
                {
                    "input_tokens": 42,
                    "output_tokens": 12,
                    "eval_duration_ns": None,
                    "total_duration_ns": None,
                },
            )
        ),
    )
    monkeypatch.setattr(
        agent,
        "ollama",
        SimpleNamespace(
            chat=lambda **kwargs: pytest.fail("Ollama must not handle Gemini")
        ),
    )

    result = agent.ai_turn(
        agent_state(model="gemini-2.5-flash", minimap="network snapshot")
    )

    assert result["status"] == "OK"
    assert calls == [
        ("gemini-2.5-flash", agent.SYSTEM_PROMPT, "network snapshot")
    ]


def test_gemini_call_uses_key_prompt_and_low_temperature(monkeypatch):
    calls = []

    class FakeConfig:
        def __init__(self, **kwargs):
            self.options = kwargs

    class FakeModels:
        @staticmethod
        def generate_content(**kwargs):
            calls.append(kwargs)
            return SimpleNamespace(
                text="  guarded JSON  ",
                usage_metadata=SimpleNamespace(
                    prompt_token_count=42,
                    candidates_token_count=12,
                ),
            )

    class FakeGenai:
        @staticmethod
        def Client(api_key):
            calls.append({"api_key": api_key})
            return SimpleNamespace(models=FakeModels())

    monkeypatch.setenv("GEMINI_API_KEY", "test-secret")
    monkeypatch.setattr(agent, "_genai", FakeGenai)
    monkeypatch.setattr(
        agent,
        "_genai_types",
        SimpleNamespace(GenerateContentConfig=FakeConfig),
    )
    monkeypatch.setattr(agent, "_GEMINI_CLIENT", None)
    monkeypatch.setattr(agent, "_GEMINI_CALL_LOCK", threading.Lock())

    text, metrics = agent._call_gemini(
        "gemini-2.5-flash", "system", "minimap"
    )

    assert text == "  guarded JSON  "
    assert calls[0] == {"api_key": "test-secret"}
    assert calls[1]["model"] == "gemini-2.5-flash"
    assert calls[1]["contents"] == "system\n\nminimap"
    assert calls[1]["config"].options == {"temperature": 0.2}
    assert metrics == {
        "input_tokens": 42,
        "output_tokens": 12,
        "eval_duration_ns": None,
        "total_duration_ns": None,
    }


def test_gemini_failure_holds_all_off(monkeypatch):
    monkeypatch.setattr(
        agent,
        "_call_gemini",
        lambda *args: (_ for _ in ()).throw(RuntimeError("cloud unavailable")),
    )
    state = agent_state(model="gemini-2.5-flash")

    failed = agent.ai_turn(state)
    guarded = agent.anti_cheat({**state, **failed})

    assert failed["status"] == "INVALID"
    assert guarded["decision"]["status"] == "HELD_ALL_OFF"
    assert guarded["decision"]["flags"] == guard.all_off_flags()


def test_gemini_timeout_holds_all_off(monkeypatch):
    class FakeConfig:
        def __init__(self, **kwargs):
            self.options = kwargs

    class SlowModels:
        @staticmethod
        def generate_content(**kwargs):
            time.sleep(0.2)
            return SimpleNamespace(text="late response")

    monkeypatch.setattr(
        agent,
        "_GEMINI_CLIENT",
        SimpleNamespace(models=SlowModels()),
    )
    monkeypatch.setattr(
        agent,
        "_genai_types",
        SimpleNamespace(GenerateContentConfig=FakeConfig),
    )
    monkeypatch.setattr(agent, "_GEMINI_CALL_LOCK", threading.Lock())
    monkeypatch.setattr(agent, "GEMINI_TIMEOUT_SECONDS", 0.02)
    state = agent_state(model="gemini-2.5-flash")

    started = time.perf_counter()
    failed = agent.ai_turn(state)
    elapsed = time.perf_counter() - started
    guarded = agent.anti_cheat({**state, **failed})

    assert elapsed < 0.15
    assert failed["status"] == "INVALID"
    assert "TimeoutError" in failed["raw_output"]
    assert guarded["decision"]["status"] == "HELD_ALL_OFF"
    assert guarded["decision"]["flags"] == guard.all_off_flags()


def test_openai_routing(monkeypatch):
    assert agent._is_openai("gpt-5") is True
    assert agent._is_openai("gpt-4.1") is True
    assert agent._is_openai("llama3.1:8b") is False
    assert agent._is_openai("gemini-2.5-flash") is False

    calls = []
    monkeypatch.setattr(
        agent,
        "_call_openai",
        lambda model, system_prompt, minimap: (
            calls.append((model, system_prompt, minimap))
            or (
                json.dumps(positional_output(reason="OpenAI route.")),
                {
                    "input_tokens": 42,
                    "output_tokens": 12,
                    "eval_duration_ns": None,
                    "total_duration_ns": None,
                },
            )
        ),
    )
    monkeypatch.setattr(
        agent,
        "ollama",
        SimpleNamespace(
            chat=lambda **kwargs: pytest.fail("Ollama must not handle OpenAI")
        ),
    )

    result = agent.ai_turn(agent_state(model="gpt-5", minimap="network snapshot"))

    assert result["status"] == "OK"
    assert calls == [("gpt-5", agent.SYSTEM_PROMPT, "network snapshot")]


def test_openai_call_uses_key_prompt_and_low_temperature(monkeypatch):
    calls = []

    class FakeMessage:
        content = "  guarded JSON  "

    class FakeChoice:
        message = FakeMessage()

    class FakeCompletions:
        @staticmethod
        def create(**kwargs):
            calls.append(kwargs)
            return SimpleNamespace(
                choices=[FakeChoice()],
                usage=SimpleNamespace(prompt_tokens=42, completion_tokens=12),
            )

    class FakeChat:
        completions = FakeCompletions()

    class FakeOpenAIClient:
        chat = FakeChat()

    class FakeOpenAI:
        @staticmethod
        def OpenAI(api_key):
            calls.append({"api_key": api_key})
            return FakeOpenAIClient()

    monkeypatch.setenv("OPENAI_API_KEY", "test-secret")
    monkeypatch.setattr(agent, "_openai", FakeOpenAI)
    monkeypatch.setattr(agent, "_OPENAI_CLIENT", None)
    monkeypatch.setattr(agent, "_OPENAI_CALL_LOCK", threading.Lock())

    text, metrics = agent._call_openai("gpt-5", "system", "minimap")

    assert text == "  guarded JSON  "
    assert calls[0] == {"api_key": "test-secret"}
    assert calls[1]["model"] == "gpt-5"
    assert calls[1]["messages"] == [
        {"role": "system", "content": "system"},
        {"role": "user", "content": "minimap"},
    ]
    assert calls[1]["temperature"] == 0.2
    assert metrics == {
        "input_tokens": 42,
        "output_tokens": 12,
        "cost_usd": None,
        "eval_duration_ns": None,
        "total_duration_ns": None,
    }


def test_grok_routing(monkeypatch):
    assert agent._is_grok("grok-4.6") is True
    assert agent._is_grok("grok-3-mini") is True
    assert agent._is_grok("gpt-5") is False
    assert agent._is_openai("grok-4.6") is False
    assert agent._is_gemini("grok-4.6") is False

    calls = []
    monkeypatch.setattr(
        agent,
        "_call_grok",
        lambda model, system_prompt, minimap: (
            calls.append((model, system_prompt, minimap))
            or (
                json.dumps(positional_output(reason="Grok route.")),
                {
                    "input_tokens": 42,
                    "output_tokens": 12,
                    "eval_duration_ns": None,
                    "total_duration_ns": None,
                },
            )
        ),
    )
    monkeypatch.setattr(
        agent, "_call_openai",
        lambda *a, **k: pytest.fail("OpenAI must not handle Grok"),
    )
    monkeypatch.setattr(
        agent,
        "ollama",
        SimpleNamespace(
            chat=lambda **kwargs: pytest.fail("Ollama must not handle Grok")
        ),
    )

    result = agent.ai_turn(agent_state(model="grok-4.6", minimap="network snapshot"))

    assert result["status"] == "OK"
    assert calls == [("grok-4.6", agent.SYSTEM_PROMPT, "network snapshot")]


def test_grok_call_uses_xai_endpoint_and_key(monkeypatch):
    """Grok goes through the openai SDK pointed at xAI's base_url, with the
    Grok key -- never the OpenAI key or the default OpenAI host."""
    calls = []

    class FakeMessage:
        content = "  guarded JSON  "

    class FakeChoice:
        message = FakeMessage()

    class FakeCompletions:
        @staticmethod
        def create(**kwargs):
            calls.append(kwargs)
            return SimpleNamespace(
                choices=[FakeChoice()],
                usage=SimpleNamespace(prompt_tokens=7, completion_tokens=3),
            )

    class FakeChat:
        completions = FakeCompletions()

    class FakeClient:
        chat = FakeChat()

    class FakeOpenAI:
        @staticmethod
        def OpenAI(api_key, base_url=None):
            calls.append({"api_key": api_key, "base_url": base_url})
            return FakeClient()

    monkeypatch.setenv("GROK_API_KEY", "grok-secret")
    monkeypatch.setenv("OPENAI_API_KEY", "openai-secret")
    monkeypatch.setattr(agent, "_openai", FakeOpenAI)
    monkeypatch.setattr(agent, "_GROK_CLIENT", None)
    monkeypatch.setattr(agent, "_GROK_CALL_LOCK", threading.Lock())

    text, metrics = agent._call_grok("grok-4.6", "system", "minimap")

    assert text == "  guarded JSON  "
    assert calls[0] == {"api_key": "grok-secret", "base_url": agent.GROK_BASE_URL}
    assert agent.GROK_BASE_URL == "https://api.x.ai/v1"
    assert calls[1]["model"] == "grok-4.6"
    assert calls[1]["temperature"] == 0.2
    assert metrics["input_tokens"] == 7 and metrics["output_tokens"] == 3


def test_grok_missing_key_holds_all_off(monkeypatch):
    monkeypatch.delenv("GROK_API_KEY", raising=False)
    monkeypatch.setattr(agent, "_openai", SimpleNamespace(OpenAI=lambda **k: None))
    monkeypatch.setattr(agent, "_GROK_CLIENT", None)
    state = agent_state(model="grok-4.6")

    failed = agent.ai_turn(state)
    guarded = agent.anti_cheat({**state, **failed})

    assert failed["status"] == "INVALID"
    assert "GROK_API_KEY" in failed["raw_output"]
    assert guarded["decision"]["status"] == "HELD_ALL_OFF"


def test_openai_failure_holds_all_off(monkeypatch):
    monkeypatch.setattr(
        agent,
        "_call_openai",
        lambda *args: (_ for _ in ()).throw(RuntimeError("cloud unavailable")),
    )
    state = agent_state(model="gpt-5")

    failed = agent.ai_turn(state)
    guarded = agent.anti_cheat({**state, **failed})

    assert failed["status"] == "INVALID"
    assert guarded["decision"]["status"] == "HELD_ALL_OFF"
    assert guarded["decision"]["flags"] == guard.all_off_flags()


def test_openai_timeout_holds_all_off(monkeypatch):
    class SlowCompletions:
        @staticmethod
        def create(**kwargs):
            time.sleep(0.2)
            return SimpleNamespace(choices=[])

    class SlowChat:
        completions = SlowCompletions()

    monkeypatch.setattr(
        agent,
        "_OPENAI_CLIENT",
        SimpleNamespace(chat=SlowChat()),
    )
    monkeypatch.setattr(agent, "_OPENAI_CALL_LOCK", threading.Lock())
    monkeypatch.setattr(agent, "OPENAI_TIMEOUT_SECONDS", 0.02)
    state = agent_state(model="gpt-5")

    started = time.perf_counter()
    failed = agent.ai_turn(state)
    elapsed = time.perf_counter() - started
    guarded = agent.anti_cheat({**state, **failed})

    assert elapsed < 0.15
    assert failed["status"] == "INVALID"
    assert "TimeoutError" in failed["raw_output"]
    assert guarded["decision"]["status"] == "HELD_ALL_OFF"
    assert guarded["decision"]["flags"] == guard.all_off_flags()


def test_openai_missing_key_holds_all_off(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.setattr(agent, "_openai", SimpleNamespace(OpenAI=lambda **k: None))
    monkeypatch.setattr(agent, "_OPENAI_CLIENT", None)
    state = agent_state(model="gpt-5")

    failed = agent.ai_turn(state)
    guarded = agent.anti_cheat({**state, **failed})

    assert failed["status"] == "INVALID"
    assert "OPENAI_API_KEY" in failed["raw_output"]
    assert guarded["decision"]["status"] == "HELD_ALL_OFF"


def test_langgraph_runs_one_complete_network_turn(tmp_path, monkeypatch):
    flags = valid_flags(dbl_route="R5_WB_B_SB")
    reason = "Open the dynamic lane for the nearest loaded bus."
    telemetry = {
        "timestamp": agent.time.time(),
        "simulation_time_seconds": 12.0,
        "routes": {
            route_id: {
                "active": True,
                "tsp_enabled": False,
                "dbl_enabled": False,
            }
            for route_id in guard.VALID_ROUTES
        },
        "network_throughput": {
            "passengers_per_minute": 20.0,
            "passengers_per_minute_recent": 30.0,
        },
        "network_summary": {"queues_passengers_est": {}},
        "signal_state": {"nodes": {}},
        "active_buses": [],
    }
    monkeypatch.setattr(agent, "_read_telemetry", lambda: telemetry)
    monkeypatch.setattr(agent, "DECISION_PATH", tmp_path / "decision.json")
    monkeypatch.setattr(
        agent,
        "ollama",
        SimpleNamespace(
            chat=lambda **kwargs: {
                "message": {
                    "content": "<think>compare passenger queues</think>\n"
                    + json.dumps(
                        positional_output(
                            dbl_route="R5_WB_B_SB", reason=reason
                        )
                    )
                }
            }
        ),
    )

    result = agent.build_graph().invoke(
        agent_state(turn=11, run_uuid="RUN-7", telemetry_frame=4242)
    )

    assert result["decision"]["status"] == "OK"
    # Through the compiled graph, not a plain dict: a state key LangGraph
    # does not declare is dropped, and the published decision would carry
    # run_uuid "" -- refused by main.merge_ai_decision as FOREIGN_DECISION.
    published = json.loads(agent.DECISION_PATH.read_text(encoding="utf-8"))
    assert published["run_uuid"] == "RUN-7" and published["telemetry_frame"] == 4242
    assert result["decision"]["flags"] == flags
    assert result["decision"]["reason"] == reason
    assert result["recent_decisions"][-1]["turn"] == 11
    assert json.loads(agent.DECISION_PATH.read_text(encoding="utf-8"))["flags"] == flags
    turn_records = [
        json.loads(line)
        for line in agent.TURN_LOG_PATH.read_text(encoding="utf-8").splitlines()
    ]
    assert len(turn_records) == 1
    assert turn_records[0]["turn"] == 11
    assert turn_records[0]["status"] == "OK"
    assert "<think>compare passenger queues</think>" in turn_records[0]["raw_output"]
    assert turn_records[0]["flags"] == flags
    assert turn_records[0]["reason"] == reason
    assert turn_records[0]["pax_per_min_recent"] == 30.0
    assert "ROUTES:" in turn_records[0]["minimap"]
    assert turn_records[0]["stale"] is False


def test_agent_failure_and_stale_telemetry_write_all_off(tmp_path, monkeypatch):
    reject_log = tmp_path / "rejects.log"
    decision_path = tmp_path / "decision.json"
    monkeypatch.setattr(guard, "REJECT_LOG_PATH", reject_log)
    monkeypatch.setattr(agent, "DECISION_PATH", decision_path)
    monkeypatch.setattr(
        agent,
        "ollama",
        SimpleNamespace(chat=lambda **kwargs: (_ for _ in ()).throw(RuntimeError("down"))),
    )

    state = agent_state(turn=9)
    failed = agent.ai_turn(state)
    guarded = agent.anti_cheat({**state, **failed})
    assert guarded["decision"]["status"] == "HELD_ALL_OFF"
    assert guarded["decision"]["flags"] == guard.all_off_flags()
    assert guarded["decision"]["reason"] == ""

    stale = agent.hold(agent_state(turn=10, status="STALE"))
    written = json.loads(decision_path.read_text(encoding="utf-8"))
    assert stale["status"] == "STALE"
    assert written["status"] == "HELD_ALL_OFF"
    assert written["flags"] == guard.all_off_flags()
    assert written["reason"] == ""
    logged = json.loads(agent.TURN_LOG_PATH.read_text(encoding="utf-8"))
    assert logged["turn"] == 10
    assert logged["status"] == "HELD_ALL_OFF"
    assert logged["reason"] == ""
    assert logged["stale"] is True


def test_turn_logging_failure_never_breaks_decision(monkeypatch, tmp_path):
    monkeypatch.setattr(agent, "TURN_LOG_PATH", tmp_path / "missing" / "turns.jsonl")
    decision = guard.safe_decision(
        json.dumps(positional_output()), turn=12, model="test-model"
    )

    agent.log_turn(
        agent_state(
            minimap="observed network",
            raw_output="<think>full trace</think>",
            locked_routes={"R1_EB_A_NB"},
        ),
        decision,
    )

    assert decision["status"] == "OK"
    assert not agent.TURN_LOG_PATH.exists()


def test_read_ai_control_is_fail_closed_and_clamps_tick(tmp_path):
    missing = agent.read_ai_control(tmp_path / "missing.json")
    assert missing == agent.DEFAULT_CONTROL

    control_path = tmp_path / "ai_control.json"
    control_path.write_text(
        json.dumps({"armed": True, "model": "m", "tick_seconds": 999}),
        encoding="utf-8",
    )
    assert agent.read_ai_control(control_path) == {
        "armed": True,
        "model": "m",
        "tick_seconds": control_panel.TICK_SECONDS_MAX,
        "simulation_running": False,
        "control_mode": "assisted",
    }
    # One decision per ~50 s signal cycle must be representable: a tick
    # below the slowest arm's latency measures a disconnected loop.
    assert control_panel.TICK_SECONDS_MAX >= 60
    control_path.write_text(
        json.dumps({"armed": True, "model": "m", "tick_seconds": 60}),
        encoding="utf-8",
    )
    assert agent.read_ai_control(control_path)["tick_seconds"] == 60
    assert control_panel._effective_tick_seconds() <= control_panel.TICK_SECONDS_MAX


def test_agent_turn_resets_on_frame_drop():
    previous_decisions = [{"turn": 66}, {"turn": 67}]

    turn, recent, last_frame, reset_detected = agent._track_run_boundary(
        67,
        previous_decisions,
        11_300,
        {"frame_number": 0, "simulation_running": True},
    )
    turn += 1

    assert reset_detected is True
    assert turn == 1
    assert recent == []
    assert last_frame == 0


def test_paused_or_old_telemetry_routes_to_stale_hold(monkeypatch):
    now = agent.time.time()
    monkeypatch.setattr(
        agent,
        "_read_telemetry",
        lambda: {"timestamp": now, "simulation_paused": True},
    )
    paused = agent.load_save(agent_state())
    assert paused["status"] == "STALE"
    assert agent._route_after_load({**agent_state(), **paused}) == "hold"

    monkeypatch.setattr(
        agent,
        "_read_telemetry",
        lambda: {"timestamp": now - agent.STALE_SECONDS - 0.1},
    )
    assert agent.load_save(agent_state())["status"] == "STALE"


def telemetry_payload(frame=60):
    return {
        "frame_number": frame,
        "simulation_time_seconds": frame / 60,
        "network_throughput": {
            "passengers_served_total": 120,
            "passengers_served_bus": 90,
            "passengers_served_car": 30,
            "buses_served": 2,
            "cars_served": 8,
            "passengers_per_minute": 120.0,
            "passengers_per_minute_recent": 135.0,
        },
        "network_summary": {
            "total_vehicles": 24,
            "queues": {"EB": 3, "WB": 4},
            "queues_passengers_est": {"EB": 12, "WB": 16},
        },
    }


def test_telemetry_log_uses_schema_keys_interval_and_ai_state(
    tmp_path, monkeypatch
):
    destination = tmp_path / "telemetry_log.jsonl"
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", destination)
    monkeypatch.setattr(main, "_last_telemetry_log_frame", None)
    monkeypatch.setitem(
        control_panel.global_config,
        "ai_runtime",
        {"armed": True, "last_status": "OK"},
    )

    assert not main.log_telemetry_sample(59, telemetry_payload(59))
    assert main.log_telemetry_sample(60, telemetry_payload(60))
    assert not main.log_telemetry_sample(60, telemetry_payload(60))
    assert not main.log_telemetry_sample(119, telemetry_payload(119))
    assert main.log_telemetry_sample(121, telemetry_payload(121))

    records = [
        json.loads(line)
        for line in destination.read_text(encoding="utf-8").splitlines()
    ]
    assert [record["frame"] for record in records] == [60, 121]
    assert records[0]["vehicles_in_network"] == 24
    assert records[0]["ai_armed"] is True
    assert records[0]["ai_last_status"] == "OK"
    assert records[0]["queues_vehicles"] == {"EB": 3, "WB": 4}
    assert records[0]["pax_per_min_recent"] == 135.0


def test_telemetry_logging_failure_is_nonfatal(tmp_path, monkeypatch):
    monkeypatch.setattr(
        main,
        "TELEMETRY_LOG_PATH",
        tmp_path / "missing" / "telemetry_log.jsonl",
    )
    monkeypatch.setattr(main, "_last_telemetry_log_frame", None)

    assert not main.log_telemetry_sample(60, telemetry_payload(60))
    assert main._last_telemetry_log_frame is None


def test_reset_clears_session_logs(tmp_path, monkeypatch):
    telemetry_log = tmp_path / "telemetry_log.jsonl"
    turn_log = tmp_path / "agent_turn_log.jsonl"
    unrelated = tmp_path / "keep.txt"
    telemetry_log.write_text("old telemetry\n", encoding="utf-8")
    turn_log.write_text("old turns\n", encoding="utf-8")
    unrelated.write_text("keep", encoding="utf-8")
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", telemetry_log)
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", turn_log)
    monkeypatch.setattr(main, "_last_telemetry_log_frame", 600)

    main.reset_session_logs()

    assert not telemetry_log.exists()
    assert not turn_log.exists()
    assert unrelated.read_text(encoding="utf-8") == "keep"
    assert main._last_telemetry_log_frame is None

    source = Path(main.__file__).read_text(encoding="utf-8")
    reset_helper = source.split("def perform_full_reset", 1)[1].split(
        "def begin_post_discharge_metering", 1
    )[0]
    assert "reset_session_logs()" in reset_helper
    main_body = source.split("def main():", 1)[1]
    assert main_body.count("perform_full_reset(vehicles, signals, telemetry)") == 2


def test_reset_isolates_runs(tmp_path, monkeypatch):
    telemetry_log = tmp_path / "telemetry_log.jsonl"
    turn_log = tmp_path / "agent_turn_log.jsonl"
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", telemetry_log)
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", turn_log)

    telemetry_log.write_text(
        json.dumps({"run": "A", "frame": 60}) + "\n",
        encoding="utf-8",
    )
    turn_log.write_text(
        json.dumps({"run": "A", "turn": 1}) + "\n",
        encoding="utf-8",
    )

    main.reset_session_logs()

    telemetry_log.write_text(
        json.dumps({"run": "B", "frame": 60}) + "\n",
        encoding="utf-8",
    )
    turn_log.write_text(
        json.dumps({"run": "B", "turn": 1}) + "\n",
        encoding="utf-8",
    )

    telemetry_rows = [
        json.loads(line)
        for line in telemetry_log.read_text(encoding="utf-8").splitlines()
    ]
    turn_rows = [
        json.loads(line)
        for line in turn_log.read_text(encoding="utf-8").splitlines()
    ]
    assert telemetry_rows == [{"run": "B", "frame": 60}]
    assert turn_rows == [{"run": "B", "turn": 1}]


def test_excel_export_builds_decisions_and_telemetry_sheets(
    tmp_path, monkeypatch
):
    from openpyxl import load_workbook

    telemetry_log = tmp_path / "telemetry_log.jsonl"
    turn_log = tmp_path / "agent_turn_log.jsonl"
    output = tmp_path / "session_export.xlsx"
    flags = valid_flags(
        tsp_route="R2_EB_B_NB", dbl_route="R5_WB_B_SB"
    )
    decision = {
        "turn": 4,
        "timestamp": 123.5,
        "model": "reasoning-model",
        "status": "OK",
        "reason": "Prioritize the highest passenger-pressure movement.",
        "pax_per_min_recent": 135.0,
        "flags": flags,
        "locked_routes": ["R1_EB_A_NB"],
        "minimap": "whole network",
        "raw_output": "<think>compare queues</think>",
    }
    turn_log.write_text(
        json.dumps(decision) + "\n{partial crash tail",
        encoding="utf-8",
    )
    sample = {
        "frame": 60,
        "sim_time_s": 1.0,
        "passengers_served_total": 120,
        "passengers_served_bus": 90,
        "passengers_served_car": 30,
        "buses_served": 2,
        "cars_served": 8,
        "pax_per_min_cumulative": 120.0,
        "pax_per_min_recent": 135.0,
        "vehicles_in_network": 24,
        "ai_armed": True,
        "ai_last_status": "OK",
        "queues_vehicles": {"EB": 3},
        "queues_passengers_est": {"EB": 12},
    }
    telemetry_log.write_text(json.dumps(sample) + "\n", encoding="utf-8")
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", telemetry_log)
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", turn_log)

    assert main.export_session_excel(output) == output

    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook.sheetnames == [
            "Decisions",
            "Telemetry",
            "AI Decision Audit",
            "Control Panel Inputs",
            "Control Panel Inputs (start)",
        ]
        inputs = workbook["Control Panel Inputs"]
        assert [cell.value for cell in inputs[1]] == [
            "section",
            "parameter",
            "value",
        ]
        assert inputs.max_row > 1
        decisions = workbook["Decisions"]
        decision_headers = [cell.value for cell in decisions[1]]
        decision_row = {
            header: decisions.cell(2, index + 1).value
            for index, header in enumerate(decision_headers)
        }
        assert decisions.max_row == 2
        assert decision_row["turn"] == 4
        assert decision_headers[4:6] == ["reason", "pax_per_min_at_turn"]
        assert decision_row["reason"] == (
            "Prioritize the highest passenger-pressure movement."
        )
        assert decision_row["pax_per_min_at_turn"] == 135.0
        assert decision_row["R2_EB_B_NB_tsp"] is True
        assert decision_row["R5_WB_B_SB_dbl"] is True
        assert decision_row["locked_routes"] == "R1_EB_A_NB"
        assert "<think>compare queues</think>" in decision_row["raw_output"]

        telemetry = workbook["Telemetry"]
        telemetry_headers = [cell.value for cell in telemetry[1]]
        telemetry_row = {
            header: telemetry.cell(2, index + 1).value
            for index, header in enumerate(telemetry_headers)
        }
        assert telemetry.max_row == 2
        assert telemetry_row["vehicles_in_network"] == 24
        assert telemetry_row["ai_armed"] is True
        assert json.loads(telemetry_row["queues_vehicles"]) == {"EB": 3}

        audit = workbook["AI Decision Audit"]
        audit_headers = [cell.value for cell in audit[1]]
        audit_row = {
            header: audit.cell(2, index + 1).value
            for index, header in enumerate(audit_headers)
        }
        assert audit_headers == main.AI_DECISION_AUDIT_HEADERS
        assert audit_row["model_reason"] == decision["reason"]
        assert audit_row["requested_tsp_routes"] == "R2_EB_B_NB"
        assert audit_row["requested_dbl_routes"] == "R5_WB_B_SB"
        assert audit_row["observation_minimap"] == "whole network"
    finally:
        workbook.close()


def test_default_excel_export_uses_dedicated_folder(tmp_path, monkeypatch):
    export_folder = tmp_path / "excel_exports"
    telemetry_log = tmp_path / "telemetry.jsonl"
    telemetry_log.write_text("{}\n", encoding="utf-8")
    monkeypatch.setattr(main, "EXCEL_EXPORT_DIR", export_folder)
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", telemetry_log)
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", tmp_path / "missing-turns.jsonl")
    monkeypatch.setitem(
        control_panel.global_config,
        "ai_runtime",
        {"model": "gemini2.5"},
    )
    monkeypatch.setitem(control_panel.global_config, "sim_time_seconds", 300)
    monkeypatch.setitem(control_panel.global_config, "random_seed", 42)

    destination = main.export_session_excel()

    assert destination is not None
    assert destination.parent == export_folder
    # An LLM arm carries its control mode; ai_runtime has none set, so the
    # session workbook falls back to assisted.
    assert re.fullmatch(
        r"gemini2\.5-assisted_5min_42seed_\d{8}_\d{6}\.xlsx",
        destination.name,
    )
    assert destination.suffix == ".xlsx"
    assert destination.exists()


def _decision_file(tmp_path, **fields):
    body = {"run_uuid": RUN_ID, "model": "new-model", "turn": 1,
            "timestamp": main.time.time(), "status": "OK",
            "flags": valid_flags(tsp_route="R1_EB_A_NB")}
    body.update(fields)
    path = tmp_path / "decision.json"
    path.write_text(json.dumps(body), encoding="utf-8")
    return path


def _all_off():
    return all(
        route[flag] is False
        for route in control_panel.bus_routes_config.values()
        for flag in ("tsp_enabled", "dbl_enabled")
    )


def test_decision_from_another_model_run_or_older_turn_is_refused(tmp_path, monkeypatch):
    runtime = {"model": "new-model", "tick_seconds": 5, "last_status": "INACTIVE", "last_turn": 0}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    tsp = control_panel.bus_routes_config["R1_EB_A_NB"]

    # The audit's probe: a fresh decision tagged with the previous model.
    assert not main.merge_ai_decision(_decision_file(tmp_path, model="old-model"))
    assert runtime["last_status"] == "FOREIGN_DECISION" and _all_off()
    # ...or made for the previous run.
    assert not main.merge_ai_decision(_decision_file(tmp_path, run_uuid="previous-run"))
    assert runtime["last_status"] == "FOREIGN_DECISION" and _all_off()

    assert main.merge_ai_decision(_decision_file(tmp_path, turn=5))
    assert tsp["tsp_enabled"] is True and runtime["last_turn"] == 5
    # Re-reading the same turn keeps it; a replayed older turn is refused.
    assert main.merge_ai_decision(_decision_file(tmp_path, turn=5))
    assert not main.merge_ai_decision(_decision_file(tmp_path, turn=4))
    assert runtime["last_status"] == "STALE_DECISION" and _all_off()


def test_full_reset_clears_flags_for_every_arm_and_restarts_turns(monkeypatch):
    monkeypatch.setitem(control_panel.global_config, "test_running", True)
    monkeypatch.setitem(control_panel.global_config, "test_model", "new-model")
    runtime = {"model": "new-model", "armed": True, "last_turn": 9}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    route = control_panel.bus_routes_config["R4_WB_A_SB"]
    monkeypatch.setitem(route, "tsp_enabled", True)
    monkeypatch.setitem(route, "dbl_enabled", True)

    main.perform_full_reset([], main.SignalController({"green_time": 240}))

    assert _all_off() and runtime["last_turn"] == 0
    assert control_panel.global_config["_run_uuid"] != RUN_ID  # a new identity


def test_second_simulator_instance_is_refused(tmp_path):
    lock = tmp_path / "simulator.lock"
    first = main.acquire_instance_lock(lock)
    with pytest.raises(SystemExit):
        main.acquire_instance_lock(lock)
    first.close()
    main.acquire_instance_lock(lock).close()


def test_armed_request_locks_the_route_before_any_grant():
    """A route is locked as soon as the controller holds a request for its
    bus, not only once green is adjusting: dropping the flag mid-request
    cancels it (FEATURE_DISABLED) and suppresses that bus's leg."""
    armed = {
        "bus_id": "ARMED",
        "route_id": "R1_EB_A_NB",
        "route_leg": {"node_x": 780, "movement": "STRAIGHT"},
        "leg_state": "APPROACHING",
        "distance_to_stop_bar_px": 120.0,
        "eta_to_stop_bar_sec_freeflow": 4.0,
        "priority_requested": True,
        "priority_granted": False,
        "priority_clearing": False,
    }
    state = agent_state(telemetry={"active_buses": [armed]})

    assert agent.check_locked(state)["locked_routes"] == {"R1_EB_A_NB"}


def test_live_request_keeps_its_route_flag_through_stale_and_held_all_off(
    tmp_path, monkeypatch
):
    """The lock is enforced at the merge boundary from the live controller:
    a route whose bus the controller already holds a request for keeps its
    flags whatever the decision says (stale, guard-held, or a decider whose
    snapshot predates the request), so a committed treatment is never
    cancelled as FEATURE_DISABLED by the decision path. Routes without a
    live request still follow the decision."""
    now = 1_900_000_000.0
    monkeypatch.setattr(main.time, "time", lambda: now)
    runtime = {"tick_seconds": 5, "last_status": "INACTIVE", "last_turn": 0}
    monkeypatch.setitem(control_panel.global_config, "ai_runtime", runtime)
    monkeypatch.setitem(control_panel.global_config, "_run_uuid", RUN_ID)
    for route in control_panel.bus_routes_config.values():
        route["tsp_enabled"] = route["dbl_enabled"] = False
    control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] = True
    control_panel.bus_routes_config["R2_EB_B_NB"]["tsp_enabled"] = True
    signals = SignalController({"green_time": 300}, yellow_time=2, red_clearance_time=2)
    bus = make_bus_for_leg("R1_EB_A_NB", NODE_A, "LOCKED_BUS")
    signals.update([bus])
    assert signals.routes_with_live_requests() == {"R1_EB_A_NB"}

    decision_path = tmp_path / "decision.json"
    stale = {"run_uuid": RUN_ID, "timestamp": now - 3600, "status": "OK",
             "flags": valid_flags()}
    decision_path.write_text(json.dumps(stale), encoding="utf-8")
    assert not main.merge_ai_decision(decision_path, signals=signals)
    assert runtime["last_status"] == "STALE_DECISION"
    assert control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] is True
    assert control_panel.bus_routes_config["R2_EB_B_NB"]["tsp_enabled"] is False

    # A fresh, valid all-off decision (what a guard HELD_ALL_OFF or a
    # snapshot-lagged decider writes) cannot withdraw it either.
    fresh = {"run_uuid": RUN_ID, "timestamp": now, "turn": 1,
             "status": "HELD_ALL_OFF", "flags": valid_flags()}
    decision_path.write_text(json.dumps(fresh), encoding="utf-8")
    assert main.merge_ai_decision(decision_path, signals=signals)
    assert control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] is True

    # Once the request is gone the same decision is applied in full.
    signals.update([])
    assert signals.routes_with_live_requests() == set()
    fresh["turn"] = 2
    decision_path.write_text(json.dumps(fresh), encoding="utf-8")
    assert main.merge_ai_decision(decision_path, signals=signals)
    assert control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] is False


def test_the_tick_floor_matches_the_agent_call_timeouts():
    """control_panel.MIN_UNSKIPPED_TICK_SECONDS exists so that no arm ever
    skips a decision grid point: the agent skips-and-counts any point that
    comes due while a call is still running, so the tick has to outlast the
    longest a call can possibly take. If a provider timeout is raised, this
    fails until the floor -- and the default tick above it -- follow.
    """
    longest_call = max(
        agent.OLLAMA_TIMEOUT_SECONDS, agent.GEMINI_TIMEOUT_SECONDS,
        agent.OPENAI_TIMEOUT_SECONDS, agent.GROK_TIMEOUT_SECONDS,
    )
    assert control_panel.AGENT_CALL_TIMEOUT_CEILING_SEC == longest_call
    assert control_panel.MIN_UNSKIPPED_TICK_SECONDS > longest_call
    assert control_panel.DEFAULT_TICK_SECONDS >= control_panel.MIN_UNSKIPPED_TICK_SECONDS
    assert control_panel.DEFAULT_TICK_SECONDS <= control_panel.TICK_SECONDS_MAX
    # Both decision paths ship at the floor, not just the Single Run card.
    assert control_panel.DEFAULT_BATCH_RUNTIME["tick_seconds"] == control_panel.DEFAULT_TICK_SECONDS


def test_ollama_leaves_physical_cores_to_the_simulator(monkeypatch):
    monkeypatch.delenv("OLLAMA_NUM_THREAD", raising=False)
    monkeypatch.setattr(agent.os, "cpu_count", lambda: 16)     # 8 physical, 2-way SMT
    assert agent.ollama_num_thread() == 8 - agent.SIM_RESERVED_CORES
    monkeypatch.setattr(agent.os, "cpu_count", lambda: 2)
    assert agent.ollama_num_thread() == 1                        # never zero
    monkeypatch.setenv("OLLAMA_NUM_THREAD", "3")
    assert agent.ollama_num_thread() == 3                        # hybrid-CPU override


def test_a_prompt_that_fills_the_context_window_fails_closed(monkeypatch):
    """Ollama truncates an over-long prompt silently and the model still
    answers valid JSON; the call must not be trusted."""
    def fake_chat(**kwargs):
        return {
            "message": {"content": json.dumps(positional_output(tsp_route="R2_EB_B_NB"))},
            "prompt_eval_count": agent.OLLAMA_NUM_CTX - 10,
        }

    monkeypatch.setattr(agent, "ollama", SimpleNamespace(chat=fake_chat))
    monkeypatch.setattr(agent, "_OLLAMA_CLIENT", None)
    state = agent_state(minimap="whole network", turn=9)
    update = agent.ai_turn(state)
    guarded = agent.anti_cheat({**state, **update})
    assert guarded["decision"]["status"] == "HELD_ALL_OFF"
    assert not any(route["tsp"] for route in guarded["decision"]["flags"].values())


def test_the_truncation_guard_uses_the_models_own_smaller_window(monkeypatch):
    """A model whose native context is below OLLAMA_NUM_CTX is capped there by
    Ollama, so a prompt that fits 8k can still be truncated."""
    class Client:
        def show(self, model):
            return SimpleNamespace(modelinfo={"llama.context_length": 4096})
        def chat(self, **kwargs):
            return {"message": {"content": json.dumps(positional_output(tsp_route="R2_EB_B_NB"))},
                    "prompt_eval_count": 4000}
    monkeypatch.setattr(agent, "_OLLAMA_CLIENT", Client())
    monkeypatch.setattr(agent, "_MODEL_CONTEXT_TOKENS", {})
    assert agent.model_context_tokens(agent._OLLAMA_CLIENT, "small-ctx") == 4096
    state = agent_state(minimap="whole network", turn=10, model="small-ctx")
    guarded = agent.anti_cheat({**state, **agent.ai_turn(state)})
    assert guarded["decision"]["status"] == "HELD_ALL_OFF"
