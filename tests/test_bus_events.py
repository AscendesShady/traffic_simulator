"""Per-bus lifecycle log: the causal evidence for treated-vs-untreated delay.

Each bus writes exactly one row to bus_events.jsonl when it leaves the
network; the row carries spawn/completion frames, per-node TSP treatment and
the frames it spent waiting. These tests drive real buses through the real
controller and geometry so the logged values come from actual simulation
state, not from stubs.
"""
import json

import pytest

import src.ui.control_panel as control_panel
import src.core.main as main
from src.telemetry.bus_event_log import BUS_EVENT_HEADERS, BusEventTracker, flatten_bus_event
from src.ui.canvas_gemini import H_Y, HEIGHT, INT_X, LANE, ROAD_W, STOP, WIDTH
from src.core.signal_controller import TSP_ACTION_EXTENDING, TSP_ACTION_NONE, SignalController
from tests.helpers import make_bus_for_leg


GREEN_FRAMES = 100


def make_controller():
    return SignalController(
        {"green_time": GREEN_FRAMES},
        yellow_time=2,
        red_clearance_time=2,
        min_green_frames=30,
    )


def read_rows(path):
    if not path.exists():
        return []
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def run_bus_to_completion(bus, controller, tracker, max_frames=3000):
    """Drive one bus through the controller until it exits the network.

    Mirrors the main loop's order: controller update, vehicle update, exit
    check (completion), then the per-frame instrumentation sample.
    """
    vehicles = [bus]
    frame = 0
    for _ in range(max_frames):
        frame += 1
        controller.update(vehicles)
        signals = controller.get_all_signals(INT_X)
        for vehicle in list(vehicles):
            vehicle.update(
                signals, INT_X, H_Y, ROAD_W, STOP, LANE, vehicles, controller
            )
            off_canvas = (
                vehicle.x < -60
                or vehicle.x > WIDTH + 60
                or vehicle.y < -60
                or vehicle.y > HEIGHT + 60
            )
            if off_canvas and vehicle.passed_nodes:
                record = tracker.complete(vehicle, frame, controller)
                vehicles.remove(vehicle)
                return record, frame
        tracker.observe(vehicles, frame, controller)
    raise AssertionError(f"bus {bus.bus_id} never completed in {max_frames} frames")


def test_bus_event_logged_on_completion(tmp_path):
    log_path = tmp_path / "bus_events.jsonl"
    tracker = BusEventTracker(log_path)
    controller = make_controller()
    controller.phase = 0
    controller.timer = 0
    bus = make_bus_for_leg("R1_EB_A_NB", 300, "DONE_BUS")

    record, completion_frame = run_bus_to_completion(bus, controller, tracker)

    rows = read_rows(log_path)
    assert len(rows) == 1
    row = rows[0]
    assert row == record
    assert row["bus_id"] == "DONE_BUS"
    assert row["route_id"] == "R1_EB_A_NB"
    assert row["passengers"] == 45
    assert row["spawn_frame"] == 1
    assert row["spawn_sim_time"] == pytest.approx(1 / 60)
    assert row["completion_frame"] == completion_frame
    assert row["completion_sim_time"] == pytest.approx(completion_frame / 60)
    assert row["completion_frame"] > row["spawn_frame"]
    assert [node["node_x"] for node in row["nodes"]] == [300]
    node = row["nodes"][0]
    assert node["arrival_frame"] == 1
    assert node["stop_bar_cross_frame"] is not None
    assert node["node_clear_frame"] is not None
    assert (
        node["arrival_frame"]
        <= node["stop_bar_cross_frame"]
        <= node["node_clear_frame"]
        <= row["completion_frame"]
    )
    # The tracker forgets a completed bus so a re-used id starts fresh.
    assert tracker.active_record("DONE_BUS") is None


def test_tsp_treated_bus_flagged(tmp_path):
    """Same scenario as test_extension_holds_green_for_bus: EW green about
    to end with the bus still upstream, so the controller extends it."""
    log_path = tmp_path / "bus_events.jsonl"
    tracker = BusEventTracker(log_path)
    control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] = True
    controller = make_controller()
    controller.phase = 0
    controller.timer = GREEN_FRAMES - 3
    bus = make_bus_for_leg("R1_EB_A_NB", 300, "TREATED_BUS")

    record, _ = run_bus_to_completion(bus, controller, tracker)

    history = controller.get_node_status(300)["terminal_history"]
    assert history[-1]["tsp_action"] == TSP_ACTION_EXTENDING
    node = record["nodes"][0]
    assert node["tsp_treated"] is True
    assert node["tsp_action"] == TSP_ACTION_EXTENDING
    assert node["tsp_adjust_frames"] == history[-1]["tsp_adjust_frames"]
    assert node["tsp_adjust_frames"] >= 1
    row = read_rows(log_path)[0]
    assert row["nodes"][0]["tsp_treated"] is True
    flat = dict(zip(BUS_EVENT_HEADERS, flatten_bus_event(row)))
    assert flat["tsp_treated_any"] is True
    assert flat["node1_tsp_action"] == TSP_ACTION_EXTENDING


def test_untreated_bus_flagged(tmp_path):
    log_path = tmp_path / "bus_events.jsonl"
    tracker = BusEventTracker(log_path)
    control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] = False
    controller = make_controller()
    controller.phase = 0
    controller.timer = GREEN_FRAMES - 3
    bus = make_bus_for_leg("R1_EB_A_NB", 300, "PLAIN_BUS")

    record, _ = run_bus_to_completion(bus, controller, tracker)

    node = record["nodes"][0]
    assert node["tsp_treated"] is False
    assert node["tsp_action"] == TSP_ACTION_NONE
    assert node["tsp_adjust_frames"] == 0
    flat = dict(zip(BUS_EVENT_HEADERS, flatten_bus_event(record)))
    assert flat["tsp_treated_any"] is False


def test_wait_frames_accumulate(tmp_path):
    tracker = BusEventTracker(tmp_path / "bus_events.jsonl")

    # Free-flowing: EW green for the whole approach, nothing ahead.
    free_controller = make_controller()
    free_controller.phase = 0
    free_controller.timer = 0
    free_bus = make_bus_for_leg("R1_EB_A_NB", 300, "FREE_BUS")
    free_record, _ = run_bus_to_completion(free_bus, free_controller, tracker)

    # Held: NS green has just started, so the EB bus must wait at the bar.
    held_controller = make_controller()
    held_controller.phase = 3
    held_controller.timer = 0
    held_bus = make_bus_for_leg("R1_EB_A_NB", 300, "HELD_BUS")
    held_record, _ = run_bus_to_completion(held_bus, held_controller, tracker)

    assert free_record["first_stop_frame"] is None
    assert free_record["total_wait_frames"] == 0
    assert free_record["nodes"][0]["node_wait_frames"] == 0

    assert held_record["first_stop_frame"] is not None
    assert held_record["nodes"][0]["node_wait_frames"] > 0
    assert held_record["total_wait_frames"] >= held_record["nodes"][0]["node_wait_frames"]
    assert (
        held_record["nodes"][0]["stop_bar_cross_frame"]
        > held_record["first_stop_frame"]
    )


def test_bus_events_cleared_on_reset(tmp_path, monkeypatch):
    bus_log = tmp_path / "bus_events.jsonl"
    bus_log.write_text('{"bus_id": "OLD"}\n', encoding="utf-8")
    monkeypatch.setattr(main, "BUS_EVENTS_LOG_PATH", bus_log)
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", tmp_path / "telemetry.jsonl")
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", tmp_path / "turns.jsonl")
    # An in-flight record must not survive RESET either.
    stale_bus = make_bus_for_leg("R1_EB_A_NB", 300, "STALE_BUS")
    main.bus_event_tracker.observe([stale_bus], 5, None)
    assert main.bus_event_tracker.active_record("STALE_BUS") is not None

    main.reset_session_logs()

    assert not bus_log.exists()
    assert main.bus_event_tracker.active_record("STALE_BUS") is None
    # The tracker follows the re-pointed path, so the next completion lands
    # in the fresh log rather than the module default.
    assert main.bus_event_tracker.path == bus_log


def test_bus_events_in_export(tmp_path, monkeypatch):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    bus_log = tmp_path / "bus_events.jsonl"
    tracker = BusEventTracker(bus_log)
    control_panel.bus_routes_config["R1_EB_A_NB"]["tsp_enabled"] = True
    controller = make_controller()
    controller.phase = 0
    controller.timer = GREEN_FRAMES - 3
    bus = make_bus_for_leg("R1_EB_A_NB", 300, "EXPORT_BUS")
    record, _ = run_bus_to_completion(bus, controller, tracker)

    monkeypatch.setattr(main, "BUS_EVENTS_LOG_PATH", bus_log)
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", tmp_path / "telemetry.jsonl")
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", tmp_path / "turns.jsonl")
    destination = tmp_path / "test_export.xlsx"

    assert main.export_test_workbook("None", 60, 1, destination) == destination

    workbook = load_workbook(destination, data_only=True)
    try:
        assert "Bus Events" in workbook.sheetnames
        sheet = workbook["Bus Events"]
        assert [cell.value for cell in sheet[1]] == BUS_EVENT_HEADERS
        assert sheet.max_row == 2
        row = dict(zip(BUS_EVENT_HEADERS, [cell.value for cell in sheet[2]]))
        assert row["bus_id"] == "EXPORT_BUS"
        assert row["passengers"] == 45
        assert row["tsp_treated_any"] is True
        assert row["node1_tsp_action"] == TSP_ACTION_EXTENDING
        assert row["node1_tsp_adjust_frames"] == record["nodes"][0]["tsp_adjust_frames"]
        assert row["completion_frame"] == record["completion_frame"]
    finally:
        workbook.close()


def test_tracker_never_raises_on_bad_input(tmp_path):
    tracker = BusEventTracker(tmp_path / "missing_dir" / "bus_events.jsonl")
    bus = make_bus_for_leg("R1_EB_A_NB", 300, "ROBUST_BUS")
    # Bad frame, missing controller, unwritable path: all swallowed.
    tracker.observe([bus, object()], "not-a-frame", None)
    tracker.observe([bus], 1, None)
    assert tracker.complete(bus, 2, None) is not None
    tracker.complete(object(), 3, None)
