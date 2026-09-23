"""Gridlock monitor: onset/clear detection, stalled-head diagnosis, report."""
from openpyxl import load_workbook

from src.telemetry.gridlock_monitor import (
    GRIDLOCK_SUSTAIN_S,
    GridlockMonitor,
    diagnose_stalled_heads,
    report_path_for,
)
from src.ui.canvas_gemini import H_Y, LANE
from tests.helpers import NODE_A, NODE_B


def wb_y(lane):
    return H_Y + (lane + 0.5) * LANE


def vehicle(sid, x, y, speed=0.0, vtype="car", **extra):
    row = {
        "snapshot_id": sid, "vehicle_type": vtype, "route_id": None,
        "x_px": x, "y_px": y, "direction": "WB", "lane_index": 0,
        "target_turn": "STRAIGHT", "speed_px_per_frame": speed,
        "passed_nodes": [NODE_B], "assigned_node_x": None,
        "lane_vacate_target": None, "must_hold_for_lane": False,
        "route_merge_hold_active": False,
    }
    row.update(extra)
    return row


def payload(t, vehicles, wb_signal="RED", wb_blocked=False):
    node = lambda: {
        "phase": "ALL_RED", "phase_timer_frames": 10, "priority_state": "NORMAL",
        "signals": {"EB": "RED", "WB": wb_signal, "NB": "RED", "SB": "RED"},
        "queues": {"EB": 0, "WB": len(vehicles), "NB": 0, "SB": 0},
        "downstream_blocked": {"EB": False, "WB": wb_blocked, "NB": False, "SB": False},
    }
    return {
        "simulation_time_seconds": t, "frame_number": int(t * 60),
        "signal_state": {"nodes": {str(NODE_A): node(), str(NODE_B): node()}},
        "network_summary": {"total_vehicles": len(vehicles), "mean_speed_px_per_frame": 0.0},
        "vehicle_positions": vehicles,
    }


def test_deadlocked_queue_head_is_diagnosed_and_a_follower_is_not():
    # The live snapshot: left turners holding 35 px short of Node A's bar,
    # a truck frozen straddling lanes 1/2, and the queue behind them.
    vehicles = [
        vehicle("car-1", 919.8, wb_y(0), target_turn="LEFT", lane_vacate_target=2, must_hold_for_lane=True),
        vehicle("truck-10", 924.6, wb_y(1) + 12.5, vtype="truck", lane_index=1, target_turn="LEFT",
                lane_vacate_target=2, must_hold_for_lane=True),
        vehicle("truck-2", 954.7, wb_y(0), vtype="truck"),
        vehicle("car-4", 989.7, wb_y(0)),
    ]
    heads = diagnose_stalled_heads(payload(100.0, vehicles))
    by_id = {h["vehicle"]: h for h in heads}
    assert set(by_id) == {"car-1", "truck-10"}
    assert by_id["truck-10"]["straddling_lanes"] is True
    assert by_id["car-1"]["straddling_lanes"] is False
    assert by_id["car-1"]["must_hold_for_lane"] is True
    assert 30 < by_id["car-1"]["dist_to_bar_px"] < 40
    assert "short of bar" in by_id["car-1"]["reason"]


def test_head_waiting_at_red_is_not_a_finding_but_spillback_hold_is():
    at_bar = [vehicle("car-9", NODE_A + 66 + 10 + 9 + 14, wb_y(0))]
    assert diagnose_stalled_heads(payload(1.0, at_bar)) == []
    held = diagnose_stalled_heads(payload(1.0, at_bar, wb_signal="GREEN", wb_blocked=True))
    assert len(held) == 1 and "spillback" in held[0]["reason"]


def test_gridlock_onset_needs_sustained_stopped_share_and_clears(tmp_path):
    monitor = GridlockMonitor()
    stopped = [vehicle(f"c{i}", 1000 + 30 * i, wb_y(i % 3)) for i in range(50)]
    monitor.sample(payload(0.0, stopped))
    monitor.sample(payload(GRIDLOCK_SUSTAIN_S - 1, stopped))
    assert monitor.state == "FREE"
    monitor.sample(payload(GRIDLOCK_SUSTAIN_S + 1, stopped))
    assert monitor.state == "GRIDLOCK" and monitor.onset_time_s == GRIDLOCK_SUSTAIN_S + 1
    moving = [dict(v, speed_px_per_frame=0.5) for v in stopped]
    monitor.sample(payload(GRIDLOCK_SUSTAIN_S + 30, moving))
    assert monitor.state == "FREE"
    assert [e[1] for e in monitor.events] == ["GRIDLOCK_ONSET", "GRIDLOCK_CLEARED"]
    assert monitor.gridlock_seconds == 29.0

    report = monitor.write_workbook(report_path_for(tmp_path / "run.xlsx"))
    assert report.name == "run_gridlock.xlsx"
    wb = load_workbook(report)
    assert wb.sheetnames == ["Summary", "Events", "Stalled Heads", "Timeline"]
    assert wb["Timeline"].max_row == 5  # header + 4 samples
    assert dict(wb["Summary"].iter_rows(min_row=2, values_only=True))["seconds_in_gridlock"] == 29.0

    monitor.reset()
    assert monitor.timeline == [] and monitor.state == "FREE"
