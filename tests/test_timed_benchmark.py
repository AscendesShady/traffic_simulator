"""Automated timed benchmark: one clean seeded run, then an auto-export.

The run length is measured on the SIMULATION clock, so a 30-minute test means
the same amount of simulated traffic whether the host renders it quickly or
slowly. The export runs inside the simulator process, reading the same
crash-safe JSONL logs the dashboard reads, so no cross-process handshake is
needed to finish a test.
"""
import inspect
import json
import re

import pytest

import control_panel
import main


DURATION_LABELS = ("5 min", "10 min", "15 min", "30 min", "1 hr", "2 hr")


def write_logs(tmp_path, monkeypatch, turns=2, telemetry_rows=2):
    """Point the exporters at throwaway logs holding a known session."""
    turn_log = tmp_path / "agent_turn_log.jsonl"
    telemetry_log = tmp_path / "telemetry_log.jsonl"
    with turn_log.open("w", encoding="utf-8") as handle:
        for turn in range(1, turns + 1):
            handle.write(
                json.dumps(
                    {
                        "turn": turn,
                        "timestamp": 1000.0 + turn,
                        "model": "nemotron-3-nano-4b",
                        "status": "OK",
                        "reason": "throughput",
                        "latency_ms": 100.0 * turn,
                        "input_tokens": 10,
                        "output_tokens": 20,
                        "tokens_per_sec": 5.0,
                        "pax_per_min_recent": 42.0,
                        "flags": {
                            route_id: {"tsp": True, "dbl": False}
                            for route_id in main.SESSION_ROUTE_IDS
                        },
                    }
                )
                + "\n"
            )
    with telemetry_log.open("w", encoding="utf-8") as handle:
        for index in range(telemetry_rows):
            handle.write(
                json.dumps(
                    {"frame": index * 60, "sim_time_s": float(index)}
                )
                + "\n"
            )
    monkeypatch.setattr(main, "AGENT_TURN_LOG_PATH", turn_log)
    monkeypatch.setattr(main, "TELEMETRY_LOG_PATH", telemetry_log)
    monkeypatch.setattr(main, "EXCEL_EXPORT_DIR", tmp_path / "excel_exports")
    return turn_log, telemetry_log


def test_duration_dropdown_offers_the_agreed_durations():
    assert list(control_panel.TEST_DURATIONS) == list(DURATION_LABELS)
    assert control_panel.TEST_DURATIONS == {
        "5 min": 300,
        "10 min": 600,
        "15 min": 900,
        "30 min": 1800,
        "1 hr": 3600,
        "2 hr": 7200,
    }


def test_duration_required(monkeypatch):
    """START TEST with no duration selected must do nothing."""
    monkeypatch.setitem(
        control_panel.global_config, "test_duration_sim_seconds", None
    )
    monkeypatch.setitem(control_panel.global_config, "start_requested", False)
    monkeypatch.setitem(control_panel.global_config, "test_running", False)

    assert control_panel.request_start_test() is None
    assert control_panel.global_config["start_requested"] is False
    assert control_panel.global_config["test_running"] is False


def test_start_test_does_full_reset(tmp_path, monkeypatch):
    """START TEST routes through the same clean-reset path as START."""
    monkeypatch.setitem(
        control_panel.global_config, "test_duration_sim_seconds", 300
    )
    monkeypatch.setitem(control_panel.global_config, "random_seed", 43)
    monkeypatch.setitem(
        control_panel.global_config, "ai_runtime", {"model": "gemma4:12b"}
    )
    monkeypatch.setitem(control_panel.global_config, "start_requested", False)

    assert control_panel.request_start_test() == 300
    # The existing full-reset-and-run path is what actually clears state.
    assert control_panel.global_config["start_requested"] is True
    assert control_panel.global_config["test_running"] is True
    assert control_panel.global_config["test_model"] == "gemma4:12b"
    assert control_panel.global_config["test_seed"] == 43

    # And that path performs a real frame-zero reset.
    turn_log, telemetry_log = write_logs(tmp_path, monkeypatch)
    monkeypatch.setattr(main, "reset_traffic_generation", lambda: None)

    class FakeSignals:
        frame_number = 900

        def reset_all_state(self):
            self.frame_number = 0

    signals = FakeSignals()
    vehicles = [object(), object()]
    for key in main.network_throughput:
        main.network_throughput[key] = 77

    assert main.perform_full_reset(vehicles, signals) == 0
    assert vehicles == []
    assert signals.frame_number == 0
    assert not turn_log.exists() and not telemetry_log.exists()
    assert set(main.network_throughput.values()) == {0}


def test_test_stops_at_sim_duration(tmp_path, monkeypatch):
    """The stop is driven by simulation_time_seconds, not wall-clock time."""
    write_logs(tmp_path, monkeypatch)
    monkeypatch.setitem(control_panel.global_config, "test_running", True)
    monkeypatch.setitem(
        control_panel.global_config, "test_duration_sim_seconds", 300
    )
    monkeypatch.setitem(control_panel.global_config, "is_running", True)
    monkeypatch.setitem(control_panel.global_config, "test_model", "None")
    monkeypatch.setitem(control_panel.global_config, "test_seed", 7)
    monkeypatch.setattr(control_panel, "write_ai_control", lambda *a, **k: None)

    # One frame short of the threshold the run continues.
    assert main.timed_test_is_complete(300 * 60 - 1) is False
    # Crossing the sim-time threshold ends it.
    assert main.timed_test_is_complete(300 * 60) is True

    # Wall-clock time is irrelevant: only the frame count moves the decision.
    monkeypatch.setattr(main.time, "time", lambda: 1e12)
    assert main.timed_test_is_complete(300 * 60 - 1) is False

    # A free run (no timed test) never auto-stops, however long it runs.
    monkeypatch.setitem(control_panel.global_config, "test_running", False)
    assert main.timed_test_is_complete(10 ** 7) is False
    monkeypatch.setitem(control_panel.global_config, "test_running", True)

    main.finish_timed_test(300 * 60)

    assert control_panel.global_config["is_running"] is False
    assert control_panel.global_config["test_running"] is False


def test_test_autoexports_with_correct_filename(tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    write_logs(tmp_path, monkeypatch)
    monkeypatch.setitem(control_panel.global_config, "test_running", True)
    monkeypatch.setitem(
        control_panel.global_config, "test_duration_sim_seconds", 1800
    )
    monkeypatch.setitem(
        control_panel.global_config, "test_model", "nemotron-3-nano-4b"
    )
    monkeypatch.setitem(control_panel.global_config, "test_seed", 43)
    monkeypatch.setattr(control_panel, "write_ai_control", lambda *a, **k: None)

    destination = main.finish_timed_test(1800 * 60)

    assert destination is not None and destination.exists()
    assert re.fullmatch(
        r"test_nemotron-3-nano-4b_30min_seed43_\d{8}_\d{6}\.xlsx",
        destination.name,
    ), destination.name
    assert control_panel.global_config["test_last_export"] == destination.name

    workbook = load_workbook(destination)
    assert workbook.sheetnames == [
        "Decisions",
        "Telemetry",
        "LLM Performance",
        "LLM Summary",
        "Control Panel Inputs",
    ]
    # The benchmark file records the exact configuration that produced it.
    inputs = workbook["Control Panel Inputs"]
    assert [cell.value for cell in inputs[1]] == ["section", "parameter", "value"]
    recorded = {
        (row[0].value, row[1].value): row[2].value
        for row in inputs.iter_rows(min_row=2)
    }
    assert ("Global", "random_seed") in recorded
    assert any(key[0].startswith("Approach: ") for key in recorded)
    assert any(key[0].startswith("Route: ") for key in recorded)
    performance = workbook["LLM Performance"]
    assert performance.max_row == 3  # header plus two logged turns
    summary = workbook["LLM Summary"]
    assert summary.cell(row=2, column=1).value == "nemotron-3-nano-4b"
    assert summary.cell(row=2, column=2).value == 2
    workbook.close()


def test_baseline_test_runs_with_model_none(tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    write_logs(tmp_path, monkeypatch)
    monkeypatch.setitem(control_panel.global_config, "test_running", True)
    monkeypatch.setitem(
        control_panel.global_config, "test_duration_sim_seconds", 600
    )
    monkeypatch.setitem(control_panel.global_config, "test_model", "None")
    monkeypatch.setitem(control_panel.global_config, "test_seed", None)
    monkeypatch.setattr(control_panel, "write_ai_control", lambda *a, **k: None)

    destination = main.finish_timed_test(600 * 60)

    assert destination is not None and destination.exists()
    assert re.fullmatch(
        r"test_baseline_10min_seedNone_\d{8}_\d{6}\.xlsx", destination.name
    ), destination.name


def test_manual_stop_during_test_no_autoexport(tmp_path, monkeypatch):
    """STOP interrupts a test: sim halts, logs survive, nothing is exported."""
    turn_log, telemetry_log = write_logs(tmp_path, monkeypatch)
    export_dir = tmp_path / "excel_exports"
    monkeypatch.setitem(control_panel.global_config, "is_running", True)
    monkeypatch.setitem(control_panel.global_config, "test_running", True)
    monkeypatch.setitem(
        control_panel.global_config, "test_duration_sim_seconds", 1800
    )

    assert control_panel.request_start_stop() == "STOPPED"

    assert control_panel.global_config["is_running"] is False
    assert control_panel.global_config["test_running"] is False
    # Logs are preserved for a manual EXPORT ALL.
    assert turn_log.exists() and telemetry_log.exists()
    # No workbook was written by the interrupted test.
    assert not export_dir.exists() or not list(export_dir.glob("test_*.xlsx"))


def test_filename_sanitises_model_separators():
    name = main.build_test_export_filename(
        "vendor/family:8b", 900, 5, timestamp="20260908_143005"
    )
    assert name == "test_vendor-family-8b_15min_seed5_20260908_143005.xlsx"


def test_arming_without_a_model_reports_baseline_not_waiting(monkeypatch):
    """Arming with model 'None' is a baseline run, not a pending decision."""
    monkeypatch.setitem(
        control_panel.global_config,
        "ai_runtime",
        {"armed": False, "model": "None", "tick_seconds": 5, "last_turn": 0},
    )
    monkeypatch.setattr(control_panel, "write_ai_control", lambda *a, **k: None)
    runtime = control_panel.global_config["ai_runtime"]

    # merge_ai_decision must not claim to be waiting for a decision, and must
    # not apply a stale decision file left over from an earlier model.
    runtime["armed"] = True
    assert main.merge_ai_decision() is False
    assert runtime["last_status"] == control_panel.BASELINE_NO_MODEL

    # Selecting None while armed reports baseline too.
    control_panel.set_active_ai_model("None", persist=False)
    assert runtime["last_status"] == control_panel.BASELINE_NO_MODEL

    # A real model still uses the ordinary waiting path.
    control_panel.set_active_ai_model("gemma4:12b", persist=False)
    assert runtime["last_status"] == "MODEL_CHANGED_WAITING"


def test_baseline_run_ignores_a_stale_decision_file(tmp_path, monkeypatch):
    """A leftover decision must never steer an unaided baseline benchmark."""
    stale = tmp_path / "decision.json"
    stale.write_text(
        json.dumps(
            {
                "timestamp": 9e12,
                "status": "OK",
                "flags": {
                    route_id: {"tsp": True, "dbl": True}
                    for route_id in main.SESSION_ROUTE_IDS
                },
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setitem(
        control_panel.global_config,
        "ai_runtime",
        {"armed": True, "model": "None", "tick_seconds": 5, "last_turn": 0},
    )
    for route_id in main.SESSION_ROUTE_IDS:
        control_panel.bus_routes_config[route_id]["tsp_enabled"] = False
        control_panel.bus_routes_config[route_id]["dbl_enabled"] = False

    assert main.merge_ai_decision(stale) is False

    for route_id in main.SESSION_ROUTE_IDS:
        assert control_panel.bus_routes_config[route_id]["tsp_enabled"] is False
        assert control_panel.bus_routes_config[route_id]["dbl_enabled"] is False


def test_countdown_counts_down_in_sim_time():
    """The countdown tracks the simulation clock, not the wall clock."""
    fmt = control_panel.format_test_countdown
    assert fmt(900, 0) == "15:00"
    assert fmt(900, 60) == "14:00"
    assert fmt(900, 899) == "00:01"
    # It never goes negative once the auto-stop threshold is reached.
    assert fmt(900, 900) == "00:00"
    assert fmt(900, 5000) == "00:00"
    # A run that has not advanced yet still shows the full duration.
    assert fmt(300, -5) == "05:00"


def test_countdown_absent_without_a_duration():
    fmt = control_panel.format_test_countdown
    assert fmt(None, 10) is None
    assert fmt(0, 10) is None
    assert fmt("bad", 10) is None
    assert fmt(900, None) is None


def test_sim_clock_is_published_for_the_countdown(monkeypatch):
    """main.py must publish the sim clock the panel counts down from."""
    monkeypatch.setitem(control_panel.global_config, "sim_time_seconds", 0.0)
    source = inspect.getsource(main.main)
    assert 'control_panel.global_config["sim_time_seconds"]' in source
    assert "master_frame_count / 60.0" in source
    # The published key exists by default, so the panel never reads a hole.
    assert "sim_time_seconds" in control_panel.global_config


def test_control_panel_input_rows_cover_every_operator_setting():
    rows = main.control_panel_input_rows()
    keys = {(section, parameter) for section, parameter, _value in rows}

    for parameter in (
        "random_seed",
        "sim_speed",
        "test_duration_sim_seconds",
        "discharge_selection",
        "llm_model",
        "llm_tick_seconds",
    ):
        assert ("Global", parameter) in keys

    # Every approach and every route is recorded.
    for name in control_panel.APPROACH_NAMES.values():
        assert (f"Approach: {name}", "inflow_rate_veh_per_min") in keys
    for route_id in control_panel.bus_routes_config:
        assert (f"Route: {route_id}", "headway_sec") in keys
        assert (f"Route: {route_id}", "tsp_enabled") in keys


def test_control_panel_inputs_reflect_live_edits(monkeypatch):
    monkeypatch.setitem(control_panel.global_config, "random_seed", 4242)
    monkeypatch.setitem(control_panel.approach_configs["EB"], "rate", 27)

    recorded = {
        (section, parameter): value
        for section, parameter, value in main.control_panel_input_rows()
    }

    assert recorded[("Global", "random_seed")] == 4242
    assert recorded[("Approach: EB Corridor", "inflow_rate_veh_per_min")] == 27
