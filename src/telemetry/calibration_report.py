"""Calibration report: python -m src.telemetry.calibration_report [CAMPAIGN_ID] [--results DIR]

One self-contained HTML page per campaign: how every regime in it was
calibrated (the Calibration sheet each run workbook carries, built by
main.calibration_rows) and whether each run stayed inside that calibration
(the validation columns of its summary row, judged against the thresholds
print_campaign_summary uses). print_campaign_summary writes it at the end of
every batch, windowed or headless; without a campaign id the CLI takes the
most recent one.
"""
import argparse
import csv
import datetime
import html
import pathlib

from src.core import main


def _number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


# (summary column, label, format, fails(value) or None = reported, not judged)
RUN_CHECKS = (
    ("network_to_calibrated_s_ratio", "In-network S / calibrated S", "{:.2f}",
     lambda v: v < main.NETWORK_S_WARN_RATIO),
    ("network_saturation_headway_samples", "Headways sampled", "{:.0f}",
     lambda v: v < main.NETWORK_S_MIN_SAMPLES),
    ("latent_demand_share_at_end", "Latent demand", "{:.1%}",
     lambda v: v > main.LATENT_DEMAND_WARN_SHARE),
    ("geh_entry_share_below_5", "Sources GEH < 5", "{:.0%}",
     lambda v: v < main.GEH_ACCEPT_SHARE),
    ("converged", "Converged", None, lambda v: v != "True"),
    ("ssm_ttc_conflicts_per_veh_hr_steady", "TTC < 1.5 s /veh-h", "{:.2f}", None),
    ("ssm_emergency_decel_per_veh_hr_steady", "Emergency braking /veh-h", "{:.2f}", None),
    ("missed_turns_total", "Missed turns", "{:.0f}", None),
)

CHECK_NOTES = (
    ("In-network S / calibrated S", f"at least {main.NETWORK_S_WARN_RATIO:.2f}: saturation headways "
     "measured at the stop bars during the run (HCM Ch. 31) against the S Webster was timed for. "
     f"Needs {main.NETWORK_S_MIN_SAMPLES} or more headways to mean anything."),
    ("Latent demand", f"at most {main.LATENT_DEMAND_WARN_SHARE:.0%} of offered vehicles still waiting "
     "at the boundary at the end (FHWA Traffic Analysis Toolbox Vol. III)."),
    ("Sources GEH < 5", f"at least {main.GEH_ACCEPT_SHARE:.0%} of sources, entered against offered "
     "volume (FHWA Traffic Analysis Toolbox Vol. III)."),
    ("Converged", "cumulative pax/min within 5 % and vehicles in the network within 10 % "
     "between 0.8 T and T."),
    ("Safety columns", "reported, not judged: TTC conflicts use FHWA SSAM's 1.5 s; emergency braking "
     "is the SUMO emergencyDecel bound; a missed turn is a vehicle that never reached its turn lane."),
)

CSS = """
:root { --bg:#f6f7f9; --card:#fff; --text:#1c2230; --muted:#5b6475; --line:#dde1e8;
        --accent:#2f6fde; --flag:#b42318; --flagbg:#fdecea; --ok:#1a7f37; --diff:#fff7e0; }
@media (prefers-color-scheme: dark) {
  :root { --bg:#12151c; --card:#1b202a; --text:#e6e9ef; --muted:#9aa3b5; --line:#2c3340;
          --accent:#6ea0ff; --flag:#ff8a80; --flagbg:#3a1d1d; --ok:#56d364; --diff:#2e2a1a; }
}
* { box-sizing: border-box; }
body { margin:0; background:var(--bg); color:var(--text);
       font:14px/1.5 "Segoe UI", system-ui, sans-serif; }
main { max-width:1200px; margin:0 auto; padding:24px 16px 48px; }
h1 { font-size:24px; margin:0 0 4px; }
h2 { font-size:18px; margin:32px 0 8px; }
h3 { font-size:15px; margin:20px 0 6px; color:var(--muted); font-weight:600; }
p.meta { color:var(--muted); margin:0 0 16px; }
.card { background:var(--card); border:1px solid var(--line); border-radius:8px;
        overflow-x:auto; margin-bottom:12px; }
table { border-collapse:collapse; width:100%; }
th, td { text-align:left; padding:6px 10px; border-bottom:1px solid var(--line); vertical-align:top; }
th { color:var(--muted); font-weight:600; background:var(--card); white-space:nowrap; }
tr:last-child td { border-bottom:none; }
td.num { font-variant-numeric: tabular-nums; white-space:nowrap; }
td.basis { color:var(--muted); }
tr.differs td { background:var(--diff); }
td.flag { color:var(--flag); background:var(--flagbg); font-weight:600; }
td.ok { color:var(--ok); }
ul { margin:4px 0 0; padding-left:20px; color:var(--muted); }
.summary { font-weight:600; margin:0 0 8px; }
"""


def _cell(value):
    if isinstance(value, bool) or value is None:
        return "" if value is None else str(value)
    if isinstance(value, float):
        return f"{value:g}"
    return str(value)


def workbook_calibration_rows(path):
    """The Calibration sheet of one run workbook, without its header."""
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Calibration" not in workbook.sheetnames:
            return []
        rows = list(workbook["Calibration"].iter_rows(values_only=True))
    finally:
        workbook.close()
    return [tuple(row) for row in rows[1:]]


def campaign_rows(campaign_id=None, results_dir=None):
    """(campaign_id, its final-checkpoint summary rows) from every summary CSV
    in the results directory; the latest campaign when none is named."""
    results_dir = pathlib.Path(results_dir or main.RESULTS_DIR)
    rows = []
    for path in sorted(results_dir.glob("experiment_summary*.csv")):
        with path.open(encoding="utf-8", newline="") as handle:
            rows.extend(
                row for row in csv.DictReader(handle)
                if row.get("campaign_id") and row.get("checkpoint_min") == row.get("test_duration_min")
            )
    if campaign_id is None and rows:
        campaign_id = max(rows, key=lambda r: r.get("run_ended_utc") or r.get("timestamp") or "")["campaign_id"]
    return campaign_id, [row for row in rows if row.get("campaign_id") == campaign_id]


def _regime_tables(regimes):
    """One table per section: parameter/value/unit/basis where every regime
    agrees, else a column per regime with the differing rows shaded."""
    order, cells = [], {}
    for index, (_, rows) in enumerate(regimes):
        for section, parameter, value, unit, basis in rows:
            key = (section, parameter)
            if key not in cells:
                order.append(key)
                cells[key] = {"unit": unit or "", "basis": basis or "", "values": {}}
            cells[key]["values"][index] = value
    sections = {}
    for key in order:
        sections.setdefault(key[0], []).append(key)
    out = []
    for section, keys in sections.items():
        many = len(regimes) > 1 and any(
            len({_cell(cells[key]["values"].get(i)) for i in range(len(regimes))}) > 1 for key in keys
        )
        out.append(f"<h3>{html.escape(str(section))}</h3><div class='card'><table>")
        if many:
            heads = "".join(f"<th>{html.escape(label)}</th>" for label, _ in regimes)
            out.append(f"<tr><th>Parameter</th><th>Unit</th>{heads}<th>Basis</th></tr>")
        else:
            out.append("<tr><th>Parameter</th><th>Value</th><th>Unit</th><th>Basis</th></tr>")
        for key in keys:
            cell = cells[key]
            values = [_cell(cell["values"].get(i)) for i in range(len(regimes))]
            same = len(set(values)) == 1
            if not many:
                value_cells = f"<td class='num'>{html.escape(values[0])}</td>"
            elif same:
                value_cells = f"<td class='num' colspan='{len(regimes)}'>{html.escape(values[0])}</td>"
            else:
                value_cells = "".join(f"<td class='num'>{html.escape(v)}</td>" for v in values)
            unit = f"<td>{html.escape(str(cell['unit']))}</td>"
            basis = f"<td class='basis'>{html.escape(str(cell['basis']))}</td>"
            name = f"<td>{html.escape(str(key[1]))}</td>"
            row_class = " class='differs'" if many and not same else ""
            out.append(
                f"<tr{row_class}>{name}{unit}{value_cells}{basis}</tr>" if many
                else f"<tr>{name}{value_cells}{unit}{basis}</tr>"
            )
        out.append("</table></div>")
    return "".join(out)


def _checks_table(runs, checks):
    heads = "".join(f"<th>{html.escape(label)}</th>" for _, label, _, _ in checks)
    body, passing = [], 0
    for label, values in runs:
        cells, failed = [], False
        for column, _, fmt, fails in checks:
            raw = values.get(column)
            number = _number(raw)
            if fmt is None:
                shown, judged = ("" if raw in (None, "") else str(raw)), raw
            elif number is None:
                shown, judged = "", None
            else:
                shown, judged = fmt.format(number), number
            css = "num"
            if fails is not None and judged not in (None, ""):
                bad = bool(fails(judged))
                failed = failed or bad
                css += " flag" if bad else " ok"
            cells.append(f"<td class='{css}'>{html.escape(shown) or '—'}</td>")
        passing += not failed
        body.append(f"<tr><td>{html.escape(label)}</td>{''.join(cells)}</tr>")
    summary = f"<p class='summary'>{passing} of {len(runs)} runs pass every judged check.</p>"
    return f"{summary}<div class='card'><table><tr><th>Run</th>{heads}</tr>{''.join(body)}</table></div>"


def render(title, meta_lines, regimes, runs, checks=RUN_CHECKS, check_notes=CHECK_NOTES):
    """The page. ``regimes`` is [(label, calibration rows)], ``runs`` is
    [(label, {column: value})] judged by ``checks``."""
    meta = "<br>".join(html.escape(line) for line in meta_lines)
    notes = "".join(f"<li><b>{html.escape(a)}</b>: {html.escape(b)}</li>" for a, b in check_notes)
    missing = [label for label, rows in regimes if not rows]
    missing_note = (
        f"<p class='meta'>No Calibration sheet for: {html.escape(', '.join(missing))} "
        "(workbook missing or older than the sheet).</p>" if missing else ""
    )
    return (
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>"
        "<meta name='viewport' content='width=device-width, initial-scale=1'>"
        f"<title>{html.escape(title)}</title><style>{CSS}</style></head><body><main>"
        f"<h1>{html.escape(title)}</h1><p class='meta'>{meta}</p>"
        "<h2>How the simulation was calibrated</h2>"
        "<p class='meta'>Every value below was frozen at the start of each run, before the clock "
        "started (main.calibrate_and_apply_webster), and is the same for every run of one regime.</p>"
        f"{missing_note}{_regime_tables([r for r in regimes if r[1]])}"
        "<h2>Did each run hold to it</h2>"
        f"{_checks_table(runs, checks)}<ul>{notes}</ul>"
        "</main></body></html>"
    )


def write_campaign_report(campaign_id=None, results_dir=None, out=None):
    """Write calibration_report_<campaign>.html beside the summary CSV and
    return its path (None when the campaign has no rows)."""
    results_dir = pathlib.Path(results_dir or main.RESULTS_DIR)
    campaign_id, rows = campaign_rows(campaign_id, results_dir)
    if not rows:
        return None
    by_hash = {}
    for row in rows:
        by_hash.setdefault(row.get("config_hash") or "?", []).append(row)
    regimes, label_of = [], {}
    for index, (config_hash, members) in enumerate(by_hash.items()):
        label = f"regime {index + 1} ({config_hash[:8]})" if len(by_hash) > 1 else f"config {config_hash[:8]}"
        calibration = []
        for row in members:
            path = results_dir / (row.get("workbook_filename") or "")
            if row.get("workbook_filename") and path.exists():
                calibration = workbook_calibration_rows(path)
                if calibration:
                    break
        regimes.append((label, calibration))
        label_of[config_hash] = label
    runs = [
        (
            f"{row.get('model')} / seed {row.get('seed')}"
            + (f" / {label_of[row.get('config_hash') or '?']}" if len(by_hash) > 1 else ""),
            row,
        )
        for row in sorted(rows, key=lambda r: (r.get("model") or "", r.get("seed") or ""))
    ]
    arms = sorted({row.get("model") for row in rows})
    seeds = sorted({row.get("seed") for row in rows})
    shas = sorted({row.get("git_sha") or "?" for row in rows})
    meta = [
        f"Campaign {campaign_id} · {len(rows)} runs · {len(arms)} arms × {len(seeds)} seeds · "
        f"{rows[0].get('test_duration_min')} min each",
        f"Arms: {', '.join(arms)}",
        f"Seeds: {', '.join(seeds)}",
        f"git {', '.join(s[:10] for s in shas)}"
        + (" (dirty tree)" if any(r.get("git_dirty") == "True" for r in rows) else "")
        + f" · {len(by_hash)} regime{'s' if len(by_hash) > 1 else ''}"
        + ("" if len(by_hash) == 1 else " -- NOT one campaign's regime"),
        f"Generated {datetime.datetime.now():%Y-%m-%d %H:%M}",
    ]
    out = pathlib.Path(out) if out else results_dir / f"calibration_report_{campaign_id}.html"
    out.write_text(render(f"Calibration report · {campaign_id[:8]}", meta, regimes, runs), encoding="utf-8")
    return out


def main_cli(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("campaign", nargs="?", default=None, help="campaign id (default: latest)")
    parser.add_argument("--results", default=None, help="results directory (default results/)")
    parser.add_argument("--out", default=None, help="output HTML path")
    args = parser.parse_args(argv)
    path = write_campaign_report(args.campaign, args.results, args.out)
    print(path or "no summary rows found for that campaign")


if __name__ == "__main__":
    main_cli()
